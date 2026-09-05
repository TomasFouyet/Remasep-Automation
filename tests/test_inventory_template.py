"""Tests del inventario de plantilla y alineación (Sprint 1.4).

Workbooks sintéticos. El valor ``SECRET_VALUE`` en una fila de datos del
generador nunca debe aparecer en los outputs.
"""

import csv
import json
import zipfile

import inventory_template as it
import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

SECRET = "SECRET_VALUE"


@pytest.fixture
def template_workbook(tmp_path):
    wb = Workbook()

    nombre = wb.active
    nombre.title = "NOMBRE"
    nombre["A1"] = "SEREMI"
    nombre["A2"] = "Establecimiento"
    nombre["A5"] = "ACTIVIDADES"  # duplica la etiqueta -> ambigüedad
    nombre["C2"] = "cod"
    nombre["I2"] = '=IF(C2="","Falta","")'
    nombre["B2"].protection = Protection(locked=False)  # input desbloqueado
    nombre.protection.sheet = True

    form = wb.create_sheet("FORM")
    form["B1"] = "SECCIÓN A: CONTROLES"
    form["B2"] = "ACTIVIDADES"
    form["C2"] = "TOTAL"
    form["A20"] = "0601101"
    form["C15"] = "=SUM(D15:F15)"
    form["D15"].protection = Protection(locked=False)
    form["E15"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    dv = DataValidation(type="list", formula1='"a,b,c"')
    form.add_data_validation(dv)
    dv.add("F15")
    form["A3"] = "merged label"
    form.merge_cells("A3:B3")
    form.protection.sheet = True

    control = wb.create_sheet("CONTROL")
    control["A1"] = "=FORM!C15"
    control["A2"] = "=NOMBRE!I2"
    control["B1"] = "unlocked pero hoja sin proteger"
    control["B1"].protection = Protection(locked=False)

    path = tmp_path / "template.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def generator_workbook(tmp_path):
    wb = Workbook()
    g1 = wb.active
    g1.title = "REMASEP 01"
    g1["B1"] = "SECCIÓN A: CONTROLES"  # exact -> FORM!B1
    g1["B2"] = "ACTIVIDADES"  # exact x2 -> ambiguous
    g1["C2"] = "TOTAL"  # exact único -> FORM!C2
    g1["D2"] = "0601101"  # exact -> FORM!A20
    g1["E2"] = "Total General"  # sin match
    g1["F2"] = "de"  # demasiado corta -> ignorada

    detail = wb.create_sheet("Atenciones - Detalles de citas")
    detail["A1"] = "ID CITA"
    detail["A2"] = SECRET
    detail["B2"] = SECRET

    path = tmp_path / "generator.xlsx"
    wb.save(path)
    return path


def _rows(path):
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _run(template, tmp_path, generator=None, **kwargs):
    out = tmp_path / "ti"
    doc = tmp_path / "docs" / "TEMPLATE_ALIGNMENT.md"
    result = it.analyze_template(
        template, out, generator_path=generator, doc_path=doc, **kwargs
    )
    return out, doc, result


# --- Parte A: summary --------------------------------------------------


def test_template_summary(template_workbook, tmp_path):
    out, _doc, _result = _run(template_workbook, tmp_path)
    summary = json.loads((out / "template_summary.json").read_text(encoding="utf-8"))

    assert summary["sheet_names"] == ["NOMBRE", "FORM", "CONTROL"]
    by_name = {s["name"]: s for s in summary["sheets"]}
    assert by_name["FORM"]["protected"] is True
    assert by_name["CONTROL"]["protected"] is False
    assert by_name["FORM"]["formula_count"] == 1
    assert by_name["CONTROL"]["formula_count"] == 2
    assert summary["totals"]["formula_count"] == 4
    assert summary["vba"]["present"] is False
    assert len(summary["sha256"]) == 64


# --- Parte B: cell inventory ---------------------------------------


def test_template_cells(template_workbook, tmp_path):
    out, _doc, _result = _run(template_workbook, tmp_path)
    rows = {(r["sheet"], r["cell"]): r for r in _rows(out / "template_cells.csv")}

    assert rows[("FORM", "C15")]["has_formula"] == "True"
    assert rows[("FORM", "C15")]["formula"] == "=SUM(D15:F15)"
    assert rows[("FORM", "C15")]["text_label"] == ""

    assert rows[("FORM", "D15")]["locked"] == "False"
    assert rows[("FORM", "E15")]["fill_rgb"] == "FFFFFF00"
    assert rows[("FORM", "F15")]["has_validation"] == "True"
    assert rows[("FORM", "A3")]["is_merged"] == "True"
    assert rows[("FORM", "B1")]["kind"] == "label"
    assert rows[("FORM", "B1")]["text_label"] == "SECCIÓN A: CONTROLES"


# --- Parte C: input candidates -----------------------------------


def test_input_candidates(template_workbook, tmp_path):
    out, _doc, _result = _run(template_workbook, tmp_path)
    cells = {(r["sheet"], r["cell"]) for r in _rows(out / "input_candidates.csv")}

    assert ("FORM", "D15") in cells  # desbloqueada en hoja protegida
    assert ("FORM", "F15") in cells  # data validation
    assert ("NOMBRE", "B2") in cells  # desbloqueada en hoja protegida

    assert ("FORM", "C15") not in cells  # tiene fórmula
    assert ("FORM", "A3") not in cells  # merged
    assert ("CONTROL", "B1") not in cells  # desbloqueada pero hoja SIN proteger

    row = next(
        r for r in _rows(out / "input_candidates.csv")
        if (r["sheet"], r["cell"]) == ("FORM", "F15")
    )
    assert "has_data_validation" in row["candidate_reasons"]


# --- Parte D: dependencies --------------------------------------


def test_template_dependencies(template_workbook, tmp_path):
    out, _doc, result = _run(template_workbook, tmp_path)

    dep_rows = _rows(out / "template_formula_dependencies.csv")
    assert any(
        r["target_sheet"] == "CONTROL"
        and r["target_cell"] == "A1"
        and r["source_sheet"] == "FORM"
        and r["is_cross_sheet"] == "True"
        for r in dep_rows
    )

    edges = {
        (r["source_sheet"], r["target_sheet"]): int(r["reference_count"])
        for r in _rows(out / "template_sheet_dependencies.csv")
    }
    assert edges[("FORM", "CONTROL")] == 1
    assert edges[("NOMBRE", "CONTROL")] == 1
    assert result.unsupported_counts == {}


# --- Parte E: fingerprint --------------------------------------


def test_fingerprint(template_workbook, tmp_path):
    out, _doc, _result = _run(template_workbook, tmp_path)
    fp = json.loads((out / "template_fingerprint.json").read_text(encoding="utf-8"))

    assert len(fp["structural_sha256"]) == 64
    assert fp["structural"]["sheet_names_ordered"] == ["NOMBRE", "FORM", "CONTROL"]
    assert fp["structural"]["formula_count_by_sheet"]["CONTROL"] == 2
    assert "SECCION A CONTROLES" in {s["label"] for s in fp["structural"]["sentinel_labels"]}
    assert fp["full_file_sha256"] == _result_sha(out)


def _result_sha(out):
    return json.loads((out / "template_summary.json").read_text(encoding="utf-8"))["sha256"]


def test_fingerprint_changes_with_structure(template_workbook, generator_workbook, tmp_path):
    _out, _doc, r1 = _run(template_workbook, tmp_path)

    wb = it.load_workbook_safely(template_workbook)
    wb.create_sheet("EXTRA")
    wb.save(tmp_path / "template2.xlsx")

    _out2, _doc2, r2 = _run(tmp_path / "template2.xlsx", tmp_path)
    assert r1.fingerprint["structural_sha256"] != r2.fingerprint["structural_sha256"]


# --- Parte F: alignment --------------------------------------


def test_alignment_candidates(template_workbook, generator_workbook, tmp_path):
    out, _doc, result = _run(template_workbook, tmp_path, generator=generator_workbook)
    rows = _rows(out / "alignment_candidates.csv")

    by_label = {}
    for r in rows:
        by_label.setdefault(r["generator_label"], []).append(r)

    assert by_label["SECCIÓN A: CONTROLES"][0]["match_type"] == "exact_label"
    assert by_label["SECCIÓN A: CONTROLES"][0]["template_cell"] == "B1"
    assert float(by_label["SECCIÓN A: CONTROLES"][0]["confidence"]) == 1.0

    assert by_label["TOTAL"][0]["match_type"] == "exact_label"
    assert by_label["0601101"][0]["match_type"] == "exact_label"
    assert by_label["0601101"][0]["template_cell"] == "A20"

    # "ACTIVIDADES" aparece 2 veces en la plantilla -> ambiguous, 2 filas.
    act = by_label["ACTIVIDADES"]
    assert len(act) == 2
    assert all(r["match_type"] == "ambiguous" for r in act)
    assert {r["template_sheet"] for r in act} == {"FORM", "NOMBRE"}

    # "Total General" no matchea; "de" es demasiado corta.
    assert "Total General" not in by_label
    assert "de" not in by_label

    types = {a.match_type for a in result.alignment_candidates}
    assert types == {"exact_label", "ambiguous"}


def test_alignment_never_silently_picks_one(template_workbook, generator_workbook, tmp_path):
    _out, _doc, result = _run(template_workbook, tmp_path, generator=generator_workbook)
    ambiguous = [a for a in result.alignment_candidates if a.generator_label == "ACTIVIDADES"]
    assert len(ambiguous) == 2  # se conservan todos, no se elige


def test_detail_sheet_excluded_from_generator_labels(
    template_workbook, generator_workbook, tmp_path
):
    with pytest.raises(it.InventoryError):
        _run(
            template_workbook,
            tmp_path,
            generator=generator_workbook,
            generator_label_sheets=["REMASEP 01", "Atenciones - Detalles de citas"],
        )


# --- Parte G: doc ----------------------------------------------


def test_alignment_doc(template_workbook, generator_workbook, tmp_path):
    _out, doc, _result = _run(template_workbook, tmp_path, generator=generator_workbook)
    text = doc.read_text(encoding="utf-8")

    for heading in (
        "## 1. Estructura de la plantilla oficial",
        "## 2. Hojas con fórmulas",
        "## 3. Hojas que se alimentan entre sí",
        "## 4. Input candidates",
        "## 5. Distribución de rellenos y protección",
        "## 6. Candidatos de alineación por hoja",
        "## 7. Matches ambiguos",
        "## 8. Observaciones que requieren revisión humana",
    ):
        assert heading in text
    assert "```mermaid" in text
    assert "ACTIVIDADES" in text  # aparece en la tabla de ambiguos


# --- Macros / VBA -------------------------------------------


def test_vba_detected_from_zip_payload(template_workbook, tmp_path):
    payload = b"FAKE-VBA-PROJECT-BINARY" * 40
    injected = tmp_path / "with_vba.xlsx"
    with zipfile.ZipFile(template_workbook) as src, zipfile.ZipFile(
        injected, "w", zipfile.ZIP_DEFLATED
    ) as dst:
        for item in src.namelist():
            dst.writestr(item, src.read(item))
        dst.writestr(it._VBA_ENTRY, payload)

    _out, _doc, result = _run(injected, tmp_path)
    assert result.vba_present is True
    assert result.vba_payload_bytes == len(payload)
    assert result.vba_payload_sha256 == it._sha256_bytes(payload)


# --- Privacidad / robustez ---------------------------------


def test_no_generator_patient_values_in_outputs(
    template_workbook, generator_workbook, tmp_path
):
    out, doc, _result = _run(template_workbook, tmp_path, generator=generator_workbook)
    targets = list(out.glob("*.csv")) + list(out.glob("*.json")) + [out / "README.md", doc]
    for path in targets:
        assert SECRET not in path.read_text(encoding="utf-8"), path.name


def test_inputs_not_modified(template_workbook, generator_workbook, tmp_path):
    t_before = template_workbook.read_bytes()
    g_before = generator_workbook.read_bytes()
    _run(template_workbook, tmp_path, generator=generator_workbook)
    assert template_workbook.read_bytes() == t_before
    assert generator_workbook.read_bytes() == g_before


def test_cli_main_ok(template_workbook, generator_workbook, tmp_path, capsys):
    code = it.main(
        [
            str(template_workbook),
            "--generator",
            str(generator_workbook),
            "--output",
            str(tmp_path / "ti"),
            "--doc",
            str(tmp_path / "doc.md"),
        ]
    )
    assert code == 0
    out = capsys.readouterr().out
    assert "VBA presente: False" in out
    assert "alignment candidates:" in out


def test_cli_missing_file(tmp_path, capsys):
    code = it.main([str(tmp_path / "nope.xlsm"), "--output", str(tmp_path / "ti")])
    assert code == 2
    assert "no existe" in capsys.readouterr().err.lower()
