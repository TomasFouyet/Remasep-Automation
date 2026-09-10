"""Tests de la capa de write readiness (Sprint 3.6). Sintéticos, sin data/local."""

from __future__ import annotations

from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.styles import Protection

from remasep.services import template_alignment as ta
from remasep.services import writable_target_mapping as wtm

POLICY = wtm.load_write_policy()


def _align_row(**kw) -> ta.AlignmentRow:
    base = {
        "source_metric_id": "LEGACY::REMASEP_OD::C6",
        "source_form": "REMASEP_OD",
        "source_kind": "BASE_AGGREGATION",
        "source_readiness": "AUTO_READY",
        "target_sheet": "REMASEP_OD",
        "target_cell": "H99",
        "target_kind": ta.DIRECT_INPUT_TARGET,
        "match_status": ta.EXACT_SEMANTIC_MATCH,
        "match_score": 1.0,
        "source_semantic_signature": "REMASEP_OD :: SIG-SRC",
        "target_semantic_signature": "REMASEP_OD :: SIG-TGT",
        "sex_match": "MATCH", "age_match": "MATCH", "procedure_match": "NOT_APPLICABLE",
        "row_match": "MATCH", "column_match": "MATCH", "section_match": "MATCH",
        "target_alignment_role": ta.ROLE_INPUT_TARGET,
        "target_locked": False,
        "target_source_expectation": ta.EXP_MEDINET,
        "needs_human_review": False,
    }
    base.update(kw)
    return ta.AlignmentRow(**base)


def _report(rows, *, ambiguous=(), excluded=(), unmatched=(), evidence=(),
            eligible=()):
    return SimpleNamespace(
        rows=list(rows), ambiguous=list(ambiguous), excluded_sources=list(excluded),
        unmatched_sources=list(unmatched), evidence=list(evidence),
        eligible_source_ids=tuple(eligible),
    )


_TFP = {"structural_template_fingerprint_id": "stf:deadbeef00000000"}


def _assess(row, **kw):
    return wtm.assess_alignment_row(
        row, source_readiness=kw.get("source_readiness", row.source_readiness),
        target_value=kw.get("target_value"), policy=POLICY,
        locked_override=kw.get("locked_override", wtm._UNSET),
    )[:2]


# --- EXACT / STRONG ------------------------------------------------


def test_exact_auto_ready_unlocked_input_is_write_ready():
    assert _assess(_align_row()) == (wtm.WRITE_READY, wtm.R_EXACT_ALL_MATCH)


def test_strong_all_required_match_is_write_ready():
    r = _align_row(match_status=ta.STRONG_MATCH, row_match="MATCH", column_match="MATCH",
                   section_match="MATCH", sex_match="MATCH", age_match="MATCH")
    assert _assess(r) == (wtm.WRITE_READY, wtm.R_STRONG_REQUIRED_MATCH)


def test_strong_missing_required_dimension_is_review_required():
    r = _align_row(match_status=ta.STRONG_MATCH, row_match="PARTIAL")
    status, reason = _assess(r)
    assert status == wtm.WRITE_REVIEW_REQUIRED
    assert reason == wtm.R_STRONG_MISSING_REQUIRED


def test_strong_sex_missing_is_review_required():
    r = _align_row(match_status=ta.STRONG_MATCH, sex_match="MISSING")
    assert _assess(r)[0] == wtm.WRITE_REVIEW_REQUIRED


def test_strong_dimension_mismatch_is_review_required():
    r = _align_row(match_status=ta.STRONG_MATCH, age_match="MISMATCH")
    status, reason = _assess(r)
    assert status == wtm.WRITE_REVIEW_REQUIRED
    assert reason == wtm.R_STRONG_DIMENSION_MISMATCH


def test_b2_strong_uses_procedure_code_required():
    r = _align_row(source_form="B2_ANEXO", source_metric_id="LEGACY::B2_ANEXO::D5",
                   match_status=ta.STRONG_MATCH, target_sheet="B2 ANEXO",
                   row_match="MATCH", procedure_match="MATCH",
                   sex_match="NOT_APPLICABLE", age_match="NOT_APPLICABLE",
                   column_match="NOT_APPLICABLE", section_match="MISSING")
    assert _assess(r) == (wtm.WRITE_READY, wtm.R_STRONG_REQUIRED_MATCH)


# --- AMBIGUOUS / CONFLICT --------------------------------------


def test_ambiguous_is_write_blocked():
    assert _assess(_align_row(match_status=ta.AMBIGUOUS)) == (
        wtm.WRITE_BLOCKED, wtm.R_MATCH_AMBIGUOUS)


def test_conflict_is_write_blocked():
    assert _assess(_align_row(match_status=ta.CONFLICT)) == (
        wtm.WRITE_BLOCKED, wtm.R_DIMENSION_CONFLICT)


# --- target guards -------------------------------------------


def test_formula_target_is_not_writable():
    r = _align_row(target_kind=ta.FORMULA_TARGET, target_alignment_role=ta.ROLE_DERIVED_TARGET)
    assert _assess(r) == (wtm.NOT_WRITABLE, wtm.R_TARGET_IS_FORMULA)


def test_structural_target_is_not_writable():
    r = _align_row(target_kind=ta.STRUCTURAL, target_alignment_role=ta.ROLE_STRUCTURAL)
    assert _assess(r)[0] == wtm.NOT_WRITABLE


def test_locked_target_is_write_blocked():
    assert _assess(_align_row(target_locked=True)) == (wtm.WRITE_BLOCKED, wtm.R_TARGET_LOCKED)
    assert _assess(_align_row(target_locked=None)) == (wtm.WRITE_BLOCKED, wtm.R_TARGET_LOCKED)


def test_locked_revalidation_override_wins():
    r = _align_row(target_locked=False)  # el alignment cree unlocked
    assert _assess(r, locked_override=True) == (wtm.WRITE_BLOCKED, wtm.R_TARGET_LOCKED)


def test_non_medinet_target_is_write_blocked():
    for exp in (ta.EXP_EGRESOS, ta.EXP_RESOURCE_CALCULATION, ta.EXP_SURGICAL_TABLE,
                ta.EXP_CONTROL_METADATA):
        assert _assess(_align_row(target_source_expectation=exp)) == (
            wtm.WRITE_BLOCKED, wtm.R_TARGET_SOURCE_NOT_MEDINET)


def test_unknown_target_expectation_is_review_required():
    r = _align_row(target_source_expectation=ta.EXP_UNKNOWN)
    assert _assess(r) == (wtm.WRITE_REVIEW_REQUIRED, wtm.R_TARGET_SOURCE_UNKNOWN)


# --- source guards ------------------------------------------


def test_review_required_source_is_not_writable():
    assert _assess(_align_row(), source_readiness="REVIEW_REQUIRED") == (
        wtm.NOT_WRITABLE, wtm.R_SOURCE_REVIEW_REQUIRED)


# --- expected value type -----------------------------------


def test_expected_value_type_integer_for_aggregate():
    evt, ok = wtm.expected_value_type_for("BASE_AGGREGATION", None, POLICY)
    assert evt == wtm.EVT_INTEGER_COUNT and ok is True
    assert wtm.expected_value_type_for("BASE_AGGREGATION", 5, POLICY) == (
        wtm.EVT_INTEGER_COUNT, True)


def test_text_value_in_count_target_is_incompatible():
    r = _align_row()
    status, reason = _assess(r, target_value="texto inesperado")
    assert status == wtm.WRITE_REVIEW_REQUIRED
    assert reason == wtm.R_EXPECTED_VALUE_TYPE_INCOMPATIBLE


# --- build_write_mapping: roll-up / excluded --------------


def test_rollup_source_is_not_writable():
    excl = SimpleNamespace(metric_id="LEGACY::REMASEP_OD::E6", form="REMASEP_OD",
                           readiness_status="NOT_APPLICABLE", reason="ROLLUP_NOT_INPUT",
                           detail="")
    rep = _report([], excluded=[excl])
    wr = wtm.build_write_mapping(rep, _TFP, POLICY)
    w = wr.readiness[0]
    assert w.write_status == wtm.NOT_WRITABLE
    assert w.write_reason == wtm.R_ROLLUP_COMPUTED_BY_TEMPLATE
    assert wr.instructions == []


def test_blocked_conflict_source_is_write_blocked():
    excl = SimpleNamespace(metric_id="LEGACY::REMASEP_01::AB84", form="REMASEP_01",
                           readiness_status="BLOCKED_CONFLICT", reason="BLOCKED_CONFLICT",
                           detail="")
    wr = wtm.build_write_mapping(_report([], excluded=[excl]), _TFP, POLICY)
    assert wr.readiness[0].write_status == wtm.WRITE_BLOCKED
    assert wr.readiness[0].write_reason == wtm.R_SOURCE_BLOCKED_CONFLICT


# --- collisions -------------------------------------------


def test_multiple_sources_same_target_collision_blocks_all():
    a = _align_row(source_metric_id="M1", source_semantic_signature="S1", target_cell="H10")
    b = _align_row(source_metric_id="M2", source_semantic_signature="S2", target_cell="H10")
    wr = wtm.build_write_mapping(_report([a, b]), _TFP, POLICY)
    assert {w.write_status for w in wr.readiness} == {wtm.WRITE_BLOCKED}
    assert all(w.write_reason == wtm.R_WRITE_COLLISION for w in wr.readiness)
    assert any(c.collision_type == wtm.COLL_MULTI_SOURCE_SAME_TARGET for c in wr.collisions)
    assert wr.instructions == []


def test_same_source_multiple_targets_collision_blocks():
    a = _align_row(source_metric_id="M1", target_cell="H10")
    b = _align_row(source_metric_id="M1", target_cell="H20")
    wr = wtm.build_write_mapping(_report([a, b]), _TFP, POLICY)
    assert any(c.collision_type == wtm.COLL_SAME_SOURCE_MULTI_TARGET for c in wr.collisions)
    assert {w.write_status for w in wr.readiness} == {wtm.WRITE_BLOCKED}


def test_no_last_write_wins():
    a = _align_row(source_metric_id="M1", source_semantic_signature="S1", target_cell="H10")
    b = _align_row(source_metric_id="M2", source_semantic_signature="S2", target_cell="H10")
    wr = wtm.build_write_mapping(_report([a, b]), _TFP, POLICY)
    # NINGUNO gana: ambos bloqueados, sin instrucción
    assert wr.instructions == []


# --- instruction identity / template fingerprint ---------


def test_instruction_id_is_stable_and_semantic_not_coordinate():
    i1 = wtm.instruction_id("SRC", "TGT", "stf:abc")
    i2 = wtm.instruction_id("SRC", "TGT", "stf:abc")
    assert i1 == i2 and i1.startswith("wi:")
    # cambia la coordenada pero NO la firma semántica -> mismo id
    a = _align_row(target_cell="H99")
    b = _align_row(target_cell="ZZ1")   # otra coordenada, misma firma
    wa = wtm.build_write_mapping(_report([a]), _TFP, POLICY)
    wb = wtm.build_write_mapping(_report([b]), _TFP, POLICY)
    assert wa.instructions[0].instruction_id == wb.instructions[0].instruction_id


def test_instruction_id_changes_with_template_fingerprint():
    a = _align_row()
    w1 = wtm.build_write_mapping(_report([a]), {"structural_template_fingerprint_id": "stf:aaa"}, POLICY)
    w2 = wtm.build_write_mapping(_report([a]), {"structural_template_fingerprint_id": "stf:bbb"}, POLICY)
    assert w1.instructions[0].instruction_id != w2.instructions[0].instruction_id


def test_structural_fingerprint_stable_under_resave(tmp_path):
    def _wb():
        wb = Workbook()
        ws = wb.active
        ws.title = "B2 ANEXO"
        ws.protection.sheet = True
        ws["A1"] = "CÓDIGOS"
        ws["A2"] = "0301002"
        ws["C2"].protection = Protection(locked=False)
        ws["C3"] = "=SUM(C2:C2)"
        return wb

    from openpyxl import load_workbook
    p1, p2 = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    _wb().save(p1)
    _wb().save(p2)  # "re-guardado" idéntico estructuralmente
    fp1 = wtm.structural_template_fingerprint(load_workbook(p1, data_only=False),
                                              file_sha256="hash-1")
    fp2 = wtm.structural_template_fingerprint(load_workbook(p2, data_only=False),
                                              file_sha256="hash-2-different")
    assert fp1["structural_sha256"] == fp2["structural_sha256"]
    assert fp1["file_sha256"] != fp2["file_sha256"]

    # cambio ESTRUCTURAL (celda antes locked ahora unlocked) -> distinto
    wb3 = _wb()
    wb3["B2 ANEXO"]["A2"].protection = Protection(locked=False)
    p3 = tmp_path / "c.xlsx"
    wb3.save(p3)
    fp3 = wtm.structural_template_fingerprint(load_workbook(p3, data_only=False))
    assert fp3["structural_sha256"] != fp1["structural_sha256"]


# --- zero / estado / period contract --------------------


def test_zero_write_policy_unresolved_does_not_block_write_ready():
    wr = wtm.build_write_mapping(_report([_align_row()]), _TFP, POLICY)
    assert wr.zero_write_policy == "UNRESOLVED"
    assert wr.readiness[0].write_status == wtm.WRITE_READY  # el cero es cosa del writer


def test_no_estado_filter_applied():
    # el módulo no aplica ningún filtro por ESTADO
    assert not hasattr(wtm, "apply_estado_filter")
    wr = wtm.build_write_mapping(_report([_align_row()]), _TFP, POLICY)
    assert wr.estado_filter_status == "PENDING_FUNCTIONAL_CONFIRMATION"


def test_metric_value_contract_has_no_pii_fields():
    import dataclasses
    fields = {f.name for f in dataclasses.fields(wtm.MetricValue)}
    assert fields == {"source_metric_id", "value", "period", "producer_version"}
    assert "run" not in {f.lower() for f in fields}
    assert "nombre" not in {f.lower() for f in fields}


# --- review clusters -----------------------------------


def test_review_clusters_group_by_cause():
    rows = [
        _align_row(source_metric_id=f"M{n}", match_status=ta.STRONG_MATCH,
                   row_match="PARTIAL", source_semantic_signature=f"S{n}",
                   target_cell=f"H{10 + n}", target_semantic_signature=f"T{n}")
        for n in range(4)
    ]
    wr = wtm.build_write_mapping(_report(rows), _TFP, POLICY)
    clusters = wtm.build_review_clusters(wr.readiness)
    assert len(clusters) == 1
    assert clusters[0].metric_count == 4
    assert "ROW_PATH" in clusters[0].missing_evidence


def test_missing_config_raises():
    with pytest.raises(wtm.WritableTargetMappingError):
        wtm.load_write_policy("/nonexistent/xyz")
