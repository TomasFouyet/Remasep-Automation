"""Tests de la capa de layout semántico (Sprint 3.1).

Hojas openpyxl construidas en memoria; sin data/local.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from remasep.services.semantic_layout import (
    CONTEXT_AMBIGUOUS,
    CONTEXT_COMPLETE,
    CONTEXT_NO_CONTEXT,
    CONTEXT_PARTIAL,
    SheetLayout,
    age_bounds_from_formula,
    age_intervals_overlap,
    compare_label_and_formula_age,
    normalize_semantic_label,
    parse_age_label,
    text_criteria_from_formula,
)

_BOLD = Font(bold=True)


def _sheet(title="REMASEP_OD"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    return ws


# --- normalización -----------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("20 A 24 AÑOS", "20 A 24 ANOS"),
        ("  Menor de \n10 años ", "MENOR DE 10 ANOS"),
        ("5010009 - VIDRIO IONÓMERO", "5010009 - VIDRIO IONOMERO"),
        (None, ""),
        ("MUJERES", "MUJERES"),
    ],
)
def test_normalize_semantic_label(raw, expected):
    assert normalize_semantic_label(raw) == expected


def test_raw_is_never_lost():
    ws = _sheet()
    ws["B1"] = "SECCIÓN A: PRUEBAS"
    ws["B1"].font = _BOLD
    ws.merge_cells("C2:D2")
    ws["C2"] = "MUJERES"
    ws["C3"] = "20 A 24 AÑOS"
    ws["A5"] = "CÓD-1"
    ws["B5"] = "Actividad Ñandú"
    ws["C5"] = 3
    ctx = SheetLayout(ws).cell_context("C5")
    assert ctx.row_labels_raw == ("CÓD-1", "Actividad Ñandú")
    assert ctx.row_labels_norm == ("COD-1", "ACTIVIDAD NANDU")
    assert "Ñandú" in ctx.row_labels_raw[1]  # raw intacto


# --- parse_age_label -------------------------------------------------


@pytest.mark.parametrize(
    ("label", "bounds"),
    [
        ("20 A 24 AÑOS", (20, 24)),
        ("20-24 años", (20, 24)),
        ("20 - 24 años ", (20, 24)),
        ("8-9 años", (8, 9)),
        ("2 años", (2, 2)),
        ("75 y más años", (75, None)),
        ("Menos de 1 año - 1 año", (0, 1)),
        ("Menos de 2 años", (None, 1)),
        ("MUJERES", None),
        ("EMBARAZADAS", None),
        ("", None),
    ],
)
def test_parse_age_label(label, bounds):
    assert parse_age_label(label) == bounds


# --- evidencia de fórmula ------------------------------------------


def test_age_bounds_from_formula_band():
    f = (
        "=COUNTIFS('Atenciones - Detalles de citas'!$AE:$AE,\"*x*\","
        "'Atenciones - Detalles de citas'!$AF:$AF,\">=20\","
        "'Atenciones - Detalles de citas'!$AF:$AF,\"<=24\")"
    )
    assert age_bounds_from_formula(f) == (20, 24)


def test_age_bounds_from_formula_open_and_strict():
    assert age_bounds_from_formula("x $AF:$AF,\"<2\"") == (None, 1)
    assert age_bounds_from_formula("x $AF:$AF,\">2\"") == (3, None)
    assert age_bounds_from_formula("x $AF:$AF,\"=2\"") == (2, 2)
    assert age_bounds_from_formula("=SUM(A1:A5)") is None


def test_text_criteria_from_formula():
    f = "=COUNTIF(D!$AE:$AE,\"*VIDRIO IONÓMERO*\"&\"*Hombre*\")"
    assert text_criteria_from_formula(f) == ("*VIDRIO IONÓMERO*", "*Hombre*")
    assert text_criteria_from_formula("=SUM(F1:F9)") == ()


# --- consistencia rótulo / fórmula ----------------------------------


def test_consistency_consistent():
    assert compare_label_and_formula_age((20, 24), (20, 24)) == "CONSISTENT"
    assert compare_label_and_formula_age((0, 1), (None, 1)) == "CONSISTENT"


def test_consistency_conflict():
    assert compare_label_and_formula_age((20, 24), (25, 29)) == "CONFLICT"


def test_consistency_no_comparable():
    assert compare_label_and_formula_age(None, (20, 24)) == "NO_COMPARABLE"
    assert compare_label_and_formula_age((20, 24), None) == "NO_COMPARABLE"


def test_age_intervals_overlap():
    assert age_intervals_overlap((20, 24), (24, 30)) is True
    assert age_intervals_overlap((20, 24), (25, 30)) is False
    assert age_intervals_overlap((75, None), (80, 84)) is True


# --- SheetLayout: encabezados -------------------------------------


def test_simple_header_complete():
    ws = _sheet()
    ws["B1"] = "SECCIÓN A: ODONTOLOGÍA"
    ws["B1"].font = _BOLD
    ws["C2"] = "HOMBRES"
    ws["A4"] = "COD-1"
    ws["B4"] = "Prestación uno"
    ws["C4"] = 5
    ctx = SheetLayout(ws).cell_context("C4")
    assert ctx.context_status == CONTEXT_COMPLETE
    assert ctx.section_labels_raw == ("SECCIÓN A: ODONTOLOGÍA",)
    assert ctx.row_labels_raw == ("COD-1", "Prestación uno")
    assert ctx.column_labels_raw == ("HOMBRES",)


def test_merged_header_and_two_levels():
    ws = _sheet()
    ws.merge_cells("C1:F1")
    ws["C1"] = "POR GRUPO DE EDAD"
    ws.merge_cells("C2:D2")
    ws["C2"] = "20 A 24 AÑOS"
    ws["C3"] = "Hombres"
    ws["D3"] = "Mujeres"
    ws["B5"] = "Actividad"
    ws["C5"] = 1
    ws["D5"] = 2
    layout = SheetLayout(ws)
    ctx_c = layout.cell_context("C5")
    ctx_d = layout.cell_context("D5")
    assert ctx_c.column_labels_raw == ("POR GRUPO DE EDAD", "20 A 24 AÑOS", "Hombres")
    assert ctx_d.column_labels_raw == ("POR GRUPO DE EDAD", "20 A 24 AÑOS", "Mujeres")


def test_three_column_levels_capped_at_three():
    ws = _sheet()
    ws["C1"] = "N1"
    ws["C2"] = "N2"
    ws["C3"] = "N3"
    ws["C4"] = "N4"
    ws["B6"] = "fila"
    ws["C6"] = 9
    ctx = SheetLayout(ws).cell_context("C6")
    assert ctx.column_labels_raw == ("N1", "N2", "N3")  # se conservan los 3 más externos


def test_nested_merged_sections():
    ws = _sheet("B2 ANEXO")  # sin la palabra "SECCIÓN"
    ws.merge_cells("B1:C1")
    ws["B1"] = "OTORRINOLARINGOLOGÍA"
    ws["B1"].font = _BOLD
    ws.merge_cells("B2:C2")
    ws["B2"] = "PROCEDIMIENTOS DIAGNÓSTICOS"
    ws["B2"].font = _BOLD
    ws["B4"] = "0601101"
    ws["C4"] = "Electrogustometría"
    ws["D4"] = 3
    ctx = SheetLayout(ws).cell_context("D4")
    assert ctx.section_labels_raw == ("OTORRINOLARINGOLOGÍA", "PROCEDIMIENTOS DIAGNÓSTICOS")


def test_bold_header_acts_as_section_when_no_keyword():
    ws = _sheet("B2 ANEXO")
    ws.merge_cells("B1:C1")
    ws["B1"] = "A. KINESIOLOGÍA"
    ws["B1"].font = _BOLD
    ws["B3"] = "0601105"
    ws["C3"] = "Atención Kinesiológica"
    ws["D3"] = 7
    ctx = SheetLayout(ws).cell_context("D3")
    assert ctx.section_labels_raw == ("A. KINESIOLOGÍA",)
    assert ctx.context_status == CONTEXT_PARTIAL  # sin encabezado de columna


def test_group_header_is_row_level_when_section_keyword_present():
    ws = _sheet()
    ws["B1"] = "SECCIÓN A: ODONTOLOGÍA"
    ws["B1"].font = _BOLD
    ws["B3"] = "   A.4. ACTIVIDADES GENERALES"
    ws["B3"].font = _BOLD
    ws["C3"] = 0
    ws["C4"] = "Hombres"
    ws["A6"] = "COD-9"
    ws["B6"] = "Obturación"
    ws["C6"] = 4
    ctx = SheetLayout(ws).cell_context("C6")
    assert ctx.section_labels_raw == ("SECCIÓN A: ODONTOLOGÍA",)
    assert ctx.row_labels_raw == ("A.4. ACTIVIDADES GENERALES", "COD-9", "Obturación")


def test_merged_row_label():
    ws = _sheet()
    ws["C1"] = "Hombres"
    ws.merge_cells("B4:C4")
    ws["B4"] = "Pediatría"
    ws["D4"] = 5
    ctx = SheetLayout(ws).cell_context("D4")
    assert ctx.row_labels_raw == ("Pediatría",)
    assert ctx.column_labels_raw == ()  # C1 no está sobre la columna D


def test_partial_row_only():
    ws = _sheet("B2 ANEXO")
    ws["B5"] = "0601105"
    ws["C5"] = "Atención"
    ws["D5"] = 9
    ctx = SheetLayout(ws).cell_context("D5")
    assert ctx.context_status == CONTEXT_PARTIAL
    assert ctx.row_labels_raw == ("0601105", "Atención")
    assert ctx.column_labels_raw == ()


def test_no_context():
    ws = _sheet()
    ws["B2"] = 7  # sin nada a la izquierda ni encima
    ctx = SheetLayout(ws).cell_context("B2")
    assert ctx.context_status == CONTEXT_NO_CONTEXT
    assert ctx.semantic_signature() == "REMASEP_OD :: <NO_CONTEXT>"


def test_ambiguous_when_header_band_gap():
    ws = _sheet()
    ws["C2"] = "TOP"
    ws["D3"] = "MID"          # sólo en la columna vecina
    ws["C4"] = "LOW"
    ws["C6"] = 9
    ctx = SheetLayout(ws).cell_context("C6")
    assert ctx.context_status == CONTEXT_AMBIGUOUS
    assert ctx.notes and "columna vecina" in ctx.notes[0]


def test_gap_rows_between_headers_and_metric_are_skipped():
    ws = _sheet()
    ws["B1"] = "SECCIÓN A: X"
    ws["B1"].font = _BOLD
    ws["C2"] = "Hombres"
    for r in range(3, 20):  # muchas filas de datos entre encabezado y métrica
        ws[f"A{r}"] = f"COD-{r}"
        ws[f"C{r}"] = r
    ctx = SheetLayout(ws).cell_context("C19")
    assert ctx.column_labels_raw == ("Hombres",)
    assert ctx.row_labels_raw == ("COD-19",)


def test_semantic_signature_is_stable_and_normalized():
    ws = _sheet()
    ws["B1"] = "SECCIÓN A: ODONTOLOGÍA"
    ws["B1"].font = _BOLD
    ws["C2"] = "20 A 24 AÑOS"
    ws["B4"] = "Pediatría"
    ws["C4"] = 1
    sig1 = SheetLayout(ws).cell_context("C4").semantic_signature()
    sig2 = SheetLayout(ws).cell_context("C4").semantic_signature()
    assert sig1 == sig2
    assert sig1 == (
        "REMASEP_OD :: SECCION A: ODONTOLOGIA :: PEDIATRIA :: 20 A 24 ANOS"
    )
