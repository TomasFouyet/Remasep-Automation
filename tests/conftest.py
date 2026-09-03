"""Configuración compartida de tests.

Fuerza Qt en modo offscreen para que la maqueta de UI pueda testearse sin
display físico, y expone un constructor de workbooks Medinet sintéticos.
"""

import os

import pytest
from openpyxl import Workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

# Encabezados equivalentes a los del workbook real (con tildes y espacios).
MEDINET_HEADERS = [
    "DIA CITA",
    "FECHA NACIMIENTO",
    "SEXO",
    "SUCURSAL",
    "ESPECIALIDAD",
    "TIPO DE CITA",
    "PRESTACIÓN",
    "ESTADO",
    "MODALIDAD",
    "PRESTACIÓN REALIZADA",
]


@pytest.fixture(scope="session")
def qapp():
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def make_medinet(tmp_path):
    """Devuelve un builder de archivos Medinet sintéticos (.xlsx).

    ``rows`` es una lista de listas alineadas a ``headers`` (por defecto
    :data:`MEDINET_HEADERS`). No usa data/local.
    """

    def _build(
        rows,
        *,
        headers=None,
        sheet="Atenciones",
        filename="medinet.xlsx",
        decoy_sheets=None,
        extra_columns=None,
    ):
        wb = Workbook()
        first = wb.active
        if decoy_sheets:
            (name0, rows0), *rest = decoy_sheets
            first.title = name0
            for decoy_row in rows0:
                first.append(list(decoy_row))
            for name, drows in rest:
                extra = wb.create_sheet(name)
                for decoy_row in drows:
                    extra.append(list(decoy_row))
            data_ws = wb.create_sheet(sheet)
        else:
            first.title = sheet
            data_ws = first

        # extra_columns: [(header, value_para_todas_las_filas), ...] — p.ej. una
        # columna derivada "arrastrada" que evita que openpyxl recorte las filas
        # semánticamente vacías al final.
        extra_columns = list(extra_columns or [])
        data_ws.append(list(headers or MEDINET_HEADERS) + [name for name, _ in extra_columns])
        tail = [value for _, value in extra_columns]
        for row in rows:
            data_ws.append(list(row) + tail)
        path = tmp_path / filename
        wb.save(path)
        return path

    return _build
