"""Tests CLI del análisis de delta + provenance (Sprint 3.4). Pair sintético."""

from __future__ import annotations

import csv
import re
import shutil
import zipfile

import analyze_final_remasep_delta as af
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from remasep.services import source_provenance as sp
from remasep.services import workbook_delta as wd

_BOLD = Font(bold=True)
_SECRET = "PACIENTE-SECRETO-XYZ"


def _inject_cache(path, cache_by_sheet):
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        blobs = {n: zf.read(n) for n in names}
    order = sorted(
        (n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)),
        key=lambda n: int(re.search(r"(\d+)", n).group(1)),
    )
    sheet_files = dict(zip(cache_by_sheet.keys(), order, strict=False))
    for sheet, cache in cache_by_sheet.items():
        xml = blobs[sheet_files[sheet]].decode("utf-8")
        for ref, value in cache.items():
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                repl = rf'<c r="{ref}"\g<a>><f>\g<f></f><v>{value}</v></c>'
            else:
                esc = str(value).replace("&", "&amp;").replace("<", "&lt;")
                repl = rf'<c r="{ref}"\g<a> t="str"><f>\g<f></f><v>{esc}</v></c>'
            pat = re.compile(
                rf'<c r="{ref}"(?P<a>[^>]*)><f>(?P<f>[^<]*)</f>(?:<v\s*/>|<v></v>)</c>'
            )
            xml, n = pat.subn(repl, xml)
            assert n == 1, f"no se pudo cachear {sheet}!{ref}"
        blobs[sheet_files[sheet]] = xml.encode("utf-8")
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in names:
            dst.writestr(name, blobs[name])
    shutil.move(tmp, path)


def _build(wb: Workbook, phase: str) -> None:
    after = phase == "after"

    r01 = wb.active
    r01.title = "REMASEP 01"
    r01["A124"] = "SECCIÓN D: CAPACIDAD INSTALADA Y UTILIZACIÓN DE QUIRÓFANOS"
    r01["A124"].font = _BOLD
    r01["B125"] = "NÚMERO DE QUIRÓFANOS EN DOTACIÓN"
    r01["C125"] = "Horas Mensuales Ocupadas de Tabla Quirúrgica"
    r01["A129"] = "De Cirugía electiva"
    r01["B129"] = None if not after else 1          # RESOURCE_CALCULATION direct
    r01["C129"] = None if not after else 5          # SURGICAL_TABLE direct
    r01["B128"] = "=SUM(B129:B129)"                 # RESOURCE_CALCULATION derived
    r01["A134"] = "SECCIÓN E: CAUSAS DE SUSPENSIÓN DE CIRUGÍAS ELECTIVAS"
    r01["A134"].font = _BOLD
    r01["A137"] = "Paciente"
    r01["C137"] = None if not after else 4          # UNKNOWN_PENDING direct
    r01["A307"] = "=SUM(B129:C129,C137:C137)"       # MIXED_DERIVED

    b1 = wb.create_sheet("REMASEP B1")
    b1["A10"] = "SECCIÓN A: RESUMEN DE INTERVENCIONES QUIRÚRGICAS POR TIPO Y GRUPO DE EDAD"
    b1["A10"].font = _BOLD
    b1["A11"] = "INTERVENCIONES"
    b1["C11"] = "POR GRUPO DE EDAD (AÑOS)"
    b1["F11"] = "POR SEXO"
    b1["C12"] = "TOTAL"
    b1["D12"] = "Menor de 10 años"
    b1["E12"] = "10 y más años"
    b1["F12"] = "Hombres"
    b1["G12"] = "Mujeres"
    b1["A13"] = "Electivas"
    b1["B13"] = "Mayor Ambulatorias"
    b1["D13"] = None if not after else 2           # EGRESOS direct (edad)
    b1["E13"] = None if not after else 1
    b1["F13"] = None if not after else 2           # EGRESOS direct (sexo)
    b1["G13"] = None if not after else 1
    b1["C13"] = "=SUM(D13:E13)"                    # EGRESOS derived  -> 3
    b1["A25"] = "SECCIÓN B: INTERVENCIONES QUIRÚRGICAS"
    b1["A25"].font = _BOLD
    b1["A29"] = "III"
    b1["B29"] = "Cirugía Otorrinológica"
    b1["C29"] = "='B2 ANEXO'!C100"                 # cross-sheet -> EGRESOS derived -> 2
    b1["A46"] = "TOTAL DE INTERVENCIONES QUIRÚRGICAS"
    b1["C46"] = "=SUM(C27:C45)"                    # -> 2  (32-vs-31 style: 3 vs 2)

    b2 = wb.create_sheet("B2 ANEXO")
    b2["B98"] = "INTERVENCIONES QUIRÚRGICAS"
    b2["B98"].font = _BOLD
    b2["A99"] = 1302008
    b2["B99"] = "Rinoplastía y/o septoplastía"
    b2["C99"] = None if not after else 2           # EGRESOS direct (procedure code)
    b2["C100"] = "=SUM(C99:C99)"                   # EGRESOS derived -> 2

    ctrl = wb.create_sheet("CONTROL")
    ctrl["D5"] = "REMSAS"
    ctrl["E5"] = "Nº ERRORES"
    ctrl["F5"] = "REMSAS SIN DATOS"
    ctrl["G5"] = "Causal (Indicar)"
    ctrl["D7"] = "REMASEP 01"
    ctrl["E7"] = 0
    ctrl["F7"] = "OK"
    ctrl["D8"] = "URGENCIAS"
    ctrl["E8"] = 0
    ctrl["F8"] = "SIN DATOS"
    ctrl["G8"] = None if not after else "No tenemos servicio de Urgencia"  # CONTROL_METADATA

    # una fila con un identificador tipo PII para el test de privacidad: NO debe
    # aparecer en ningún artefacto (no cambia entre versiones -> no es delta).
    r01["Z400"] = _SECRET


_CACHE_BEFORE = {
    "REMASEP 01": {"B128": 0, "A307": 0},
    "REMASEP B1": {"C13": 0, "C29": 0, "C46": 0},
    "B2 ANEXO": {"C100": 0},
}
_CACHE_AFTER = {
    "REMASEP 01": {"B128": 1, "A307": 10},
    "REMASEP B1": {"C13": 3, "C29": 2, "C46": 2},
    "B2 ANEXO": {"C100": 2},
}


@pytest.fixture
def pair_paths(tmp_path):
    wb_b, wb_a = Workbook(), Workbook()
    _build(wb_b, "before")
    _build(wb_a, "after")
    pb, pa = tmp_path / "before.xlsx", tmp_path / "after.xlsx"
    wb_b.save(pb)
    wb_a.save(pa)
    _inject_cache(pb, _CACHE_BEFORE)
    _inject_cache(pa, _CACHE_AFTER)
    return pb, pa


@pytest.fixture
def result(pair_paths, tmp_path):
    pb, pa = pair_paths
    return af.run(str(pb), str(pa), str(tmp_path / "out")), tmp_path / "out"


# ---------------------------------------------------------------------------


def test_regression_shape(result):
    res, _ = result
    s = res["summary"]
    assert s["formula_expression_changes"] == 0
    assert s["direct_input_changes"] + s["formula_result_changes"] == s["total_changed_cells"]
    assert s["source_hash_verified_unchanged"] is True
    assert s["candidate_golden_status"] == "CANDIDATE_PENDING_SOURCE_COMPLETENESS"
    assert s["reference_before"]["status"] == "INCOMPLETE_PRE_SURGERY"
    assert s["reference_after"]["status"] == "FINAL_PER_CLIENT"
    assert s["reference_after"]["approval_status"] == "UNCONFIRMED"


def test_direct_vs_formula_classification(result):
    res, _ = result
    by = {d.key: d for d in res["deltas"]}
    assert by[("REMASEP B1", "D13")].change_kind == wd.DIRECT_INPUT_CHANGE
    assert by[("REMASEP B1", "C13")].change_kind == wd.FORMULA_RESULT_CHANGE
    assert by[("B2 ANEXO", "C99")].change_kind == wd.DIRECT_INPUT_CHANGE
    assert by[("B2 ANEXO", "C100")].change_kind == wd.FORMULA_RESULT_CHANGE


def test_provenance_egresos_direct_and_derived(result):
    res, _ = result
    prov = res["prov"]
    assert prov[("B2 ANEXO", "C99")].source == sp.SRC_EGRESOS
    assert prov[("B2 ANEXO", "C99")].status == sp.ST_CLIENT_CONFIRMED
    assert prov[("B2 ANEXO", "C99")].change_origin == sp.ORIGIN_DIRECT_INPUT
    # propagación B2 -> B1
    assert prov[("B2 ANEXO", "C100")].source == sp.SRC_EGRESOS
    assert prov[("B2 ANEXO", "C100")].status == sp.ST_DERIVED_FROM_DEPENDENCIES
    assert prov[("REMASEP B1", "C29")].source == sp.SRC_EGRESOS
    assert prov[("REMASEP B1", "C29")].change_origin == sp.ORIGIN_FORMULA_PROPAGATION
    assert prov[("REMASEP B1", "D13")].source == sp.SRC_EGRESOS


def test_provenance_section_d_and_e(result):
    res, _ = result
    prov = res["prov"]
    assert prov[("REMASEP 01", "B129")].source == sp.SRC_RESOURCE_CALCULATION
    assert prov[("REMASEP 01", "C129")].source == sp.SRC_SURGICAL_TABLE
    assert prov[("REMASEP 01", "C137")].source == sp.SRC_UNKNOWN_PENDING
    assert prov[("REMASEP 01", "C137")].status == sp.ST_UNKNOWN


def test_mixed_derived_when_multiple_sources_feed_a_formula(result):
    res, _ = result
    assert res["prov"][("REMASEP 01", "A307")].source == sp.SRC_MIXED_DERIVED


def test_control_metadata(result):
    res, _ = result
    prov = res["prov"]
    assert prov[("CONTROL", "G8")].source == sp.SRC_CONTROL_METADATA
    assert res["control"]["total_errores"] == 0
    assert res["control"]["control_status"] == "PASS_INTERNAL_VALIDATION"
    assert res["control"]["approval_status"] == "UNCONFIRMED"
    assert len(res["control"]["causales"]) == 1


def test_consistency_check_32_vs_31_style_is_open_question_not_fail(result):
    res, _ = result
    checks = {c["check_id"]: c for c in res["checks"]}
    grid = checks["B1_GRID_TOTAL_VS_CODED_INTERVENTIONS"]
    assert grid["left_value"] == 3 and grid["right_value"] == 2
    assert grid["difference"] == 1
    assert grid["status"] == "OPEN_FUNCTIONAL_QUESTION"
    assert all(c["status"] != "FAIL" for c in res["checks"])
    age_sex = checks["B1_SECCION_A_AGE_TOTAL_VS_SEX_TOTAL"]
    assert age_sex["left_value"] == age_sex["right_value"] == 3
    assert age_sex["status"] == "OK"


def test_artifacts_written_and_named(result):
    _res, out = result
    names = {p.name for p in out.iterdir()}
    assert names == {
        "workbook_structure_comparison.json", "formula_expression_changes.csv",
        "delta_cells.csv", "direct_input_changes.csv", "formula_result_changes.csv",
        "delta_dependency_edges.csv", "provenance_map.csv", "source_summary.csv",
        "source_consistency_checks.csv", "control_summary.csv", "summary.json",
        "README.md",
    }


def test_formula_expression_changes_csv_empty(result):
    _res, out = result
    with open(out / "formula_expression_changes.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows == []


def test_provenance_map_has_row_per_delta_with_review_flag(result):
    res, out = result
    with open(out / "provenance_map.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == len(res["deltas"])
    assert {r["needs_human_review"] for r in rows} <= {"yes", "no"}
    # los UNKNOWN / Sección D piden revisión; EGRESOS confirmados no
    by_cell = {(r["sheet"], r["cell"]): r for r in rows}
    assert by_cell[("REMASEP 01", "C137")]["needs_human_review"] == "yes"
    assert by_cell[("B2 ANEXO", "C99")]["needs_human_review"] == "no"


def test_no_pii_in_any_artifact(result):
    _res, out = result
    for p in out.iterdir():
        assert _SECRET not in p.read_text(encoding="utf-8"), p.name


def test_source_files_unmodified(pair_paths, tmp_path):
    pb, pa = pair_paths
    before = (wd.compute_sha256(pb), wd.compute_sha256(pa))
    af.run(str(pb), str(pa), str(tmp_path / "out2"))
    assert (wd.compute_sha256(pb), wd.compute_sha256(pa)) == before


def test_cli_main_smoke(pair_paths, tmp_path, capsys):
    pb, pa = pair_paths
    code = af.main([str(pb), str(pa), "--output", str(tmp_path / "cli")])
    assert code == 0
    out = capsys.readouterr().out
    assert "candidate_golden_status: CANDIDATE_PENDING_SOURCE_COMPLETENESS" in out
    assert "OPEN_FUNCTIONAL_QUESTION" in out


def test_cli_main_bad_file(tmp_path, capsys):
    bad = tmp_path / "nope.xlsx"
    code = af.main([str(bad), str(bad)])
    assert code == 2


# --- FORMULA_EXPRESSION_CHANGE: provenance UNKNOWN / needs review -----


@pytest.fixture
def expr_change_pair(tmp_path):
    def build(wb, phase):
        ws = wb.active
        ws.title = "REMASEP 01"
        ws["A1"] = 3
        ws["B1"] = "=A1*2" if phase == "before" else "=A1*3"  # expresión reescrita
        ws["A124"] = "SECCIÓN D: CAPACIDAD INSTALADA Y UTILIZACIÓN DE QUIRÓFANOS"

    wb_b, wb_a = Workbook(), Workbook()
    build(wb_b, "before")
    build(wb_a, "after")
    pb, pa = tmp_path / "b.xlsx", tmp_path / "a.xlsx"
    wb_b.save(pb)
    wb_a.save(pa)
    _inject_cache(pb, {"REMASEP 01": {"B1": 6}})
    _inject_cache(pa, {"REMASEP 01": {"B1": 9}})
    return pb, pa


def test_expression_change_provenance_is_unknown_pending(expr_change_pair, tmp_path):
    pb, pa = expr_change_pair
    res = af.run(str(pb), str(pa), str(tmp_path / "out"))
    deltas = {d.key: d for d in res["deltas"]}
    key = ("REMASEP 01", "B1")
    assert deltas[key].change_kind == wd.FORMULA_EXPRESSION_CHANGE
    prov = res["prov"][key]
    assert prov.source == sp.SRC_UNKNOWN_PENDING
    assert prov.status == sp.ST_UNKNOWN
    assert prov.change_origin == sp.ORIGIN_FORMULA_EXPRESSION_CHANGE
    assert prov.needs_human_review is True

    # summary lo cuenta aparte, no como directo ni derivado
    s = res["summary"]
    assert s["formula_expression_value_changes"] == 1
    assert s["formula_expression_changes"] == 1  # diff_formulas también

    # no aparece en direct_input_changes.csv ni formula_result_changes.csv
    with open(tmp_path / "out" / "direct_input_changes.csv", encoding="utf-8", newline="") as fh:
        assert "B1" not in {r["cell"] for r in csv.DictReader(fh)}
    with open(tmp_path / "out" / "formula_result_changes.csv", encoding="utf-8", newline="") as fh:
        assert "B1" not in {r["cell"] for r in csv.DictReader(fh)}
    # sí en delta_cells.csv con before_formula y formula_change_type
    with open(tmp_path / "out" / "delta_cells.csv", encoding="utf-8", newline="") as fh:
        row = next(r for r in csv.DictReader(fh) if r["cell"] == "B1")
    assert row["change_kind"] == "FORMULA_EXPRESSION_CHANGE"
    assert row["before_formula"] == "=A1*2" and row["formula"] == "=A1*3"
    assert row["formula_change_type"] == wd.FORMULA_MODIFIED
