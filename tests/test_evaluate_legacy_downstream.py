"""Tests del cierre completo BASE/DERIVED/DOWNSTREAM_TOTAL/VALIDATION (Sprint 2.5).

Workbook sintético: hoja de detalle en posiciones reales + `REMASEP_OD` con
BASE, DERIVED, `SUM` y `IF`, con valores cacheados parcheados en el XML. No se
usa data/local.
"""

from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
import zipfile

import close_legacy_aggregations as clo
import evaluate_legacy_downstream as ev
import pytest
from openpyxl import Workbook

_DETAIL = "Atenciones - Detalles de citas"
_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}


def _sheet_xml_map(path) -> dict[str, str]:
    ns_m = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as zf:
        wbxml = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_target = {rel.get("Id"): rel.get("Target") for rel in rels}
    out = {}
    for sheet in wbxml.find(f"{{{ns_m}}}sheets"):
        target = rid_target[sheet.get(f"{{{ns_r}}}id")].lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        out[sheet.get("name")] = target
    return out


@pytest.fixture
def build_downstream_workbook(tmp_path, inject_formula_cache):
    def _build(detail_rows, sheet_formulas, cache=None, *, filename="downstream.xlsx"):
        cache = cache or {}
        wb = Workbook()
        detail = wb.active
        detail.title = _DETAIL
        for letter, header in _HDR.items():
            detail[f"{letter}1"] = header
        for i, row in enumerate(detail_rows, start=2):
            for sem, letter in _COLS.items():
                if row.get(sem) is not None:
                    detail[f"{letter}{i}"] = row[sem]
            detail[f"BA{i}"] = "x"
        for sheet_name, formulas in sheet_formulas.items():
            ws = wb.create_sheet(sheet_name)
            for coord, formula in formulas.items():
                ws[coord] = formula
        path = tmp_path / filename
        wb.save(path)

        xml_map = _sheet_xml_map(path)
        for sheet_name, cells in cache.items():
            patch = {(int(c[1:]), c[0]): v for c, v in cells.items()}
            if patch:
                inject_formula_cache(path, patch, sheet_xml=xml_map[sheet_name])
        return path

    return _build


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2000", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "COMPOSITE", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


_AE_H = '"*COMPOSITE*"&"*Odontología General*"&"*Hombre*"'
_AE_ANY = '"*COMPOSITE*"&"*Odontología General*"'


def _countifs(ae: str, af_crit: str) -> str:
    return f"=COUNTIFS('{_DETAIL}'!$AE:$AE,{ae},'{_DETAIL}'!$AF:$AF,\"{af_crit}\")"


@pytest.fixture
def scenario(build_downstream_workbook):
    rows = [
        _row(FECHA_NACIMIENTO="01-01-2025"),  # AF 1, Hombre
        _row(FECHA_NACIMIENTO="01-06-2025"),  # AF 1, Hombre
        _row(FECHA_NACIMIENTO="01-01-2020", SEXO="Mujer"),  # AF 6, Mujer
        {},  # fila estructuralmente vacía
    ]
    od = {
        # --- base / derived (Sprint 2.3/2.4) ---------------------------
        "F30": _countifs(_AE_H, "<2"),                       # -> 2 (Hombre, AF<2)
        "F46": _countifs(_AE_ANY, "<2") + "-F30",            # 2 - 2 = 0 (Mujer<2 -> 0)
        # --- SUM downstream (Sprint 2.5) -------------------------------
        "P10": "=SUM(F30:F46)",                              # rango, incluye blancos -> 2+0=2
        "P11": "=SUM(P10)",                                  # SUM de SUM (profundidad+1)
        "P12": "=SUM(P11)",                                  # cadena profundidad > 1
        # --- IF validation (Sprint 2.5) --------------------------------
        "U1": "=IF(P10<100,1,0)",                            # ref < literal
        "U2": "=IF(P10<=P10,\"yes\",\"no\")",                # ref <= ref
        "U3": "=IF(P11>P10,1,0)",                            # ref > ref (depende de SUM)
        "U4": "=IF(P11>=P10,1,0)",                           # ref >= ref
        "U5": "=IF(F46=0,\"zero\",\"nonzero\")",             # ref = literal (depende de DERIVED); F46=0 -> "zero"
        "U6": "=IF(F46<>0,\"nonzero\",\"zero\")",            # ref <> literal; F46=0 -> "zero"
        "W1": "=IF(F30<>0, IF(F46=0,\"clean\",\"\"), \"\")",  # IF anidado; F30<>0 y F46=0 -> "clean"
        # --- dependencia no disponible ----------------------------------
        "Y1": "=SUM(F30*2)",                                  # UNSUPPORTED_SUM
        "Y2": "=IF(Y1>0,1,0)",                                # depende de Y1 -> DEPENDENCY_UNAVAILABLE
        # --- IF no soportado ---------------------------------------------
        "Z1": "=IF(F30>0,1,0,0)",                             # UNSUPPORTED_IF (4 args)
        # --- ciclo (referencia como resultado) ---------------------------
        "CY1": "=IF(F30>0,CY2,0)",
        "CY2": "=IF(F30>0,CY1,0)",
    }
    cache = {
        "REMASEP_OD": {
            "F30": 2, "F46": 0, "P10": 2, "P11": 2, "P12": 2,
            "U1": 1, "U2": "yes", "U3": 0, "U4": 1,
            "U5": "zero", "U6": "nonzero",  # U6 deliberadamente "mal" -> DIFFERENCE (real: "zero")
            "W1": "clean",
            # Y1, Y2, Z1, CY1, CY2 sin cache -> UNAVAILABLE si algún día se evaluaran
        }
    }
    path = build_downstream_workbook(rows, {"REMASEP_OD": od}, cache)
    return ev.evaluate_full_closure(path), path


def _nodes(result):
    return {(n.sheet, n.cell): n for n in result.medinet_nodes()}


# --- SUM -----------------------------------------------------------


def test_sum_of_range_including_derived(scenario):
    result, _ = scenario
    nodes = _nodes(result)
    p10 = nodes[("REMASEP_OD", "P10")]
    assert p10.kind == clo.KIND_TOTAL
    assert p10.python_value == 2  # F30(2) + F46(0) + blancos(0)
    assert p10.evaluation_status == ev.EVAL_SUPPORTED


def test_sum_of_sum_chain_and_depth(scenario):
    result, _ = scenario
    nodes = _nodes(result)
    p10, p11, p12 = nodes[("REMASEP_OD", "P10")], nodes[("REMASEP_OD", "P11")], nodes[("REMASEP_OD", "P12")]
    assert p11.python_value == 2
    assert p12.python_value == 2
    assert p10.depth < p11.depth < p12.depth
    assert p12.depth >= 3  # F30(0) -> F46(1) -> P10(2) -> P11(3) -> P12(4)


def test_unsupported_sum(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "Y1")]
    assert node.evaluation_status == ev.EVAL_UNSUPPORTED_SUM
    assert node.python_value is None


# --- IF --------------------------------------------------------------


@pytest.mark.parametrize(
    ("cell", "expected"),
    [("U1", 1), ("U2", "yes"), ("U3", 0), ("U4", 1), ("U5", "zero"), ("U6", "zero")],
)
def test_if_operators_evaluate(scenario, cell, expected):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", cell)]
    assert node.evaluation_status == ev.EVAL_SUPPORTED
    assert node.python_value == expected


def test_if_depends_on_sum_and_derived(scenario):
    result, _ = scenario
    nodes = _nodes(result)
    assert nodes[("REMASEP_OD", "U3")].kind == clo.KIND_VALIDATION
    assert nodes[("REMASEP_OD", "U5")].depends_on_age is True  # via F46 <- F30 (AF)


def test_nested_if_depends_on_base_and_derived(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "W1")]
    assert node.evaluation_status == ev.EVAL_SUPPORTED
    # F30=2<>0 (True) -> nested IF(F46=0,"clean","") ; F46=0 -> "clean"
    assert node.python_value == "clean"
    assert node.cache_status == clo.CACHE_MATCH


def test_unsupported_if(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "Z1")]
    assert node.evaluation_status == ev.EVAL_UNSUPPORTED_IF


# --- dependencia no disponible / ciclo -----------------------------


def test_dependency_unavailable_propagates(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "Y2")]
    assert node.evaluation_status == ev.EVAL_DEPENDENCY_UNAVAILABLE
    assert node.python_value is None


def test_cycle_not_evaluated(scenario):
    result, _ = scenario
    nodes = _nodes(result)
    cy1, cy2 = nodes[("REMASEP_OD", "CY1")], nodes[("REMASEP_OD", "CY2")]
    assert cy1.evaluation_status == ev.EVAL_CYCLE
    assert cy2.evaluation_status == ev.EVAL_CYCLE
    assert cy1.python_value is None


def test_downstream_does_not_break_other_evaluations(scenario):
    result, _ = scenario
    summary = ev._summary_dict(result)
    assert summary["base_supported"] == summary["base_total"]
    assert summary["derived_supported"] == summary["derived_total"]


# --- depends_on_AF propagado ----------------------------------------


def test_depends_on_af_propagates_through_sum_and_if(scenario):
    result, _ = scenario
    nodes = _nodes(result)
    assert nodes[("REMASEP_OD", "F30")].depends_on_age is True
    assert nodes[("REMASEP_OD", "P10")].depends_on_age is True   # SUM incluye F30
    assert nodes[("REMASEP_OD", "U1")].depends_on_age is True    # IF depende de P10


# --- cache -------------------------------------------------------------


def test_cache_match(scenario):
    result, _ = scenario
    assert _nodes(result)[("REMASEP_OD", "P10")].cache_status == clo.CACHE_MATCH


def test_cache_difference(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "U6")]
    assert node.cache_status == clo.CACHE_DIFFERENCE
    assert node.python_value == "zero"
    assert node.cached_value == "nonzero"


def test_cache_unavailable(scenario):
    result, _ = scenario
    node = _nodes(result)[("REMASEP_OD", "CY1")]  # nunca se evalúa -> no se compara cache
    assert node.cache_status == ""
    node2 = _nodes(result)[("REMASEP_OD", "Y1")]
    assert node2.cache_status == ""  # UNSUPPORTED_SUM tampoco se compara


# --- artefactos / privacidad / metric_id ---------------------------


def test_full_equivalence_has_all_medinet_nodes(scenario, tmp_path):
    result, _ = scenario
    ev.write_outputs(result, tmp_path / "out")
    with open(tmp_path / "out" / "full_equivalence.csv", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == len(result.medinet_nodes())
    by_id = {r["metric_id"]: r for r in rows}
    assert "LEGACY::REMASEP_OD::P10" in by_id
    assert by_id["LEGACY::REMASEP_OD::P10"]["kind"] == "DOWNSTREAM_TOTAL"
    assert by_id["LEGACY::REMASEP_OD::P10"]["evaluation_status"] == "SUPPORTED"


def test_validation_summary_only_contains_validation_kind(scenario, tmp_path):
    result, _ = scenario
    ev.write_outputs(result, tmp_path / "out")
    with open(tmp_path / "out" / "validation_summary.csv", encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    cells = {r["cell"] for r in rows}
    assert {"U1", "U2", "U3", "U4", "U5", "U6", "W1"} <= cells
    assert "P10" not in cells  # P10 es DOWNSTREAM_TOTAL, no VALIDATION


def test_summary_json_separates_support_and_cache(scenario, tmp_path):
    result, _ = scenario
    ev.write_outputs(result, tmp_path / "out")
    summary = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))
    assert summary["unsupported_cells"] >= 2  # Y1 (SUM) + Z1 (IF)
    assert summary["dependency_unavailable_cells"] >= 1  # Y2
    assert summary["cycles"] >= 2  # CY1, CY2
    assert summary["formula_support_status"] == "PARTIAL"
    assert summary["cache_consistency_status"] == "DIFFERENCES"
    assert summary["fully_evaluated_cells"] < summary["total_medinet_dependent_cells"]


def test_artifacts_contain_no_individual_patient_data(build_downstream_workbook, tmp_path):
    secret = "SECRET-PATIENT-XYZ"
    rows = [
        _row(FECHA_NACIMIENTO="01-01-2025", ESTADO=secret, PRESTACION_REALIZADA=secret),
        _row(FECHA_NACIMIENTO="01-01-2020", ESTADO=secret),
    ]
    od = {
        "F30": _countifs(_AE_H, "<2"),
        "P10": "=SUM(F30:F30)",
        "U1": "=IF(P10>0,1,0)",
    }
    path = build_downstream_workbook(rows, {"REMASEP_OD": od}, {"REMASEP_OD": {"F30": 1, "P10": 1, "U1": 1}})
    result = ev.evaluate_full_closure(path)
    files = ev.write_outputs(result, tmp_path / "out")
    for f in files:
        assert secret not in f.read_text(encoding="utf-8"), f.name


# --- errores de entrada / CLI ---------------------------------------


def test_missing_file(tmp_path):
    with pytest.raises(clo.ClosureError):
        ev.evaluate_full_closure(tmp_path / "nope.xlsx")


def test_cli_main(scenario, tmp_path, capsys):
    _result, path = scenario
    code = ev.main([str(path), "--output", str(tmp_path / "out")])
    assert code == 0  # PARTIAL no es FAIL
    out = capsys.readouterr().out
    assert "formula_support_status: PARTIAL" in out
    assert "cache_consistency_status: DIFFERENCES" in out
