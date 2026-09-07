"""Tests CLI de la validación / readiness semántica (Sprint 3.3).

Workbook sintético (mismo patrón que `test_map_semantic_metrics.py`): produce las
cuatro readiness (`AUTO_READY`, `REVIEW_REQUIRED`, `BLOCKED_CONFLICT`,
`NOT_APPLICABLE`). Sin data/local.
"""

from __future__ import annotations

import csv
import json

import pytest
import validate_semantic_mapping as vs
from openpyxl import Workbook
from openpyxl.styles import Font

from remasep.services import semantic_validation as sv

_DETAIL = "Atenciones - Detalles de citas"
_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}
_BOLD = Font(bold=True)


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2003", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "VIDRIO IONÓMERO", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _countifs(sex_pat: str, lo: int, hi: int) -> str:
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*VIDRIO IONÓMERO*\"&\"*{sex_pat}*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


def _countifs_nosex(lo: int, hi: int) -> str:
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*VIDRIO IONÓMERO*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


@pytest.fixture
def scenario(tmp_path):
    wb = Workbook()
    detail = wb.active
    detail.title = _DETAIL
    for letter, header in _HDR.items():
        detail[f"{letter}1"] = header
    rows = [
        _row(FECHA_NACIMIENTO="01-01-2003"),
        _row(FECHA_NACIMIENTO="01-01-2003", SEXO="Mujer"),
        _row(FECHA_NACIMIENTO="01-01-1999", SEXO="Mujer", ESTADO="SECRET-XYZ"),
        {},
    ]
    for i, r in enumerate(rows, start=2):
        for sem, letter in _COLS.items():
            if r.get(sem) is not None:
                detail[f"{letter}{i}"] = r[sem]
        detail[f"BA{i}"] = "x"

    od = wb.create_sheet("REMASEP_OD")
    od["B1"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["B1"].font = _BOLD
    od.merge_cells("C2:D2")
    od["C2"] = "20 A 24 AÑOS"
    od["E2"] = "TOTAL"
    od["C3"] = "Hombres"
    od["D3"] = "Mujeres"
    od["A6"] = "5010009 - VIDRIO IONÓMERO"
    od["B6"] = "Obturación de vidrio ionómero"
    od["C6"] = _countifs("Hombre", 20, 24)     # CONFIRMED  -> AUTO_READY
    od["D6"] = _countifs("Hombre", 20, 24)     # label Mujeres / formula Hombre -> CONFLICT
    od["E6"] = "=SUM(C6+D6)"                    # columna TOTAL -> NOT_APPLICABLE
    od["A7"] = "5010010 - OTRO INSUMO"
    od["C7"] = _countifs("Mujer", 20, 24)      # label Hombres / formula Mujer -> CONFLICT
    od["A8"] = "5010011 - TERCER INSUMO"
    od["C8"] = _countifs_nosex(20, 24)         # label Hombres, formula sin sexo -> LABEL_ONLY

    b2 = wb.create_sheet("B2 ANEXO")
    b2["B3"] = "A. KINESIOLOGÍA"
    b2["B3"].font = _BOLD
    b2["D3"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"x\")"
    b2["B5"] = "0601105"
    b2["C5"] = "Atención Kinesiológica Integral"
    b2["D5"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"y\")"

    path = tmp_path / "mapping.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def built(scenario):
    return vs.build_validation(scenario)


def _readiness_by_cell(result):
    return {(r.sheet, r.cell): r for r in result.readiness}


# --- readiness ----------------------------------------------------


def test_four_readiness_states_present(built):
    _inv, _cfg, _metrics, _pol, _ovr, result = built
    by = _readiness_by_cell(result)
    assert by[("REMASEP_OD", "C6")].readiness_status == sv.AUTO_READY
    assert by[("REMASEP_OD", "D6")].readiness_status == sv.BLOCKED_CONFLICT
    assert by[("REMASEP_OD", "C7")].readiness_status == sv.BLOCKED_CONFLICT
    assert by[("REMASEP_OD", "E6")].readiness_status == sv.NOT_APPLICABLE
    assert by[("REMASEP_OD", "C8")].readiness_status == sv.REVIEW_REQUIRED
    assert by[("B2 ANEXO", "D5")].readiness_status == sv.AUTO_READY


def test_conflict_metric_never_auto_ready(built):
    *_, result = built
    for r in result.readiness:
        if any(i.severity == sv.SEV_BLOCKING for i in r.issues):
            assert r.readiness_status == sv.BLOCKED_CONFLICT


# --- artefactos --------------------------------------------------


def test_write_outputs_file_set(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    files = vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    names = {f.name for f in files}
    assert names == {
        "semantic_metric_readiness.csv", "review_queue.csv", "conflicts.csv",
        "review_reason_summary.csv", "review_clusters.csv", "readiness_coverage.csv",
        "readiness_manual_review_sample.csv", "summary.json", "README.md",
    }


def test_conflicts_csv_records_not_resolves(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "conflicts.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert {r["cell"] for r in rows} == {"D6", "C7"}
    for r in rows:
        assert r["issue_type"] == "SEX_LABEL_FORMULA_MISMATCH"
        assert r["recommended_action"] == "HUMAN_REVIEW"
        assert r["status"] == "OPEN"
        # se registra la evidencia; NO se propone un valor
        assert {r["label_evidence"], r["formula_evidence"]} == {"MALE", "FEMALE"}
        blob = " ".join(r.values()).lower()
        assert "change to" not in blob


def test_review_queue_only_blocking_and_review(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "review_queue.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows  # no vacío
    assert all(r["severity"] in (sv.SEV_BLOCKING, sv.SEV_REVIEW) for r in rows)
    assert all(r["recommended_action"] == "HUMAN_REVIEW" for r in rows)


def test_review_clusters_group_equivalent_issues(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "review_clusters.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    sex_conflict = next(r for r in rows if r["issue_type"] == "SEX_LABEL_FORMULA_MISMATCH")
    assert int(sex_conflict["metric_count"]) == 2  # D6 y C7 en un solo cluster
    assert sex_conflict["severity"] == sv.SEV_BLOCKING


def test_readiness_coverage_has_total_row(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "readiness_coverage.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    total = next(r for r in rows if r["scope"] == "TOTAL")
    assert int(total["metric_count"]) == len(result.readiness)
    summed = sum(int(total[s]) for s in (sv.AUTO_READY, sv.REVIEW_REQUIRED,
                                         sv.BLOCKED_CONFLICT, sv.NOT_APPLICABLE))
    assert summed == len(result.readiness)


def test_manual_review_sample_deterministic(built, tmp_path):
    *_, result = built
    s1 = vs._manual_review_sample(result)
    s2 = vs._manual_review_sample(result)
    assert [r.metric_id for r in s1] == [r.metric_id for r in s2]
    # todos los BLOCKED_CONFLICT en la muestra
    blocked = {r.metric_id for r in result.readiness
               if r.readiness_status == sv.BLOCKED_CONFLICT}
    assert blocked <= {r.metric_id for r in s1}


def test_summary_consistency(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    summary = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))
    assert summary["validation_status"] == "PRELIMINARY_NOT_VALIDATED"
    assert summary["total_metrics"] == len(result.readiness)
    assert (summary["auto_ready"] + summary["review_required"]
            + summary["blocked_conflict"] + summary["not_applicable"]) == len(result.readiness)
    assert summary["blocked_conflict"] == 2
    assert summary["override_count"] == 0
    # el tamaño reportado coincide con la muestra escrita
    with open(tmp_path / "out" / "readiness_manual_review_sample.csv",
              encoding="utf-8", newline="") as fh:
        sample_rows = list(csv.DictReader(fh))
    assert summary["readiness_manual_review_sample_size"] == len(sample_rows)


def test_no_individual_patient_data_in_artifacts(built, tmp_path):
    inv, _cfg, metrics, pol, ovr, result = built
    files = vs.write_outputs(inv, result, pol, ovr, metrics, tmp_path / "out")
    for f in files:
        assert "SECRET-XYZ" not in f.read_text(encoding="utf-8"), f.name


def test_overrides_empty_and_no_silent_resolution(built):
    *_, ovr, result = built
    assert ovr == {}
    assert all(r.override_matched is False for r in result.readiness)


# --- CLI --------------------------------------------------------


def test_cli_main(scenario, tmp_path, capsys):
    code = vs.main([str(scenario), "--output", str(tmp_path / "out")])
    assert code == 0
    out = capsys.readouterr().out
    assert "AUTO_READY" in out and "BLOCKED_CONFLICT" in out
    assert (tmp_path / "out" / "summary.json").is_file()


def test_cli_bad_workbook_returns_2(tmp_path, capsys):
    bad = tmp_path / "nope.xlsx"
    bad.write_text("not a workbook", encoding="utf-8")
    code = vs.main([str(bad), "--output", str(tmp_path / "out")])
    assert code == 2
