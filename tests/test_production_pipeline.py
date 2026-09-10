"""Sprint 3.8 — contrato de runtime limpio: producción sin workbook legacy.

El camino de PRODUCCIÓN construye 1122 MetricValues/PendingWrites usando **sólo**
el export Medinet + los assets de runtime versionados. No abre
``GENERACION DATOS REMASEP.xlsx`` ni lee ``artifacts/``.
"""

from __future__ import annotations

import builtins
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from remasep.core.errors import RemasepError
from remasep.services.common import Period
from remasep.services.production_pipeline import (
    build_production_metric_values,
    build_production_pending_writes,
)
from remasep.services.runtime_assets import RuntimeAssetError

_REAL_ROOT = Path(__file__).resolve().parent.parent / "config" / "runtime_2026"
_LEGACY = Path("data/local/GENERACION DATOS REMASEP.xlsx")
_MEDINET_REAL = Path("data/local/detalle_citas - 2026-09-07T123630.940.xlsx")

_ROW = [
    date(2026, 7, 10), date(1990, 1, 1), "Mujer", "Santiago", "ODONTOLOGIA",
    "CONSULTA GENERAL", "EVALUACIÓN ODONTOLÓGICA", "Atendido", "Presencial", "",
]


@pytest.fixture
def _forbid_legacy_and_artifacts(monkeypatch):
    """Cualquier intento de abrir el workbook legacy o algo bajo artifacts/ falla."""
    forbidden = ("GENERACION DATOS REMASEP.xlsx",)
    art = f"{Path('artifacts').resolve()}"
    real_open = builtins.open
    real_load = openpyxl.load_workbook

    def guarded_open(file, *a, **k):
        s = str(Path(file).resolve()) if not str(file).startswith("<") else str(file)
        if any(f in str(file) for f in forbidden) or s.startswith(art):
            raise AssertionError(f"producción intentó abrir {file!r}")
        return real_open(file, *a, **k)

    def guarded_load(fn, *a, **k):
        if any(f in str(fn) for f in forbidden):
            raise AssertionError(f"producción intentó openpyxl.load_workbook({fn!r})")
        return real_load(fn, *a, **k)

    monkeypatch.setattr(builtins, "open", guarded_open)
    monkeypatch.setattr(openpyxl, "load_workbook", guarded_load)


# ---------------------------------------------------------------------------
# Clean runtime contract (sintético, sin data/local)
# ---------------------------------------------------------------------------


def test_production_builds_1122_pending_writes_from_runtime_only(
    make_medinet, _forbid_legacy_and_artifacts
):
    medinet = make_medinet([_ROW, _ROW, list(_ROW[:6]) + ["OTRA", "Atendido", "Presencial", ""]])
    result = build_production_pending_writes(medinet, Period(7, 2026))

    assert result.bundle_version == "runtime_2026"
    assert result.template_fingerprint_id == "stf:dc624775927d4d4d"
    assert result.zero_write_policy == "WRITE_ZERO"
    assert result.estado_filter_status == "PENDING_FUNCTIONAL_CONFIRMATION"
    # cada instrucción WRITE_READY tiene exactamente un MetricValue
    assert len(result.run.metric_values) == 1122
    assert len(result.pending_writes) == 1122
    assert result.completeness.ok is True
    assert not result.completeness.missing
    assert not result.completeness.duplicate
    assert not result.completeness.orphan
    assert result.run.value_type_conflicts == []
    assert result.run.unsupported == []
    # todos los valores son enteros >= 0 (INTEGER_COUNT)
    for pw in result.pending_writes:
        assert isinstance(pw.value, int) and not isinstance(pw.value, bool)
        assert pw.value >= 0
    # targets únicos
    cells = [(pw.target_sheet, pw.target_cell) for pw in result.pending_writes]
    assert len(cells) == len(set(cells))


def test_production_does_not_apply_estado_filter(make_medinet, _forbid_legacy_and_artifacts):
    """Modo producción: input_scope PERIOD_ONLY, sin filtro por ESTADO."""
    rows = [_ROW]  # Atendido
    cancelled = list(_ROW)
    cancelled[7] = "Cancelado"
    rows.append(cancelled)
    medinet = make_medinet(rows)
    run, _bundle = build_production_metric_values(medinet, Period(7, 2026))
    assert run.mode == "PRODUCTION_PERIOD_SCOPE"
    assert run.estado_filter_applied is False
    assert run.input_scope == "PERIOD_ONLY"
    # ambas filas entran al alcance (no se filtra Cancelado)
    assert run.scope_records == 2


def test_out_of_period_medinet_rows_do_not_count(make_medinet, _forbid_legacy_and_artifacts):
    in_july = _ROW
    other_month = list(_ROW)
    other_month[0] = date(2026, 8, 10)
    medinet = make_medinet([in_july, other_month])
    run, _b = build_production_metric_values(medinet, Period(7, 2026))
    assert run.scope_records == 1


def test_missing_runtime_bundle_is_controlled_error(monkeypatch, make_medinet):
    # sin config/runtime_2026 resoluble -> RuntimeAssetError (RemasepError), no crudo
    import remasep.services.runtime_assets as ra

    def _boom(explicit=None):
        raise RuntimeAssetError("RUNTIME_ASSET_MISSING", "simulado")

    monkeypatch.setattr(ra, "resolve_runtime_root", _boom)
    with pytest.raises(RuntimeAssetError) as exc:
        build_production_pending_writes(make_medinet([_ROW]), Period(7, 2026))
    assert exc.value.code == "RUNTIME_ASSET_MISSING"
    assert isinstance(exc.value, RemasepError)
    assert "instalación no contiene" in exc.value.user_message


# ---------------------------------------------------------------------------
# Equivalencia OLD (legacy workbook) vs NEW (runtime assets) — gated
# ---------------------------------------------------------------------------

_local_available = _LEGACY.is_file() and _MEDINET_REAL.is_file()


@pytest.mark.skipif(not _local_available, reason="requiere data/local de desarrollo")
def test_old_vs_new_produce_identical_1122_pending_writes():
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
    from build_metric_values import (  # type: ignore[import-not-found]
        build_legacy_reference,
        load_manifest,
        run_producer,
    )

    from remasep.services.medinet_analysis import processing_scope_frame
    from remasep.services.metric_value_producer import join_metric_values_with_manifest

    period = Period(7, 2026)

    reference = build_legacy_reference(_LEGACY)
    old_manifest = load_manifest(
        Path("artifacts/writable_target_mapping/write_manifest_ready.csv")
    )
    old_frame = processing_scope_frame(_MEDINET_REAL, period)
    old_run = run_producer(
        old_frame, reference, old_manifest, period, mode="PRODUCTION_PERIOD_SCOPE"
    )
    old_pending = join_metric_values_with_manifest(old_run.metric_values, old_manifest)

    new = build_production_pending_writes(_MEDINET_REAL, period)

    assert len(old_pending) == len(new.pending_writes) == 1122

    old_map = {
        p.instruction_id: (p.source_metric_id, p.target_sheet, p.target_cell, p.value)
        for p in old_pending
    }
    new_map = {
        p.instruction_id: (p.source_metric_id, p.target_sheet, p.target_cell, p.value)
        for p in new.pending_writes
    }
    assert old_map == new_map

    old_mv = {mv.source_metric_id: mv.value for mv in old_run.metric_values}
    new_mv = {mv.source_metric_id: mv.value for mv in new.run.metric_values}
    assert old_mv == new_mv
