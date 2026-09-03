"""Tests del inventario de dependencias (Sprint 1.2, Partes B-E).

El workbook se construye sintéticamente en cada test: sin datos reales.
"""

import csv

import analyze_dependencies as ad
import pytest
from openpyxl import Workbook


@pytest.fixture
def synthetic_workbook(tmp_path):
    wb = Workbook()

    detail = wb.active
    detail.title = "Detail"
    for col in "ABCDEFGH":
        detail[f"{col}1"] = f"h_{col}"
    for row in (2, 3, 4):
        for col in "ABCDEFGH":
            detail[f"{col}{row}"] = row
        detail[f"I{row}"] = f"=IF(A{row}>0,B{row},C{row})"
        detail[f"J{row}"] = f'=DATEDIF(D{row},E{row},"Y")'
        detail[f"K{row}"] = f"=A{row}&F{row}"

    out1 = wb.create_sheet("Out1")
    out1["A1"] = '=COUNTIF(\'Detail\'!$I:$I,">0")'
    out1["A2"] = '=SUMIFS(\'Detail\'!$K:$K,\'Detail\'!$A:$A,"x")'
    out1["A3"] = "=Out2!B1+1"
    out1["A4"] = '=INDIRECT("Detail!A1")'

    out2 = wb.create_sheet("Out2")
    out2["B1"] = "=SUM('Detail'!J2:J4)"
    out2["B2"] = "=B1*2"

    path = tmp_path / "synthetic.xlsx"
    wb.save(path)
    return path


def _rows(path):
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _run(workbook, tmp_path, **kwargs):
    out = tmp_path / "inv"
    doc = tmp_path / "docs" / "FLOW.md"
    kwargs.setdefault("detail_sheet", "Detail")
    result = ad.analyze_dependencies(workbook, out, flow_doc=doc, **kwargs)
    return out, doc, result


# --- Parte B ---------------------------------------------------------------


def test_derived_columns(synthetic_workbook, tmp_path):
    out, _doc, result = _run(synthetic_workbook, tmp_path)
    rows = {r["target_column"]: r for r in _rows(out / "derived_columns.csv")}

    assert set(rows) == {"I", "J", "K"}
    assert result.derived_column_count == 3

    assert rows["I"]["formula_count"] == "3"
    assert rows["I"]["first_formula_cell"] == "I2"
    assert rows["I"]["example_formula"] == "=IF(A2>0,B2,C2)"
    assert rows["I"]["source_columns"] == "A|B|C"
    assert rows["I"]["source_sheets"] == "(local)"
    assert rows["I"]["function_names"] == "IF"
    assert rows["I"]["pattern_count"] == "1"

    assert rows["J"]["source_columns"] == "D|E"
    assert rows["J"]["function_names"] == "DATEDIF"
    assert rows["K"]["source_columns"] == "A|F"
    assert rows["K"]["function_names"] == ""


def test_derived_columns_only_cover_detail_sheet(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    # Out1/Out2 no aparecen como columnas derivadas.
    cells = {r["first_formula_cell"] for r in _rows(out / "derived_columns.csv")}
    assert cells == {"I2", "J2", "K2"}


# --- Parte C ---------------------------------------------------------------


def test_formula_dependencies_cross_sheet(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = _rows(out / "formula_dependencies.csv")

    def has(**kw):
        return any(all(r[k] == v for k, v in kw.items()) for r in rows)

    assert has(
        target_sheet="Out1",
        target_cell="A1",
        source_sheet="Detail",
        source_reference_type="whole_column",
        source_column_start="I",
        source_column_end="I",
        is_cross_sheet="True",
    )
    assert has(
        target_sheet="Out2",
        target_cell="B1",
        source_sheet="Detail",
        source_reference="'Detail'!J2:J4",
        source_reference_type="range",
        source_column_start="J",
        source_column_end="J",
        is_cross_sheet="True",
    )
    # Referencia local: misma hoja, no cross-sheet.
    assert has(
        target_sheet="Detail",
        target_cell="I2",
        source_sheet="Detail",
        source_reference="A2",
        is_cross_sheet="False",
    )
    # Cross-sheet entre dos hojas output.
    assert has(target_sheet="Out1", target_cell="A3", source_sheet="Out2", is_cross_sheet="True")


def test_formula_dependencies_symbolic_whole_column(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = _rows(out / "formula_dependencies.csv")
    whole_cols = [r["source_reference"] for r in rows if r["source_reference_type"] == "whole_column"]
    # A:A no se expande: se guarda simbólicamente.
    assert "'Detail'!$I:$I" in whole_cols
    assert all(":" in ref for ref in whole_cols)


# --- Parte D ---------------------------------------------------------------


def test_sheet_dependencies(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    edges = {(r["source_sheet"], r["target_sheet"]): int(r["reference_count"]) for r in _rows(out / "sheet_dependencies.csv")}

    assert edges[("Detail", "Out1")] == 3  # I:I  +  K:K, A:A
    assert edges[("Detail", "Out2")] == 1  # J2:J4
    assert edges[("Out2", "Out1")] == 1  # =Out2!B1+1
    assert edges[("Detail", "Detail")] == 21  # 3*(3 + 2 + 2)
    assert edges[("Out2", "Out2")] == 1  # =B1*2
    assert ("Out1", "Out1") not in edges  # A4 INDIRECT: literal en string, sin refs locales


def test_flow_doc_has_mermaid_and_only_observed_edges(synthetic_workbook, tmp_path):
    _out, doc, _result = _run(synthetic_workbook, tmp_path)
    text = doc.read_text(encoding="utf-8")

    assert "```mermaid" in text
    assert 'graph LR' in text
    for label in ("Detail", "Out1", "Out2"):
        assert f'"{label}"' in text
    # Aristas observadas presentes; nada inventado hacia/desde hojas inexistentes.
    assert "|3|" in text and "|1|" in text
    assert "Nonexistent" not in text


# --- Parte E ---------------------------------------------------------------


def test_output_column_dependencies(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = _rows(out / "output_column_dependencies.csv")
    by_key = {(r["target_sheet"], r["source_column"]): r for r in rows}

    assert set(by_key) == {
        ("Out1", "A"),
        ("Out1", "I"),
        ("Out1", "K"),
        ("Out2", "J"),
    }
    assert by_key[("Out1", "I")]["formulas_using_column"] == "1"
    assert by_key[("Out1", "I")]["example_target_cell"] == "A1"
    assert by_key[("Out2", "J")]["example_target_cell"] == "B1"
    # Orden: por hoja y luego por índice de columna (A antes que I antes que K).
    out1_cols = [r["source_column"] for r in rows if r["target_sheet"] == "Out1"]
    assert out1_cols == ["A", "I", "K"]


def test_output_sheets_override(synthetic_workbook, tmp_path):
    out, _doc, result = _run(synthetic_workbook, tmp_path, output_sheets=["Out1"])
    assert result.output_sheets == ["Out1"]
    rows = _rows(out / "output_column_dependencies.csv")
    assert {r["target_sheet"] for r in rows} == {"Out1"}


# --- Unsupported constructs ----------------------------------------------


def test_unsupported_constructs_reported_not_inferred(synthetic_workbook, tmp_path):
    out, _doc, result = _run(synthetic_workbook, tmp_path)
    rows = _rows(out / "unsupported_constructs.csv")

    indirect = [r for r in rows if r["construct_type"] == "indirect"]
    assert len(indirect) == 1
    assert indirect[0]["sheet"] == "Out1"
    assert indirect[0]["example_cell"] == "A4"
    assert indirect[0]["occurrences"] == "1"
    assert result.unsupported_rows and result.unsupported_rows[0][1] == "indirect"


# --- Robustez -----------------------------------------------------------------


def test_input_workbook_is_not_modified(synthetic_workbook, tmp_path):
    before = synthetic_workbook.read_bytes()
    _run(synthetic_workbook, tmp_path)
    assert synthetic_workbook.read_bytes() == before


def test_no_patient_values_in_outputs(synthetic_workbook, tmp_path):
    """Los CSV solo llevan coordenadas/columnas/hojas/fórmulas, nunca valores de celda."""
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    for name in (
        "derived_columns.csv",
        "formula_dependencies.csv",
        "sheet_dependencies.csv",
        "output_column_dependencies.csv",
    ):
        text = (out / name).read_text(encoding="utf-8")
        # Los valores de datos de Detail son "h_A".. y el entero de fila; no deben
        # aparecer salvo dentro del texto de una fórmula (columna example_formula).
        assert "h_A" not in text


def test_cli_main_ok(synthetic_workbook, tmp_path, capsys):
    code = ad.main(
        [
            str(synthetic_workbook),
            "--output",
            str(tmp_path / "inv"),
            "--flow-doc",
            str(tmp_path / "FLOW.md"),
            "--detail-sheet",
            "Detail",
        ]
    )
    assert code == 0
    stdout = capsys.readouterr().out
    assert "columnas derivadas detectadas: 3" in stdout


def test_cli_missing_file(tmp_path, capsys):
    code = ad.main([str(tmp_path / "nope.xlsx"), "--output", str(tmp_path / "inv")])
    assert code == 2
    assert "no existe" in capsys.readouterr().err.lower()


def test_missing_detail_sheet_is_a_warning_not_an_error(synthetic_workbook, tmp_path, capsys):
    out, _doc, result = _run(synthetic_workbook, tmp_path, detail_sheet="Nonexistent")
    assert result.detail_sheet_present is False
    assert result.derived_column_count == 0

    derived = (out / "derived_columns.csv").read_text(encoding="utf-8").splitlines()
    assert len(derived) == 1  # solo cabecera
    assert _rows(out / "output_column_dependencies.csv") == []
