"""Tests del comparador de agregaciones legacy (Sprint 2.3).

Workbook sintético: hoja de detalle con columnas en las posiciones reales
(D, G, H, K, M, O, AA, P, J, AB) + hojas ``REMASEP 01`` / ``B2 ANEXO`` con
fórmulas COUNTIF/COUNTIFS y valores cacheados parcheados en el XML.
No se usa data/local.
"""

from __future__ import annotations

import csv
import json
import xml.etree.ElementTree as ET
import zipfile

import compare_legacy_aggregations as agg
import pytest
from openpyxl import Workbook

_DETAIL = "Atenciones - Detalles de citas"
_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}
_GENETICA = "CONSULTA MEDICA DE ESPECIALIDAD EN GENETICA CLINICA - COD.0101325"
_KINE = "ATENCION KINESIOLOGICA INTEGRAL AMBULATORIA - COD.0601105"


def _sheet_xml_map(path) -> dict[str, str]:
    ns_m = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as zf:
        wbxml = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid_target = {rel.get("Id"): rel.get("Target") for rel in rels}
    out = {}
    for sheet in wbxml.find(f"{{{ns_m}}}sheets"):
        rid = sheet.get(f"{{{ns_r}}}id")
        target = rid_target[rid].lstrip("/")
        if not target.startswith("xl/"):
            target = "xl/" + target
        out[sheet.get("name")] = target
    return out


@pytest.fixture
def build_agg_workbook(tmp_path, inject_formula_cache):
    """Construye el workbook sintético y devuelve (path, sheet_xml_map)."""

    def _build(detail_rows, sheets, *, filename="agg.xlsx"):
        wb = Workbook()
        detail = wb.active
        detail.title = _DETAIL
        for letter, header in _HDR.items():
            detail[f"{letter}1"] = header
        for i, row in enumerate(detail_rows, start=2):
            for sem, letter in _COLS.items():
                if row.get(sem) is not None:
                    detail[f"{letter}{i}"] = row[sem]
            detail[f"BA{i}"] = "x"

        for sheet_name, spec in sheets.items():
            ws = wb.create_sheet(sheet_name)
            for coord, value in spec.get("labels", {}).items():
                ws[coord] = value
            for coord, formula in spec["formulas"].items():
                ws[coord] = formula

        path = tmp_path / filename
        wb.save(path)

        xml_map = _sheet_xml_map(path)
        for sheet_name, spec in sheets.items():
            cache = {}
            for coord, value in spec.get("cache", {}).items():
                m = coord[0], int(coord[1:])
                cache[(m[1], m[0])] = value
            if cache:
                inject_formula_cache(path, cache, sheet_xml=xml_map[sheet_name])
        return path

    return _build


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2000", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "X", "TIPO_DE_CITA": "OTRA",
        "PRESTACION": "p", "ESTADO": "Atendido", "MODALIDAD": "Fonasa B",
        "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _q(col: str, crit: str) -> str:
    return f"=COUNTIF('{_DETAIL}'!${col}:${col},{crit})"


@pytest.fixture
def scenario(build_agg_workbook):
    rows = [
        _row(TIPO_DE_CITA=_GENETICA, SEXO="Hombre", FECHA_NACIMIENTO="01-01-2010"),  # AG H, AF 16
        _row(TIPO_DE_CITA=_GENETICA, SEXO="Mujer", FECHA_NACIMIENTO="01-01-2020"),  # AG M, AF 6
        _row(TIPO_DE_CITA=_GENETICA, SEXO="Hombre", FECHA_NACIMIENTO="01-01-2024"),  # AF 2
        _row(TIPO_DE_CITA="CONTROL PERIODONCIA", SEXO="Hombre", FECHA_NACIMIENTO="01-01-2000"),  # AH H, AF 26
        _row(TIPO_DE_CITA="OTRA", PRESTACION="pre CONTROL APARATO REMOVIBLE post", SEXO="Mujer",
             FECHA_NACIMIENTO="01-01-2015"),  # AJ "1Mujer"
        _row(TIPO_DE_CITA=_KINE, SUCURSAL="Santiago"),  # AC = tipo+Santiago
        _row(TIPO_DE_CITA=_KINE, SUCURSAL="Santiago"),
        {},  # fila estructuralmente vacía
    ]
    remasep01 = {
        "labels": {"A5": "CONTROL PERIODONCIA", "A6": "control odon esp"},
        "formulas": {
            "B1": _q("AG", '"*consulta médico*Hombre"'),  # -> 2
            "B2": (f"=COUNTIFS('{_DETAIL}'!$AG:$AG,\"*consulta médico*\","
                   f"'{_DETAIL}'!$AF:$AF,\">=10\",'{_DETAIL}'!$AF:$AF,\"<=20\")"),  # -> 1
            "B5": (f"=COUNTIFS('{_DETAIL}'!$O:$O,$A5,'{_DETAIL}'!$AF:$AF,\">20\")"),  # -> 1
            "B6": _q("AH", '$A6&"Hombre"'),  # -> 1
            "B7": _q("AJ", '"1Mujer"'),  # -> 1
            "B8": _q("AC", f'"{_KINE}Santiago"') + "+" + _q("AC", '"nada"').removeprefix("="),  # -> 2
            "B9": _q("AF", '"<10"'),  # age-only -> padding sensitive; activas AF<10: rows 2,3 -> 2
            "B10": _q("AG", '"*x*"') + "-A1",  # UNSUPPORTED
            "B11": _q("AG", '"*consulta médico*Hombre"'),  # cache difference (no AF)
            "B12": _q("AH", '"control odon espHombre"'),  # cache unavailable (sin cache)
            "B13": (f"=COUNTIFS('{_DETAIL}'!$AG:$AG,\"*consulta médico*\","
                    f"'{_DETAIL}'!$AF:$AF,\">=10\",'{_DETAIL}'!$AF:$AF,\"<=20\")"),  # AF-dep diff
        },
        "cache": {
            "B1": 2, "B2": 1, "B5": 1, "B6": 1, "B7": 1, "B8": 2, "B9": 2,
            "B10": 5, "B11": 99, "B13": 42,
        },
    }
    b2 = {
        "formulas": {"D1": f"=COUNTIF('{_DETAIL}'!O:O,\"{_KINE}\")"},  # -> 2
        "cache": {"D1": 2},
    }
    path = build_agg_workbook(rows, {"REMASEP 01": remasep01, "B2 ANEXO": b2})
    return agg.compare_legacy_aggregations(path), path


def _cells(result):
    return {(c.sheet, c.cell): c for c in result.cells}


def _read_rows(path):
    with open(path, encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


# --- estructura general -------------------------------------------------


def test_dataset_and_classification(scenario):
    result, _ = scenario
    assert result.physical_rows == 8
    assert result.structural_empty_rows == 1
    assert result.active_records == 7
    assert result.detail_sheet == _DETAIL

    cells = _cells(result)
    assert cells[("REMASEP 01", "B1")].classification == "DIRECT_SOURCE_AGGREGATION"
    assert cells[("REMASEP 01", "B10")].classification == "UNSUPPORTED"
    assert "tras el COUNTIF" in cells[("REMASEP 01", "B10")].unsupported_reason


@pytest.mark.parametrize(
    ("sheet", "cell", "value"),
    [
        ("REMASEP 01", "B1", 2),   # COUNTIF wildcard
        ("REMASEP 01", "B2", 1),   # COUNTIFS + banda edad (>=, <=)
        ("REMASEP 01", "B5", 1),   # COUNTIFS + cell-ref + ">"
        ("REMASEP 01", "B6", 1),   # cell-ref & "Hombre"
        ("REMASEP 01", "B7", 1),   # COUNTIF exacto sobre AJ
        ("REMASEP 01", "B8", 2),   # suma de COUNTIF
        ("REMASEP 01", "B9", 2),   # COUNTIF solo edad ("<10")
        ("B2 ANEXO", "D1", 2),     # whole-column O:O sin '$'
    ],
)
def test_python_evaluation(scenario, sheet, cell, value):
    result, _ = scenario
    assert _cells(result)[(sheet, cell)].python_value == value


# --- cache: separación de estados -----------------------------------


def test_cache_match(scenario):
    result, _ = scenario
    c = _cells(result)[("REMASEP 01", "B1")]
    assert c.cache_status == "MATCH"
    assert c.notes == ""


def test_cache_difference_without_af(scenario):
    result, _ = scenario
    c = _cells(result)[("REMASEP 01", "B11")]
    assert c.cache_status == "CACHE_DIFFERENCE"
    assert c.python_value == 2
    assert c.cached_value == 99
    assert "depends on AF" not in c.notes  # no depende de AF


def test_cache_unavailable(scenario):
    result, _ = scenario
    c = _cells(result)[("REMASEP 01", "B12")]
    assert c.cache_status == "CACHE_UNAVAILABLE"
    assert c.python_value == 1


def test_cache_difference_with_af_gets_note(scenario):
    result, _ = scenario
    c = _cells(result)[("REMASEP 01", "B13")]
    assert c.cache_status == "CACHE_DIFFERENCE"
    assert c.depends_on_age is True
    assert "Formula depends on AF" in c.notes


# --- padding sensitivity --------------------------------------------


def test_padding_sensitivity(scenario, tmp_path):
    result, _ = scenario
    cells = _cells(result)
    # con filtro de categoría -> no sensible
    assert cells[("REMASEP 01", "B2")].padding_can_affect is False
    # COUNTIF solo sobre AF -> una fila vacía (AF=0) contaría
    assert cells[("REMASEP 01", "B9")].padding_can_affect is True

    agg.write_outputs(result, tmp_path / "out")
    rows = _read_rows(tmp_path / "out" / "padding_sensitivity.csv")
    by = {r["target_cell"]: r for r in rows}
    assert by["B9"]["padding_can_affect_result"] == "True"
    assert by["B2"]["padding_can_affect_result"] == "False"


# --- artefactos + resumen -----------------------------------------


def test_summary_json(scenario, tmp_path):
    result, _ = scenario
    files = agg.write_outputs(result, tmp_path / "out")
    summary = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))

    assert summary["physical_rows"] == 8
    assert summary["structural_empty_rows"] == 1
    assert summary["active_records"] == 7
    assert summary["unsupported_formula_cells"] == 1
    assert summary["formula_support_status"] == "PARTIAL"
    assert summary["cache_consistency_status"] == "DIFFERENCES"
    assert summary["padding_sensitive"] is True
    assert summary["padding_sensitive_cells"] == 1
    assert summary["cache_differences"] >= 2
    assert summary["cache_differences_depending_on_AF"] == 1
    assert "COUNTIF" in summary["observed_formula_families"]
    assert {f.name for f in files} >= {
        "summary.json", "aggregation_equivalence.csv", "sheet_summary.csv",
        "direct_aggregation_inventory.csv", "unsupported_formulas.csv",
        "padding_sensitivity.csv", "legacy_metrics_long.csv", "README.md",
    }


def test_sheet_summary_csv(scenario, tmp_path):
    result, _ = scenario
    agg.write_outputs(result, tmp_path / "out")
    rows = {r["sheet"]: r for r in _read_rows(tmp_path / "out" / "sheet_summary.csv")}
    assert rows["B2 ANEXO"]["supported_cells"] == "1"
    assert rows["B2 ANEXO"]["cache_matches"] == "1"
    assert int(rows["REMASEP 01"]["unsupported_cells"]) == 1
    assert int(rows["REMASEP 01"]["cache_differences"]) >= 2


def test_metric_ids_are_stable_and_technical(scenario, tmp_path):
    result, _ = scenario
    agg.write_outputs(result, tmp_path / "out")
    rows = _read_rows(tmp_path / "out" / "legacy_metrics_long.csv")
    by = {r["metric_id"]: r for r in rows}
    assert "LEGACY::REMASEP_01::B1" in by
    assert by["LEGACY::REMASEP_01::B1"]["value"] == "2"
    assert by["LEGACY::REMASEP_01::B1"]["generator_sheet"] == "REMASEP 01"
    # las celdas no soportadas NO aparecen como métricas
    assert "LEGACY::REMASEP_01::B10" not in by


# --- privacidad ---------------------------------------------------


def test_artifacts_have_no_individual_patient_data(build_agg_workbook, tmp_path):
    secret = "SECRET-PATIENT-XYZ"
    rows = [
        _row(TIPO_DE_CITA=_GENETICA, SEXO="Hombre", ESTADO=secret,
             PRESTACION_REALIZADA=secret, MODALIDAD=secret),
        _row(TIPO_DE_CITA=_GENETICA, SEXO="Mujer", ESTADO=secret),
    ]
    sheets = {
        "REMASEP 01": {
            "formulas": {"B1": _q("AG", '"*consulta médico*"')},
            "cache": {"B1": 2},
        }
    }
    path = build_agg_workbook(rows, sheets)
    result = agg.compare_legacy_aggregations(path)
    files = agg.write_outputs(result, tmp_path / "out")
    for f in files:
        assert secret not in f.read_text(encoding="utf-8"), f.name


# --- errores de entrada ----------------------------------------


def test_missing_file(tmp_path):
    with pytest.raises(agg.AggregationCompareError):
        agg.compare_legacy_aggregations(tmp_path / "nope.xlsx")


def test_cli_main(scenario, tmp_path, capsys):
    _result, path = scenario
    code = agg.main([str(path), "--output", str(tmp_path / "out")])
    assert code == 0
    out = capsys.readouterr().out
    assert "formula_support_status: PARTIAL" in out
    assert "cache_consistency_status: DIFFERENCES" in out
