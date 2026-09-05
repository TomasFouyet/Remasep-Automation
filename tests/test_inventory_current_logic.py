"""Tests del inventario de lógica actual (Sprint 1.3).

Workbooks sintéticos únicamente. Los valores "raw" llevan el centinela
``SECRET_VALUE`` para verificar que nunca se exportan.
"""

import csv

import inventory_current_logic as icl
import pytest
from openpyxl import Workbook

SECRET = "SECRET_VALUE"


@pytest.fixture
def synthetic_workbook(tmp_path):
    wb = Workbook()
    detail = wb.active
    detail.title = "Detail"

    headers = {
        "A": "DIA CITA",
        "B": "FECHA NACIMIENTO",
        "C": "SEXO",
        "D": "TIPO DE CITA",
        "E": "SUCURSAL",
        "F": "PRESTACION",
        "G": "ESTADO",
        "H": "MODALIDAD",
        "I": "PRESTACION REALIZADA",
    }
    for col, text in headers.items():
        detail[f"{col}1"] = text
    detail["J1"] = "J concat"
    detail["K1"] = "K edad"
    detail["L1"] = "L map"
    detail["M1"] = "M pattern"
    detail["N1"] = "N pattern2"

    for row in (2, 3, 4, 5):
        for col in "ABCDEFGHI":
            detail[f"{col}{row}"] = SECRET
        detail[f"J{row}"] = f"=D{row}&E{row}"
        detail[f"K{row}"] = f'=IF(A{row}>B{row},0,DATEDIF(B{row},A{row},"Y"))'
        detail[f"L{row}"] = f'=IFS(D{row}="X","mapx",D{row}="Y","mapy")&C{row}'
        detail[f"M{row}"] = f'=(COUNTIF(F{row},"*ALPHA*")+COUNTIF(F{row},"*BETA*"))&C{row}'
        detail[f"N{row}"] = f'=(COUNTIF(F{row},"PREFIX*"))&C{row}'

    out1 = wb.create_sheet("Out1")
    out1["A1"] = '=COUNTIF(\'Detail\'!J:J,"x")'
    out1["A2"] = '=COUNTIF(\'Detail\'!D:D,"y")'

    out2 = wb.create_sheet("Out2")
    out2["B1"] = "=SUM('Detail'!K2:K9)"
    out2["B2"] = '=COUNTIF(\'Detail\'!L:L,"z")'
    out2["B3"] = '=COUNTIF(\'Detail\'!M:M,"w")'
    out2["B4"] = '=COUNTIF(\'Detail\'!N:N,"q")'

    path = tmp_path / "synthetic.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def workbook_with_issues(tmp_path):
    wb = Workbook()
    detail = wb.active
    detail.title = "Detail"
    detail["A1"] = "H1"
    detail["B1"] = "H2"
    detail["C1"] = "DER"
    for row in (2, 3, 4, 5):
        detail[f"A{row}"] = SECRET
        detail[f"B{row}"] = SECRET
    detail["C2"] = "=A2&B2"
    detail["C3"] = "=A3&B3"
    detail["C4"] = "=A4+B4"  # patrón normalizado distinto -> pattern_count = 2
    # C5 vacío -> hueco (la fila 5 tiene datos en A5/B5)

    out1 = wb.create_sheet("Out1")
    out1["A1"] = '=COUNTIF(\'Detail\'!C:C,"x")'

    path = tmp_path / "issues.xlsx"
    wb.save(path)
    return path


def _rows(path):
    with open(path, newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _run(workbook, tmp_path, **kwargs):
    out = tmp_path / "cl"
    doc = tmp_path / "docs" / "CURRENT_LOGIC.md"
    kwargs.setdefault("detail_sheet", "Detail")
    result = icl.analyze_current_logic(workbook, out, doc_path=doc, **kwargs)
    return out, doc, result


# --- Parte A: source schema ---------------------------------------------


def test_source_schema(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = {r["column"]: r for r in _rows(out / "source_schema.csv")}

    assert set(rows) == set("ABCDEFGHIJKLMN")
    assert rows["A"]["header"] == "DIA CITA"
    assert rows["A"]["kind"] == "raw"
    assert rows["J"]["kind"] == "derived"
    assert rows["J"]["has_formulas"] == "True"
    assert rows["J"]["formula_count"] == "4"

    # D lo referencia Out1 directamente.
    assert rows["D"]["directly_used_by_outputs"] == "True"
    assert rows["J"]["directly_used_by_outputs"] == "True"
    assert rows["A"]["directly_used_by_outputs"] == "False"

    # Transitivo: A..F sí; ESTADO/MODALIDAD/PRESTACION REALIZADA no.
    for col in "ABCDEF":
        assert rows[col]["transitively_used_by_outputs"] == "True", col
    for col in ("G", "H", "I"):
        assert rows[col]["transitively_used_by_outputs"] == "False", col


def test_transitive_raw_columns_reported(synthetic_workbook, tmp_path):
    _out, _doc, result = _run(synthetic_workbook, tmp_path)
    assert [letter for letter, _h in result.transitive_raw_columns] == list("ABCDEF")
    assert dict(result.transitive_raw_columns)["A"] == "DIA CITA"


# --- Parte B: transformaciones ----------------------------------------


def test_derived_transformations(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = {r["target_column"]: r for r in _rows(out / "derived_transformations.csv")}

    assert set(rows) == set("JKLMN")
    assert rows["J"]["transformation_type"] == "CONCAT"
    assert rows["J"]["source_columns"] == "D|E"
    assert rows["J"]["source_headers"] == "TIPO DE CITA|SUCURSAL"
    assert rows["J"]["pattern_count"] == "1"

    assert rows["K"]["transformation_type"] == "AGE_DATEDIF"
    assert rows["K"]["source_columns"] == "A|B"

    assert rows["L"]["transformation_type"] == "EXACT_MAP"
    assert rows["L"]["source_columns"] == "C|D"

    assert rows["M"]["transformation_type"] == "PATTERN_FLAG"
    assert rows["M"]["source_columns"] == "C|F"
    assert rows["N"]["transformation_type"] == "PATTERN_FLAG"


# --- Parte C: reglas candidatas --------------------------------------


def test_rule_candidates(synthetic_workbook, tmp_path):
    out, _doc, result = _run(synthetic_workbook, tmp_path)
    rows = _rows(out / "derived_rule_candidates.csv")

    assert {r["target_column"] for r in rows} == {"L", "M", "N"}
    assert len(rows) == 5  # L:2  M:2  N:1

    lrules = [r for r in rows if r["target_column"] == "L"]
    assert [r["rule_index"] for r in lrules] == ["1", "2"]
    assert lrules[0]["source_column"] == "D"
    assert lrules[0]["source_header"] == "TIPO DE CITA"
    assert lrules[0]["operator"] == "equals"
    assert lrules[0]["criterion"] == "X"
    assert lrules[0]["output_literal"] == "mapx"
    assert lrules[0]["appended_columns"] == "C"
    assert lrules[0]["original_formula"].startswith("=IFS(")

    mrules = [r for r in rows if r["target_column"] == "M"]
    assert mrules[0]["source_column"] == "F"
    assert mrules[0]["operator"] == "contains"
    assert mrules[0]["criterion"] == "ALPHA"  # sin los '*' exteriores
    assert mrules[0]["output_literal"] == ""
    assert mrules[0]["appended_columns"] == "C"

    nrule = next(r for r in rows if r["target_column"] == "N")
    assert nrule["operator"] == "wildcard_match"
    assert nrule["criterion"] == "PREFIX*"  # patrón no '*x*' -> se conserva

    assert len(result.rule_candidates) == 5


# --- Parte D: dependencias transitivas ------------------------------


def test_transitive_output_dependencies(synthetic_workbook, tmp_path):
    out, _doc, _result = _run(synthetic_workbook, tmp_path)
    rows = {(r["target_sheet"], r["raw_source_column"]): r for r in _rows(out / "transitive_output_dependencies.csv")}

    assert set(rows) == {
        ("Out1", "D"),
        ("Out1", "E"),
        ("Out2", "A"),
        ("Out2", "B"),
        ("Out2", "C"),
        ("Out2", "D"),
        ("Out2", "F"),
    }
    assert rows[("Out1", "D")]["dependency_paths"] == "D|J->D"  # directa + vía J
    assert rows[("Out1", "E")]["dependency_paths"] == "J->E"
    assert rows[("Out2", "C")]["dependency_paths"] == "L->C|M->C|N->C"
    assert rows[("Out2", "F")]["dependency_paths"] == "M->F|N->F"
    assert rows[("Out2", "A")]["raw_source_header"] == "DIA CITA"
    # Ninguna columna derivada aparece como raw.
    assert not {k[1] for k in rows} & set("JKLMN")


# --- Parte E: documento --------------------------------------------


def test_doc_sections_and_observations(synthetic_workbook, tmp_path):
    _out, doc, _result = _run(synthetic_workbook, tmp_path)
    text = doc.read_text(encoding="utf-8")

    assert "## Observaciones que requieren validación funcional" in text
    assert "Conjunto total de columnas raw utilizadas transitivamente" in text
    assert "**ESTADO**" in text and "**MODALIDAD**" in text
    assert "**PRESTACIÓN REALIZADA**" in text
    # G/H/I señaladas como que no participan.
    assert "no participa" in text
    # El conjunto raw aparece.
    for letter in "ABCDEF":
        assert f"`{letter}`" in text


# --- Invariantes ----------------------------------------------------


def test_clean_workbook_has_no_invariant_warnings(synthetic_workbook, tmp_path):
    out, _doc, result = _run(synthetic_workbook, tmp_path)
    assert result.invariant_warnings == []
    invariants = (out / "invariants.csv").read_text(encoding="utf-8").splitlines()
    assert len(invariants) == 1  # solo cabecera


def test_issues_workbook_reports_but_does_not_fail(workbook_with_issues, tmp_path):
    out, _doc, result = _run(workbook_with_issues, tmp_path)
    checks = {w.check for w in result.invariant_warnings}
    assert "pattern_count" in checks
    assert "gaps" in checks
    assert "formula_count_vs_data_rows" in checks

    rows = _rows(out / "invariants.csv")
    assert all(r["severity"] == "warning" for r in rows)
    gap_row = next(r for r in rows if r["check"] == "gaps")
    assert gap_row["target_column"] == "C"
    assert "5" in gap_row["detail"]


def test_cli_main_ok_and_exit_zero_with_issues(workbook_with_issues, tmp_path, capsys):
    code = icl.main(
        [
            str(workbook_with_issues),
            "--output",
            str(tmp_path / "cl"),
            "--doc",
            str(tmp_path / "CURRENT_LOGIC.md"),
            "--detail-sheet",
            "Detail",
        ]
    )
    assert code == 0  # análisis, no procesamiento: no falla por invariantes
    assert "advertencias de invariantes: 3" in capsys.readouterr().out


# --- Privacidad / robustez ----------------------------------------


def test_no_patient_values_in_any_output(synthetic_workbook, tmp_path):
    out, doc, _result = _run(synthetic_workbook, tmp_path)
    targets = list(out.glob("*.csv")) + [out / "README.md", doc]
    for path in targets:
        assert SECRET not in path.read_text(encoding="utf-8"), path.name


def test_input_workbook_not_modified(synthetic_workbook, tmp_path):
    before = synthetic_workbook.read_bytes()
    _run(synthetic_workbook, tmp_path)
    assert synthetic_workbook.read_bytes() == before


def test_missing_file(tmp_path, capsys):
    code = icl.main([str(tmp_path / "nope.xlsx"), "--output", str(tmp_path / "cl")])
    assert code == 2
    assert "no existe" in capsys.readouterr().err.lower()


def test_missing_detail_sheet_is_warning(synthetic_workbook, tmp_path, capsys):
    out, _doc, result = _run(synthetic_workbook, tmp_path, detail_sheet="Nope")
    assert result.detail_sheet_present is False
    assert _rows(out / "source_schema.csv") == []
    assert result.transitive_raw_columns == []
