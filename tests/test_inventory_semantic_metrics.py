"""Tests del inventario semántico (Sprint 3.1), con workbook sintético.

Hoja de detalle en posiciones reales + `REMASEP_OD` con layout (sección /
encabezados de columna combinados / rótulos de fila) y métricas COUNTIFS, un
`SUM` y un `IF`. Sin data/local.
"""

from __future__ import annotations

import csv
import json

import inventory_semantic_metrics as inv
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

_DETAIL = "Atenciones - Detalles de citas"
_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}
_BOLD = Font(bold=True)


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2003", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "COMPOSITE", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _countifs(sex_pat: str, lo: int, hi: int) -> str:
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*COMPOSITE*\"&\"*{sex_pat}*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


def _countif(sex_pat: str) -> str:
    return f"=COUNTIF('{_DETAIL}'!$AE:$AE,\"*COMPOSITE*\"&\"*{sex_pat}*\")"


@pytest.fixture
def scenario(tmp_path):
    wb = Workbook()
    detail = wb.active
    detail.title = _DETAIL
    for letter, header in _HDR.items():
        detail[f"{letter}1"] = header
    rows = [
        _row(FECHA_NACIMIENTO="01-01-2003"),                 # AF 23, Hombre
        _row(FECHA_NACIMIENTO="01-01-2003", SEXO="Mujer"),   # AF 23, Mujer
        _row(FECHA_NACIMIENTO="01-01-1999"),                 # AF 27, Hombre
        _row(FECHA_NACIMIENTO="01-01-1999", SEXO="Mujer", ESTADO="SECRET-XYZ"),  # AF 27
        {},                                                   # fila vacía
    ]
    for i, r in enumerate(rows, start=2):
        for sem, letter in _COLS.items():
            if r.get(sem) is not None:
                detail[f"{letter}{i}"] = r[sem]
        detail[f"BA{i}"] = "x"

    od = wb.create_sheet("REMASEP_OD")
    od["B1"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["B1"].font = _BOLD
    od.merge_cells("C2:D2")
    od["C2"] = "20 A 24 AÑOS"
    od.merge_cells("E2:F2")
    od["E2"] = "25 A 29 AÑOS"
    od.merge_cells("I2:J2")
    od["I2"] = "AMBOS SEXOS"
    for col in ("C", "E"):
        od[f"{col}3"] = "Hombres"
    for col in ("D", "F"):
        od[f"{col}3"] = "Mujeres"
    # I/J: sólo el encabezado combinado "AMBOS SEXOS" (sin subnivel) -> I6 y J6
    # comparten sección + fila + columna -> misma semantic_signature.
    od["A6"] = "COD-1"
    od["B6"] = "Prestación uno"
    od["C6"] = _countifs("Hombre", 20, 24)   # label 20-24 / formula 20-24 -> CONSISTENT
    od["D6"] = _countifs("Mujer", 20, 24)
    od["E6"] = _countifs("Hombre", 25, 29)   # label 25-29 / formula 25-29 -> CONSISTENT
    od["F6"] = _countifs("Mujer", 25, 29)
    od["I6"] = _countif("Hombre")            # duplicate signature con J6
    od["J6"] = _countif("Hombre")
    od["A7"] = "COD-2"
    od["B7"] = "Prestación dos"
    od["C7"] = _countifs("Hombre", 25, 29)   # label 20-24 (col C) / formula 25-29 -> CONFLICT
    od["K6"] = "=SUM(C6+D6+E6+F6)"           # DOWNSTREAM_TOTAL
    od["L6"] = "=IF(C6<D6,1,0)"              # VALIDATION

    path = tmp_path / "semantic.xlsx"
    wb.save(path)
    return inv.build_inventory(path), path


def _metrics(result):
    return {(m.sheet, m.cell): m for m in result.metrics}


# --- estructura general ------------------------------------------------


def test_source_is_medinet_and_counts(scenario):
    result, _ = scenario
    metrics = _metrics(result)
    assert all(m.source == "MEDINET" for m in result.metrics)
    kinds = {(k): 0 for k in ("BASE_AGGREGATION", "DERIVED_AGGREGATION", "DOWNSTREAM_TOTAL")}
    for m in result.metrics:
        kinds[m.kind] += 1
    assert kinds["BASE_AGGREGATION"] == 7   # C6 D6 E6 F6 I6 J6 C7
    assert kinds["DOWNSTREAM_TOTAL"] == 1   # K6
    assert ("REMASEP_OD", "L6") not in metrics  # IF va a validations


def test_validation_is_separate(scenario):
    result, _ = scenario
    ids = {v["cell"] for v in result.validations}
    assert ids == {"L6"}
    v = next(v for v in result.validations if v["cell"] == "L6")
    assert v["row_context"] == "COD-1 :: Prestación uno"
    assert v["section_context"] == "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"


# --- contexto / raw vs normalized -------------------------------


def test_row_column_section_context(scenario):
    result, _ = scenario
    c6 = _metrics(result)[("REMASEP_OD", "C6")]
    assert c6.context_status == inv.CONTEXT_COMPLETE
    assert c6.section_labels_raw == ("SECCIÓN A: PRUEBAS ODONTOLÓGICAS",)
    assert c6.row_labels_raw == ("COD-1", "Prestación uno")
    assert c6.column_labels_raw == ("20 A 24 AÑOS", "Hombres")


def test_semantic_signature_raw_preserved_normalized_separate(scenario):
    result, _ = scenario
    c6 = _metrics(result)[("REMASEP_OD", "C6")]
    assert "AÑOS" in c6.column_labels_raw[0]  # raw con tilde
    assert c6.semantic_signature == (
        "REMASEP_OD :: SECCION A: PRUEBAS ODONTOLOGICAS :: COD-1 :: "
        "PRESTACION UNO :: 20 A 24 ANOS :: HOMBRES"
    )


# --- evidencia de fórmula ------------------------------------------


def test_formula_evidence_age_and_text(scenario):
    result, _ = scenario
    c6 = _metrics(result)[("REMASEP_OD", "C6")]
    assert c6.age_lower_bound == 20
    assert c6.age_upper_bound == 24
    assert c6.formula_evidence_status == "AGE_AND_TEXT"
    assert any("COMPOSITE" in t for t in c6.text_criteria)
    assert "AE" in c6.referenced_detail_columns and "AF" in c6.referenced_detail_columns


def test_downstream_total_has_no_criteria(scenario):
    result, _ = scenario
    k6 = _metrics(result)[("REMASEP_OD", "K6")]
    assert k6.kind == "DOWNSTREAM_TOTAL"
    assert k6.formula_evidence_status == "AGGREGATE_ONLY"
    assert k6.text_criteria == ()


# --- consistencia rótulo / fórmula --------------------------------


def test_label_formula_consistent(scenario):
    result, _ = scenario
    assert _metrics(result)[("REMASEP_OD", "C6")].consistency_status == "CONSISTENT"
    assert _metrics(result)[("REMASEP_OD", "E6")].consistency_status == "CONSISTENT"


def test_label_formula_conflict_is_recorded_not_fixed(scenario):
    result, _ = scenario
    c7 = _metrics(result)[("REMASEP_OD", "C7")]
    assert c7.column_labels_raw[0] == "20 A 24 AÑOS"
    assert (c7.age_lower_bound, c7.age_upper_bound) == (25, 29)
    assert c7.consistency_status == "CONFLICT"


# --- duplicados ----------------------------------------------------


def test_duplicate_semantic_signatures_detected_not_resolved(scenario, tmp_path):
    result, _ = scenario
    inv.write_outputs(result, tmp_path / "out")
    with open(tmp_path / "out" / "duplicate_semantic_signatures.csv", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 1
    assert rows[0]["metric_count"] == "2"
    assert set(rows[0]["cells"].split("|")) == {"REMASEP_OD!I6", "REMASEP_OD!J6"}
    assert rows[0]["status"] == "UNRESOLVED"


# --- artefactos / summary / privacidad ----------------------------


def test_summary_json(scenario, tmp_path):
    result, _ = scenario
    inv.write_outputs(result, tmp_path / "out")
    summary = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))
    assert summary["total_metric_candidates"] == 8
    assert summary["metrics_by_kind"]["BASE_AGGREGATION"] == 7
    assert summary["duplicate_semantic_signatures"] == 1
    assert summary["label_formula_conflicts"] == 1
    assert summary["label_formula_consistent"] >= 2
    assert summary["validation_total"] == 1
    assert summary["max_column_hierarchy_depth"] == 2
    assert summary["semantic_layer_status"] == "PRELIMINARY_NOT_VALIDATED"


def test_layout_csv_written_per_sheet(scenario, tmp_path):
    result, _ = scenario
    files = inv.write_outputs(result, tmp_path / "out")
    names = {f.name for f in files}
    assert "layout_REMASEP_OD.csv" in names
    assert "semantic_metric_candidates.csv" in names


def test_artifacts_have_no_individual_patient_data(scenario, tmp_path):
    result, _ = scenario
    files = inv.write_outputs(result, tmp_path / "out")
    for f in files:
        assert "SECRET-XYZ" not in f.read_text(encoding="utf-8"), f.name


def test_metric_id_is_stable(scenario):
    result, _ = scenario
    assert _metrics(result)[("REMASEP_OD", "C6")].metric_id == "LEGACY::REMASEP_OD::C6"


def test_cli_main(scenario, tmp_path, capsys):
    _result, path = scenario
    code = inv.main([str(path), "--output", str(tmp_path / "out")])
    assert code == 0
    out = capsys.readouterr().out
    assert "métricas candidatas: 8" in out
