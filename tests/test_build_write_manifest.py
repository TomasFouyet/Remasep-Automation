"""Tests CLI del write manifest (Sprint 3.6). Generador + plantilla sintéticos."""

from __future__ import annotations

import csv
import hashlib
import json

import build_write_manifest as bwm
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, Protection

_DETAIL = "Atenciones - Detalles de citas"
_BOLD = Font(bold=True)
_UNLOCKED = Protection(locked=False)
_SECRET = "PACIENTE-SECRETO-XYZ"

_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}


def _drow(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2003", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "VIDRIO IONÓMERO", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _countifs(sex_pat, lo, hi):
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*VIDRIO IONÓMERO*\"&\"*{sex_pat}*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


def _build_generator(path):
    wb = Workbook()
    detail = wb.active
    detail.title = _DETAIL
    for letter, header in _HDR.items():
        detail[f"{letter}1"] = header
    rows = [_drow(), _drow(SEXO="Mujer"),
            _drow(FECHA_NACIMIENTO="01-01-1999", ESTADO=_SECRET), {}]
    for i, r in enumerate(rows, start=2):
        for sem, letter in _COLS.items():
            if r.get(sem) is not None:
                detail[f"{letter}{i}"] = r[sem]
        detail[f"BA{i}"] = "x"

    od = wb.create_sheet("REMASEP_OD")
    od["B1"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["B1"].font = _BOLD
    od.merge_cells("C2:D2")
    od["C2"] = "20 A 24 AÑOS"
    od["E2"] = "TOTAL"
    od["C3"] = "Hombres"
    od["D3"] = "Mujeres"
    od["A6"] = "5010009 - VIDRIO IONÓMERO"
    od["B6"] = "Obturación de vidrio ionómero"
    od["C6"] = _countifs("Hombre", 20, 24)    # AUTO_READY -> EXACT
    od["D6"] = _countifs("Hombre", 20, 24)    # CONFLICT source -> excluido
    od["A7"] = "5010010 - OTRO INSUMO"
    od["C7"] = _countifs("Mujer", 20, 24)     # CONFLICT source -> excluido

    b2 = wb.create_sheet("B2 ANEXO")
    b2["B3"] = "A. KINESIOLOGÍA"
    b2["B3"].font = _BOLD
    b2["B5"] = "0601105"
    b2["C5"] = "Atención Kinesiológica Integral Ambulatoria"
    b2["D5"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"x\")"
    wb.save(path)


def _build_template(path):
    wb = Workbook()
    od = wb.active
    od.title = "REMASEP_OD"
    od.protection.sheet = True
    od["A20"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["A20"].font = _BOLD
    od.merge_cells("D21:E21")
    od["D21"] = "20 A 24 AÑOS"
    od["D22"] = "Hombres"
    od["E22"] = "Mujeres"
    od["A25"] = "5010009 - VIDRIO IONÓMERO"
    od["B25"] = "Obturación de vidrio ionómero"
    od["D25"].protection = _UNLOCKED      # writable input -> WRITE_READY candidate
    od["F25"] = "=D25"                    # formula target
    od["G21"] = "TOTAL"
    od["G22"] = "Ambos sexo"

    b1 = wb.create_sheet("REMASEP B1")
    b1.protection.sheet = True
    b1["A1"] = "SECCIÓN A: INTERVENCIONES QUIRÚRGICAS"
    b1["A1"].font = _BOLD
    b1.merge_cells("C2:D2")
    b1["C2"] = "POR GRUPO DE EDAD"
    b1["C3"] = "Menor de 10 años"
    b1["D3"] = "Hombres"
    b1["A5"] = "Electivas"
    b1["B5"] = "Mayor Ambulatorias"
    b1["C5"].protection = _UNLOCKED       # EGRESOS region -> should never be WRITE_READY

    b2 = wb.create_sheet("B2 ANEXO")
    b2.protection.sheet = True
    b2["A1"] = "CÓDIGOS"
    b2["B1"] = "GLOSA"
    b2["C1"] = "TOTAL"
    b2["A2"] = "A. KINESIOLOGÍA"
    b2["A2"].font = _BOLD
    b2["A3"] = "0601105"
    b2["B3"] = "Atención Kinesiológica Integral Ambulatoria"
    b2["C3"].protection = _UNLOCKED       # writable, matched by procedure code
    wb.save(path)


@pytest.fixture
def pair(tmp_path):
    gen, tpl = tmp_path / "generator.xlsx", tmp_path / "template.xlsx"
    _build_generator(gen)
    _build_template(tpl)
    return gen, tpl


@pytest.fixture
def outcome(pair, tmp_path):
    gen, tpl = pair
    analysis = bwm.build_write_analysis(str(gen), str(tpl))
    files = bwm.write_outputs(analysis, tmp_path / "out")
    return analysis, files, tmp_path / "out"


# ---------------------------------------------------------------------------


def test_files_written(outcome):
    _a, files, _out = outcome
    names = {f.name for f in files}
    assert names == {
        "write_readiness.csv", "write_manifest_ready.csv", "write_review_queue.csv",
        "write_review_clusters.csv", "write_blocked.csv", "write_collisions.csv",
        "write_evidence.csv", "coverage_by_form.csv", "coverage_by_match_status.csv",
        "template_fingerprint.json", "summary.json",
        "write_mapping_manual_review_sample.csv", "README.md",
    }


def test_source_files_unchanged(pair, tmp_path):
    gen, tpl = pair
    before = (hashlib.sha256(gen.read_bytes()).hexdigest(),
              hashlib.sha256(tpl.read_bytes()).hexdigest())
    bwm.build_write_analysis(str(gen), str(tpl))
    after = (hashlib.sha256(gen.read_bytes()).hexdigest(),
             hashlib.sha256(tpl.read_bytes()).hexdigest())
    assert before == after


def test_exact_medinet_input_is_write_ready(outcome):
    _a, _f, out = outcome
    with open(out / "write_readiness.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    od = next(r for r in rows if r["source_metric_id"].endswith("REMASEP_OD::C6"))
    assert od["match_status"] == "EXACT_SEMANTIC_MATCH"
    assert od["write_status"] == "WRITE_READY"
    assert od["target_locked"] == "unlocked"
    assert od["target_cell"] != "C6"  # identidad semántica, no coordenada


def test_manifest_only_write_ready_no_pii_no_values(outcome):
    _a, _f, out = outcome
    text = (out / "write_manifest_ready.csv").read_text(encoding="utf-8")
    assert _SECRET not in text
    with open(out / "write_manifest_ready.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows  # al menos una instrucción
    for r in rows:
        assert r["instruction_id"].startswith("wi:")
        assert "value" not in {k.lower() for k in r}
    # cada instruction_id del manifest debe ser WRITE_READY en readiness
    with open(out / "write_readiness.csv", encoding="utf-8", newline="") as fh:
        ready_ids = {rr["source_metric_id"] for rr in csv.DictReader(fh)
                     if rr["write_status"] == "WRITE_READY"}
    assert {r["source_metric_id"] for r in rows} <= ready_ids


def test_conflict_sources_are_blocked_not_in_manifest(outcome):
    _a, _f, out = outcome
    with open(out / "write_readiness.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    for coord in ("D6", "C7"):
        r = next(x for x in rows if x["source_metric_id"].endswith(coord))
        assert r["write_status"] in ("WRITE_BLOCKED", "NOT_WRITABLE")
    with open(out / "write_manifest_ready.csv", encoding="utf-8", newline="") as fh:
        mids = {r["source_metric_id"] for r in csv.DictReader(fh)}
    assert not any(m.endswith(("D6", "C7")) for m in mids)


def test_formula_target_not_in_manifest(outcome):
    _a, _f, out = outcome
    # F25 es fórmula; ninguna instrucción debe apuntar a una celda de fórmula
    with open(out / "write_manifest_ready.csv", encoding="utf-8", newline="") as fh:
        cells = {(r["target_sheet"], r["target_cell"]) for r in csv.DictReader(fh)}
    assert ("REMASEP_OD", "F25") not in cells


def test_template_fingerprint_json(outcome):
    _a, _f, out = outcome
    fp = json.loads((out / "template_fingerprint.json").read_text(encoding="utf-8"))
    assert fp["structural_template_fingerprint_id"].startswith("stf:")
    assert fp["file_sha256"] != fp["structural_sha256"]
    assert "REMASEP_OD" in fp["per_sheet"]


def test_sample_has_target_locked_and_expected_value_type(outcome):
    _a, _f, out = outcome
    with open(out / "write_mapping_manual_review_sample.csv", encoding="utf-8",
              newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert rows
    header = rows[0].keys()
    # columnas nuevas presentes, columnas previas conservadas
    for col in ("source_metric_id", "source_context", "match_status", "target_sheet",
                "target_cell", "target_context", "target_kind", "target_alignment_role",
                "target_source_expectation", "write_status", "write_reason",
                "evidence_summary", "target_locked", "expected_value_type"):
        assert col in header
    for r in rows:
        assert r["target_locked"] in ("locked", "unlocked", "unknown")
        if r["write_status"] == "WRITE_READY":
            assert r["expected_value_type"] == "INTEGER_COUNT"
            assert r["target_locked"] == "unlocked"


def test_summary_formula_metrics_are_explicit(outcome):
    _a, _f, out = outcome
    s = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    # métrica ambigua renombrada / eliminada
    assert "formula_targets_excluded" not in s
    assert "template_formula_targets" in s
    assert "source_rows_blocked_by_formula_target" in s
    # celdas físicas de fórmula >= source rows bloqueados por ellas
    assert s["template_formula_targets"] >= s["source_rows_blocked_by_formula_target"]


def test_summary_period_and_estado_contract(outcome):
    _a, _f, out = outcome
    s = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    assert s["write_status"] == "PRELIMINARY_NOT_VALIDATED"
    assert s["estado_filter_status"] == "PENDING_FUNCTIONAL_CONFIRMATION"
    assert s["zero_write_policy"] == "UNRESOLVED"
    assert "processing_scope_records" in s["period_contract"]
    assert s["source_hash_verified_unchanged"] is True
    total = sum(s["write_status_counts"].values())
    assert total == s["total_source_metrics"]


def test_no_pii_in_any_artifact(outcome):
    _a, files, _out = outcome
    for f in files:
        assert _SECRET not in f.read_text(encoding="utf-8"), f.name


def test_cli_main_smoke(pair, tmp_path, capsys):
    gen, tpl = pair
    code = bwm.main([str(gen), str(tpl), "--output", str(tmp_path / "cli")])
    assert code == 0
    out = capsys.readouterr().out
    assert "write_status:" in out
    assert "structural_template_fingerprint_id:" in out
    assert "estado_filter_status: PENDING_FUNCTIONAL_CONFIRMATION" in out


def test_cli_main_bad_file(tmp_path):
    assert bwm.main([str(tmp_path / "no.xlsx"), str(tmp_path / "no2.xlsx")]) == 2
