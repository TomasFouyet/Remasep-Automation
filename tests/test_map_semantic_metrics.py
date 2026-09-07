"""Tests del mapeo semántico end-to-end (Sprint 3.2), con workbook sintético.

Hoja de detalle en posiciones reales + `REMASEP_OD` (con layout) + `B2 ANEXO`
(lista plana con códigos). Sin data/local.
"""

from __future__ import annotations

import csv
import json

import map_semantic_metrics as mp
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from remasep.services import semantic_mapping as sm

_DETAIL = "Atenciones - Detalles de citas"
_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}
_BOLD = Font(bold=True)


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2003", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "VIDRIO IONÓMERO", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _countifs(sex_pat: str, lo: int, hi: int) -> str:
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*VIDRIO IONÓMERO*\"&\"*{sex_pat}*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


@pytest.fixture
def scenario(tmp_path):
    wb = Workbook()
    detail = wb.active
    detail.title = _DETAIL
    for letter, header in _HDR.items():
        detail[f"{letter}1"] = header
    rows = [
        _row(FECHA_NACIMIENTO="01-01-2003"),                                  # AF 23 Hombre
        _row(FECHA_NACIMIENTO="01-01-2003", SEXO="Mujer"),                    # AF 23 Mujer
        _row(FECHA_NACIMIENTO="01-01-1999", SEXO="Mujer", ESTADO="SECRET-XYZ"),
        {},
    ]
    for i, r in enumerate(rows, start=2):
        for sem, letter in _COLS.items():
            if r.get(sem) is not None:
                detail[f"{letter}{i}"] = r[sem]
        detail[f"BA{i}"] = "x"

    od = wb.create_sheet("REMASEP_OD")
    od["B1"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["B1"].font = _BOLD
    od.merge_cells("C2:D2")
    od["C2"] = "20 A 24 AÑOS"
    od["E2"] = "TOTAL"
    od["C3"] = "Hombres"
    od["D3"] = "Mujeres"
    od["A6"] = "5010009 - VIDRIO IONÓMERO"
    od["B6"] = "Obturación de vidrio ionómero"
    od["C6"] = _countifs("Hombre", 20, 24)    # label Hombres + formula *Hombre* -> CONFIRMED
    od["D6"] = _countifs("Hombre", 20, 24)    # label Mujeres pero formula *Hombre* -> CONFLICT
    od["E6"] = "=SUM(C6+D6)"                   # DOWNSTREAM_TOTAL, columna TOTAL
    od["A7"] = "Prestación sin código"
    od["C7"] = _countifs("Mujer", 20, 24)     # label Hombres + formula *Mujer* -> CONFLICT

    b2 = wb.create_sheet("B2 ANEXO")
    b2["B3"] = "A. KINESIOLOGÍA"
    b2["B3"].font = _BOLD
    b2["D3"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"x\")"
    b2["B5"] = "0601105"
    b2["C5"] = "Atención Kinesiológica Integral"
    b2["D5"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"y\")"

    path = tmp_path / "mapping.xlsx"
    wb.save(path)
    return mp.build_mapping(path), path


def _by_cell(metrics):
    return {(m.sheet, m.cell): m for m in metrics}


# --- estructura general ------------------------------------------


def test_source_and_form(scenario):
    (_inv, _cfg, metrics), _ = scenario
    assert all(m.source == "MEDINET" for m in metrics)
    by = _by_cell(metrics)
    assert by[("REMASEP_OD", "C6")].form == "REMASEP_OD"
    assert by[("B2 ANEXO", "D5")].form == "B2_ANEXO"


def test_sex_confirmed_and_conflict(scenario):
    (_inv, _cfg, metrics), _ = scenario
    by = _by_cell(metrics)
    c6 = by[("REMASEP_OD", "C6")]
    assert (c6.sex_value, c6.sex_status) == (sm.SEX_MALE, sm.CONFIRMED)
    d6 = by[("REMASEP_OD", "D6")]
    assert d6.sex_status == sm.CONFLICT
    assert d6.mapping_status == "CONFLICT"
    assert any(cf.dimension == "SEX" for cf in d6.conflicts)


def test_age_confirmed(scenario):
    (_inv, _cfg, metrics), _ = scenario
    c6 = _by_cell(metrics)[("REMASEP_OD", "C6")]
    assert (c6.age_min_years, c6.age_max_years, c6.age_status) == (20, 24, sm.CONFIRMED)


def test_procedure_code_extracted(scenario):
    (_inv, _cfg, metrics), _ = scenario
    c6 = _by_cell(metrics)[("REMASEP_OD", "C6")]
    assert (c6.procedure_code_raw, c6.procedure_label_raw) == ("5010009", "VIDRIO IONÓMERO")
    assert c6.procedure_status == sm.PROC_EXPLICIT
    c7 = _by_cell(metrics)[("REMASEP_OD", "C7")]
    assert c7.procedure_code_raw == ""  # "Prestación sin código"


def test_aggregation_scope_total_not_from_kind(scenario):
    (_inv, _cfg, metrics), _ = scenario
    e6 = _by_cell(metrics)[("REMASEP_OD", "E6")]
    assert e6.kind == "DOWNSTREAM_TOTAL"
    assert e6.aggregation_scope == sm.SCOPE_TOTAL  # por el rótulo "TOTAL", no por el kind


def test_b2_anexo_row_only(scenario):
    (_inv, _cfg, metrics), _ = scenario
    by = _by_cell(metrics)
    d5 = by[("B2 ANEXO", "D5")]
    assert d5.sex_status == sm.NOT_APPLICABLE
    assert d5.age_status == sm.NOT_APPLICABLE
    assert d5.procedure_code_raw == "0601105"
    assert d5.mapping_status == "CONFIRMED"
    d3 = by[("B2 ANEXO", "D3")]  # fila-sección "A. KINESIOLOGÍA" -> subtotal
    assert d3.aggregation_scope == sm.SCOPE_SUBTOTAL


def test_c6_mapping_confirmed(scenario):
    (_inv, _cfg, metrics), _ = scenario
    assert _by_cell(metrics)[("REMASEP_OD", "C6")].mapping_status == "CONFIRMED"


# --- artefactos --------------------------------------------------


def test_outputs_and_summary(scenario, tmp_path):
    (inv, cfg, metrics), _ = scenario
    files = mp.write_outputs(inv, cfg, metrics, tmp_path / "out")
    names = {f.name for f in files}
    assert names >= {
        "semantic_metrics.csv", "dimension_evidence.csv", "mapping_coverage.csv",
        "mapping_conflicts.csv", "procedure_codes.csv", "manual_review_sample.csv",
        "summary.json", "README.md",
    }
    summary = json.loads((tmp_path / "out" / "summary.json").read_text(encoding="utf-8"))
    assert summary["total_semantic_metrics"] == len(metrics)
    assert summary["mapping_conflict"] == 2  # D6 y C7
    assert summary["mapping_config_version"] == "semantic_mapping_2026"
    assert summary["mapping_status"] == "PRELIMINARY_NOT_VALIDATED"


def test_dimension_evidence_preserves_raw_and_normalized(scenario, tmp_path):
    (inv, cfg, metrics), _ = scenario
    mp.write_outputs(inv, cfg, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "dimension_evidence.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    sex_col = [r for r in rows if r["dimension"] == "SEX" and r["evidence_type"] == "COLUMN_LABEL"]
    assert any(r["raw_value"] == "Hombres" and r["normalized_value"] == "HOMBRES" for r in sex_col)


def test_conflicts_csv_records_not_resolves(scenario, tmp_path):
    (inv, cfg, metrics), _ = scenario
    mp.write_outputs(inv, cfg, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "mapping_conflicts.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert {r["metric_id"].split("::")[-1] for r in rows} == {"D6", "C7"}
    assert all(r["conflict_type"] == "SEX_LABEL_FORMULA_MISMATCH" for r in rows)


def test_manual_review_sample_is_deterministic_and_stratified(scenario, tmp_path):
    (_inv, _cfg, metrics), _ = scenario
    s1 = mp._manual_review_sample(metrics)
    s2 = mp._manual_review_sample(metrics)
    assert [m.metric_id for m in s1] == [m.metric_id for m in s2]
    # todos los CONFLICT presentes
    conflict_ids = {m.metric_id for m in metrics if m.mapping_status == "CONFLICT"}
    assert conflict_ids <= {m.metric_id for m in s1}


def test_procedure_codes_csv_only_explicit(scenario, tmp_path):
    (inv, cfg, metrics), _ = scenario
    mp.write_outputs(inv, cfg, metrics, tmp_path / "out")
    with open(tmp_path / "out" / "procedure_codes.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert {r["procedure_code_raw"] for r in rows} == {"5010009", "0601105"}


def test_artifacts_have_no_individual_patient_data(scenario, tmp_path):
    (inv, cfg, metrics), _ = scenario
    files = mp.write_outputs(inv, cfg, metrics, tmp_path / "out")
    for f in files:
        assert "SECRET-XYZ" not in f.read_text(encoding="utf-8"), f.name


def test_cli_main(scenario, tmp_path, capsys):
    (_inv, _cfg, _metrics), path = scenario
    code = mp.main([str(path), "--output", str(tmp_path / "out")])
    assert code == 0
    out = capsys.readouterr().out
    assert "SemanticMetric:" in out and "CONFLICT" in out
