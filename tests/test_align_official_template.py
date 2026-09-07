"""Tests CLI de la alineación semántica generador → plantilla (Sprint 3.5).

Generador y plantilla sintéticos con **layouts distintos** (coordenadas
diferentes, misma semántica). Sin data/local.
"""

from __future__ import annotations

import csv
import hashlib
import json

import align_official_template as af
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, Protection

from remasep.services import template_alignment as ta

_DETAIL = "Atenciones - Detalles de citas"
_BOLD = Font(bold=True)
_UNLOCKED = Protection(locked=False)
_SECRET = "PACIENTE-SECRETO-XYZ"

_COLS = {
    "DIA_CITA": "D", "FECHA_NACIMIENTO": "G", "SEXO": "H", "SUCURSAL": "K",
    "ESPECIALIDAD": "M", "TIPO_DE_CITA": "O", "PRESTACION": "AA",
    "ESTADO": "P", "MODALIDAD": "J", "PRESTACION_REALIZADA": "AB",
}
_HDR = {
    "D": "DIA CITA", "G": "FECHA NACIMIENTO", "H": "SEXO", "K": "SUCURSAL",
    "M": "ESPECIALIDAD", "O": "TIPO DE CITA", "AA": "PRESTACIÓN",
    "P": "ESTADO", "J": "MODALIDAD", "AB": "PRESTACIÓN REALIZADA", "BA": "MARKER",
}


def _row(**kw):
    base = {
        "DIA_CITA": "01-07-2026", "FECHA_NACIMIENTO": "01-01-2003", "SEXO": "Hombre",
        "SUCURSAL": "Santiago", "ESPECIALIDAD": "Odontología General",
        "TIPO_DE_CITA": "OTRA", "PRESTACION": "VIDRIO IONÓMERO", "ESTADO": "Atendido",
        "MODALIDAD": "Fonasa B", "PRESTACION_REALIZADA": "",
    }
    base.update(kw)
    return base


def _countifs(sex_pat, lo, hi):
    return (
        f"=COUNTIFS('{_DETAIL}'!$AE:$AE,\"*VIDRIO IONÓMERO*\"&\"*{sex_pat}*\","
        f"'{_DETAIL}'!$AF:$AF,\">={lo}\",'{_DETAIL}'!$AF:$AF,\"<={hi}\")"
    )


def _build_generator(path):
    wb = Workbook()
    detail = wb.active
    detail.title = _DETAIL
    for letter, header in _HDR.items():
        detail[f"{letter}1"] = header
    rows = [
        _row(),
        _row(SEXO="Mujer"),
        _row(FECHA_NACIMIENTO="01-01-1999", ESTADO=_SECRET),  # sentinel PII
        {},
    ]
    for i, r in enumerate(rows, start=2):
        for sem, letter in _COLS.items():
            if r.get(sem) is not None:
                detail[f"{letter}{i}"] = r[sem]
        detail[f"BA{i}"] = "x"

    od = wb.create_sheet("REMASEP_OD")
    od["B1"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["B1"].font = _BOLD
    od.merge_cells("C2:D2")
    od["C2"] = "20 A 24 AÑOS"
    od["E2"] = "TOTAL"
    od["C3"] = "Hombres"
    od["D3"] = "Mujeres"
    od["A6"] = "5010009 - VIDRIO IONÓMERO"
    od["B6"] = "Obturación de vidrio ionómero"
    od["C6"] = _countifs("Hombre", 20, 24)    # AUTO_READY
    od["D6"] = _countifs("Hombre", 20, 24)    # CONFLICT -> excluido
    od["A7"] = "5010010 - OTRO INSUMO"
    od["C7"] = _countifs("Mujer", 20, 24)     # CONFLICT -> excluido

    b2 = wb.create_sheet("B2 ANEXO")
    b2["B3"] = "A. KINESIOLOGÍA"
    b2["B3"].font = _BOLD
    b2["B5"] = "0601105"
    b2["C5"] = "Atención Kinesiológica Integral Ambulatoria"
    b2["D5"] = "=COUNTIF('" + _DETAIL + "'!$AC:$AC,\"x\")"
    wb.save(path)


def _build_template(path):
    """Mismo vocabulario semántico, coordenadas totalmente distintas.

    Como la plantilla oficial real: hojas **protegidas**; sólo las celdas de dato
    van DESBLOQUEADAS (`Protection(locked=False)`). Los rótulos/códigos quedan
    bloqueados -> `STRUCTURAL`, no `DIRECT_INPUT_TARGET`.
    """
    wb = Workbook()
    od = wb.active
    od.title = "REMASEP_OD"
    od.protection.sheet = True
    od["A20"] = "SECCIÓN A: PRUEBAS ODONTOLÓGICAS"
    od["A20"].font = _BOLD
    od.merge_cells("D21:E21")
    od["D21"] = "20 A 24 AÑOS"
    od["D22"] = "Hombres"
    od["E22"] = "Mujeres"
    od["A25"] = "5010009 - VIDRIO IONÓMERO"
    od["B25"] = "Obturación de vidrio ionómero"
    od["D25"].protection = _UNLOCKED       # celda de dato -> DIRECT_INPUT_TARGET
    od["F25"] = "=D25"                     # fórmula -> FORMULA_TARGET
    od["G21"] = "TOTAL"
    od["G22"] = "Ambos sexo"

    b1 = wb.create_sheet("REMASEP B1")
    b1.protection.sheet = True
    b1["A1"] = "SECCIÓN A: INTERVENCIONES QUIRÚRGICAS"
    b1["A1"].font = _BOLD
    b1.merge_cells("C2:D2")
    b1["C2"] = "POR GRUPO DE EDAD"
    b1["C3"] = "Menor de 10 años"
    b1["D3"] = "Hombres"
    b1["A5"] = "Electivas"
    b1["B5"] = "Mayor Ambulatorias"
    b1["C5"].protection = _UNLOCKED        # celda de dato, región EGRESOS

    b2 = wb.create_sheet("B2 ANEXO")
    b2.protection.sheet = True
    b2["A1"] = "CÓDIGOS"
    b2["B1"] = "GLOSA"
    b2["C1"] = "TOTAL"
    b2["A2"] = "A. KINESIOLOGÍA"
    b2["A2"].font = _BOLD
    b2["A3"] = "0601105"                   # código: bloqueado -> STRUCTURAL
    b2["B3"] = "Atención Kinesiológica Integral Ambulatoria"
    b2["C3"].protection = _UNLOCKED        # celda de dato (match por código)
    wb.save(path)


@pytest.fixture
def paths(tmp_path):
    gen = tmp_path / "generator.xlsx"
    tpl = tmp_path / "template.xlsx"
    _build_generator(gen)
    _build_template(tpl)
    return gen, tpl


@pytest.fixture
def analysis(paths, tmp_path):
    gen, tpl = paths
    a = af.build_alignment_analysis(str(gen), str(tpl))
    files = af.write_outputs(a, tmp_path / "out")
    return a, files, tmp_path / "out"


# ---------------------------------------------------------------------------


def test_runs_and_writes_expected_files(analysis):
    _a, files, out = analysis
    names = {f.name for f in files}
    assert names == {
        "target_metric_inventory.csv", "source_metric_inventory.csv",
        "semantic_alignment.csv", "alignment_evidence.csv", "ambiguous_matches.csv",
        "unmatched_sources.csv", "unmatched_targets.csv", "source_exclusions.csv",
        "target_source_expectations.csv", "alignment_coverage.csv",
        "alignment_manual_review_sample.csv", "summary.json", "README.md",
    }
    assert (out / "summary.json").is_file()


def test_source_and_template_hash_unchanged(paths):
    gen, tpl = paths
    before = (hashlib.sha256(gen.read_bytes()).hexdigest(),
              hashlib.sha256(tpl.read_bytes()).hexdigest())
    af.build_alignment_analysis(str(gen), str(tpl))
    after = (hashlib.sha256(gen.read_bytes()).hexdigest(),
             hashlib.sha256(tpl.read_bytes()).hexdigest())
    assert before == after


def test_matching_is_not_coordinate_based(analysis):
    a, _files, _out = analysis
    rows = a["report"].rows
    od = [r for r in rows if r.source_metric_id.endswith("REMASEP_OD::C6")]
    assert od, "la métrica AUTO_READY de OD debería emparejarse"
    r = od[0]
    assert r.match_status in (ta.EXACT_SEMANTIC_MATCH, ta.STRONG_MATCH)
    assert r.target_cell != "C6"  # coordenada distinta a la del source


def test_blocked_conflicts_excluded(analysis):
    _a, _files, out = analysis
    with open(out / "source_exclusions.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    reasons = {r["reason"] for r in rows}
    assert ta.EXCL_BLOCKED_CONFLICT in reasons
    # los CONFLICT (D6, C7) no entran a semantic_alignment.csv como match
    with open(out / "semantic_alignment.csv", encoding="utf-8", newline="") as fh:
        aligned_ids = {r["source_metric_id"] for r in csv.DictReader(fh)}
    assert not any(i.endswith(("D6", "C7")) for i in aligned_ids)


def test_procedure_code_match_b2(analysis):
    _a, _files, out = analysis
    with open(out / "semantic_alignment.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    b2 = [r for r in rows if r["source_form"] == "B2_ANEXO"]
    assert b2, "la prestación B2 con código debería emparejarse"
    assert b2[0]["procedure_match"] == ta.EV_MATCH
    assert b2[0]["target_cell"] != "D5"


def test_b1_input_targets_are_non_medinet_region(analysis):
    _a, _files, out = analysis
    with open(out / "unmatched_targets.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    b1_inputs = [
        r for r in rows
        if r["target_sheet"] == "REMASEP B1"
        and r["target_alignment_role"] == ta.ROLE_INPUT_TARGET
    ]
    assert b1_inputs
    assert all(r["expected_source"] == ta.EXP_EGRESOS for r in b1_inputs)
    assert all(r["reason"] == ta.TGT_NON_MEDINET_REGION for r in b1_inputs)
    assert all(r["is_genuine_unmatched_medinet_input"] == "no" for r in b1_inputs)


def test_summary_consistency(analysis):
    _a, _files, out = analysis
    s = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    assert s["alignment_status"] == "PRELIMINARY_NOT_VALIDATED"
    assert s["template_approval_status"] == "UNCONFIRMED"
    assert s["candidate_golden_status"] == "CANDIDATE_PENDING_SOURCE_COMPLETENESS"
    assert s["source_hash_verified_unchanged"] is True
    matched = sum(s["match_status_counts"].values())
    assert matched + s["unmatched_sources"] == s["eligible_source_metrics"]
    assert s["eligible_source_metrics"] + s["excluded_source_metrics"] == s["total_source_metrics"]


def test_summary_separates_target_concepts(analysis):
    _a, _files, out = analysis
    s = json.loads((out / "summary.json").read_text(encoding="utf-8"))
    # identidades de cobertura: no se mezclan cantidades
    assert (s["input_targets"]
            == s["medinet_input_targets_expected"]
            + s["non_medinet_input_targets"]
            + s["unknown_source_input_targets"])
    assert (s["medinet_input_targets_expected"]
            == s["matched_medinet_input_targets"]
            + s["genuine_unmatched_medinet_input_targets"])
    assert (s["target_inventory_total"]
            == s["input_targets"] + s["formula_targets"] + s["structural_targets"]
            + s["target_by_alignment_role"].get("VALIDATION", 0)
            + s["target_by_alignment_role"].get("UNKNOWN", 0))
    # STRUCTURAL nunca es un target MEDINET faltante
    assert s["structural_targets"] > 0
    assert s["genuine_unmatched_medinet_input_targets"] <= s["medinet_input_targets_expected"]


def test_unmatched_targets_reason_taxonomy(analysis):
    _a, _files, out = analysis
    with open(out / "unmatched_targets.csv", encoding="utf-8", newline="") as fh:
        rows = list(csv.DictReader(fh))
    reasons = {r["reason"] for r in rows}
    assert reasons <= {
        ta.TGT_NO_MEDINET_SOURCE_FOUND, ta.TGT_NON_MEDINET_REGION,
        ta.TGT_DERIVED_NOT_INPUT, ta.TGT_STRUCTURAL_NOT_ALIGNMENT,
        ta.TGT_VALIDATION_NOT_ALIGNMENT, ta.TGT_UNKNOWN_SOURCE,
    }
    # sólo NO_MEDINET_SOURCE_FOUND es "genuino"
    for r in rows:
        genuine = r["is_genuine_unmatched_medinet_input"] == "yes"
        assert genuine == (r["reason"] == ta.TGT_NO_MEDINET_SOURCE_FOUND)
    # STRUCTURAL / DERIVED nunca son genuinos
    for r in rows:
        if r["reason"] in (ta.TGT_STRUCTURAL_NOT_ALIGNMENT, ta.TGT_DERIVED_NOT_INPUT):
            assert r["is_genuine_unmatched_medinet_input"] == "no"


def test_no_pii_in_any_artifact(analysis):
    _a, files, _out = analysis
    for f in files:
        assert _SECRET not in f.read_text(encoding="utf-8"), f.name


def test_cli_main_smoke(paths, tmp_path, capsys):
    gen, tpl = paths
    code = af.main([str(gen), str(tpl), "--output", str(tmp_path / "cli")])
    assert code == 0
    out = capsys.readouterr().out
    assert "candidate_golden_status: CANDIDATE_PENDING_SOURCE_COMPLETENESS" in out
    assert "match_status" in out


def test_cli_main_bad_file(tmp_path):
    code = af.main([str(tmp_path / "nope.xlsx"), str(tmp_path / "nope2.xlsx")])
    assert code == 2
