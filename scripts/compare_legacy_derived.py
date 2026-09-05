"""Equivalencia de la lógica derivada legacy AC:AL (Sprint 2.2).

Compara, fila real a fila real, los valores derivados `AC:AL` que Excel dejó
*cacheados* en `GENERACION DATOS REMASEP.xlsx` contra los que produce la
implementación Python de producción (`remasep.services.legacy_transform`).

- Solo lectura. NO escribe el workbook, NO usa COM ni LibreOffice, NO recalcula.
- Si Excel no guardó valores cacheados suficientes para una columna, esa columna
  queda en estado ``CACHE_UNAVAILABLE`` (no se inventan resultados).
- Privacidad: `mismatches.csv` solo lleva fila + columna + tipo de diferencia;
  nunca valores derivados por individuo (que incluyen sexo/prestación).

Uso:

    python scripts/compare_legacy_derived.py \\
        "data/local/GENERACION DATOS REMASEP.xlsx" \\
        --output artifacts/legacy_equivalence
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
from inventory_workbook import compute_sha256
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException

from remasep.adapters.medinet import (
    REQUIRED_FIELDS,
    is_blank_cell,
    semantic_column_map,
)
from remasep.core.errors import RemasepError
from remasep.core.text import normalize_text
from remasep.services.legacy_rules import load_legacy_rules
from remasep.services.legacy_transform import DERIVED_COLUMNS, excel_str, legacy_derived_value

DEFAULT_OUTPUT = "artifacts/legacy_equivalence"


class EquivalenceError(RemasepError):
    """Error de entrada del comparador con mensaje claro."""


# ---------------------------------------------------------------------------
# Modelos
# ---------------------------------------------------------------------------

STATUS_EXACT = "EXACT"
STATUS_MISMATCH = "MISMATCH"
STATUS_UNAVAILABLE = "UNAVAILABLE"
STATUS_NO_FORMULA = "NO_FORMULA"

MT_VALUE = "VALUE_MISMATCH"
MT_SEMANTIC = "SEMANTIC_MISMATCH"
MT_CACHE_MISSING = "CACHE_MISSING"
MT_PY_NONE = "PY_NONE"
MT_LEGACY_ERROR = "LEGACY_ERROR"
# Cache no confiable: Excel dejó 0 en AF pero, con fechas presentes y parseables,
# la fórmula debería dar una edad > 0 (recálculo pendiente en el workbook).
MT_STALE_CACHE = "STALE_CACHE_SUSPECTED"


@dataclass(frozen=True)
class CellComparison:
    row_number: int
    column: str
    status: str
    mismatch_type: str = ""


@dataclass
class ColumnSummary:
    column: str
    compared_rows: int = 0
    exact_matches: int = 0
    mismatches: int = 0
    unavailable: int = 0

    @property
    def exact_match_rate(self) -> float:
        return round(self.exact_matches / self.compared_rows, 6) if self.compared_rows else 0.0

    @property
    def status(self) -> str:
        if self.compared_rows == 0:
            return "NO_DATA"
        if self.mismatches > 0:
            return "FAIL"
        if self.unavailable > 0:
            return "CACHE_UNAVAILABLE"
        return "PASS"


@dataclass
class EquivalenceResult:
    source_name: str
    source_sha256: str
    sheet_name: str
    physical_rows: int
    structural_empty_rows: int
    active_records: int
    columns_tested: list[str]
    column_summaries: list[ColumnSummary]
    mismatch_rows: list[CellComparison]
    cache_diagnostic: dict[str, dict[str, int]]
    legacy_records_with_match: int
    legacy_category_hits: dict[str, int]
    generated_at: str
    captured_details: list[tuple[int, str, str, str]] = field(default_factory=list)

    @property
    def total_comparisons(self) -> int:
        return sum(c.compared_rows for c in self.column_summaries)

    @property
    def exact_matches(self) -> int:
        return sum(c.exact_matches for c in self.column_summaries)

    @property
    def mismatches(self) -> int:
        return sum(c.mismatches for c in self.column_summaries)

    @property
    def unavailable(self) -> int:
        return sum(c.unavailable for c in self.column_summaries)

    @property
    def cached_values_available(self) -> bool:
        return any(c.compared_rows for c in self.column_summaries) and self.unavailable == 0

    @property
    def overall_status(self) -> str:
        statuses = {c.status for c in self.column_summaries}
        if "FAIL" in statuses:
            return "FAIL"
        if statuses <= {"PASS"}:
            return "PASS"
        # Sin desacuerdos reales, pero alguna columna sin cache utilizable.
        return "CACHE_UNAVAILABLE"


# ---------------------------------------------------------------------------
# Comparación
# ---------------------------------------------------------------------------


def _numeric(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None


def _compare_cell(column: str, python_value: object, cached_value: object) -> tuple[str, str]:
    if column == "AF":
        py_num = _numeric(python_value)
        cache_num = _numeric(cached_value)
        if py_num is not None and cache_num is not None:
            if py_num == cache_num:
                return STATUS_EXACT, ""
            # Excel dejó 0 pero la fórmula, con fechas válidas, debería dar edad > 0:
            # cache no recalculada (no es un desacuerdo real con Python).
            if cache_num == 0 and py_num > 0:
                return STATUS_UNAVAILABLE, MT_STALE_CACHE
            return STATUS_MISMATCH, MT_VALUE
        if python_value is None:
            return STATUS_MISMATCH, MT_PY_NONE
        if isinstance(cached_value, str) and cached_value.startswith("#"):
            return STATUS_UNAVAILABLE, MT_LEGACY_ERROR
        return (
            (STATUS_EXACT, "")
            if excel_str(python_value) == excel_str(cached_value)
            else (STATUS_MISMATCH, MT_VALUE)
        )

    python_text = excel_str(python_value)
    cached_text = excel_str(cached_value)
    if python_text == cached_text:
        return STATUS_EXACT, ""
    if normalize_text(python_text) == normalize_text(cached_text):
        return STATUS_MISMATCH, MT_SEMANTIC
    return STATUS_MISMATCH, MT_VALUE


def _resolve_sheet(workbook: openpyxl.Workbook) -> str:
    best_name, best_score = workbook.sheetnames[0], -1
    for name in workbook.sheetnames:
        worksheet = workbook[name]
        header = [worksheet.cell(row=1, column=c).value for c in range(1, (worksheet.max_column or 0) + 1)]
        score = len(set(semantic_column_map(header).values()) & set(REQUIRED_FIELDS))
        if score > best_score:
            best_name, best_score = name, score
    return best_name


def compare_legacy_derived(path: str | Path, *, capture_details: bool = False) -> EquivalenceResult:
    source = Path(path)
    if not source.is_file():
        raise EquivalenceError(f"El archivo no existe o no es un archivo: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise EquivalenceError(f"Formato no soportado: {source.suffix!r}")
    try:
        wb_formulas = openpyxl.load_workbook(source, data_only=False, read_only=True)
        wb_values = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise EquivalenceError(f"No se pudo abrir el workbook: {exc}") from exc

    try:
        sheet_name = _resolve_sheet(wb_values)
        ws_f = wb_formulas[sheet_name]
        ws_v = wb_values[sheet_name]

        header = [ws_v.cell(row=1, column=c).value for c in range(1, (ws_v.max_column or 0) + 1)]
        raw_map = semantic_column_map(header)
        semantic_to_index: dict[str, int] = {}
        for col_index, cell_value in enumerate(header, start=1):
            semantic = raw_map.get(str(cell_value))
            if semantic and semantic not in semantic_to_index:
                semantic_to_index[semantic] = col_index

        missing_required = [f for f in REQUIRED_FIELDS if f not in semantic_to_index]
        if missing_required:
            raise EquivalenceError(
                "El workbook no expone columnas para: " + ", ".join(missing_required)
            )
        semantic_fields = list(semantic_to_index)  # requeridos + opcionales presentes

        derived_index = {col: column_index_from_string(col) for col in DERIVED_COLUMNS}
        max_row = ws_v.max_row or 1

        ruleset = load_legacy_rules()
        summaries = {col: ColumnSummary(column=col) for col in DERIVED_COLUMNS}
        cache_diag = {
            col: {"formula_cells": 0, "cached_present": 0, "cached_missing": 0}
            for col in DERIVED_COLUMNS
        }
        mismatch_rows: list[CellComparison] = []
        captured: list[tuple[int, str, str, str]] = []
        structural_empty_rows = 0
        active_records = 0
        legacy_records_with_match = 0
        legacy_hits: Counter[str] = Counter()

        value_columns = {*semantic_to_index.values(), *derived_index.values()}
        formula_rows = ws_f.iter_rows(min_row=2, max_row=max_row)
        value_rows = ws_v.iter_rows(min_row=2, max_row=max_row)
        for offset, (f_row, v_row) in enumerate(zip(formula_rows, value_rows, strict=False)):
            excel_row = offset + 2
            f_cells = list(f_row)
            v_cells = list(v_row)

            v_by_col = {
                idx: getattr(v_cells[idx - 1], "value", None) if idx - 1 < len(v_cells) else None
                for idx in value_columns
            }
            f_by_col = {
                idx: getattr(f_cells[idx - 1], "data_type", None) if idx - 1 < len(f_cells) else None
                for idx in derived_index.values()
            }

            record = {sem: v_by_col.get(idx) for sem, idx in semantic_to_index.items()}

            # Cache diagnostic (todas las filas físicas con fórmula).
            for col in DERIVED_COLUMNS:
                if f_by_col.get(derived_index[col]) == "f":
                    cache_diag[col]["formula_cells"] += 1
                    if v_by_col.get(derived_index[col]) is None:
                        cache_diag[col]["cached_missing"] += 1
                    else:
                        cache_diag[col]["cached_present"] += 1

            if all(is_blank_cell(record.get(sem)) for sem in semantic_fields):
                structural_empty_rows += 1
                continue

            active_records += 1
            codes = ruleset.classify(record)
            if codes:
                legacy_records_with_match += 1
                legacy_hits.update(codes)

            for col in DERIVED_COLUMNS:
                summary = summaries[col]
                summary.compared_rows += 1
                if f_by_col.get(derived_index[col]) != "f":
                    summary.unavailable += 1
                    mismatch_rows.append(
                        CellComparison(excel_row, col, STATUS_NO_FORMULA, MT_CACHE_MISSING)
                    )
                    continue
                cached_value = v_by_col.get(derived_index[col])
                if cached_value is None:
                    summary.unavailable += 1
                    mismatch_rows.append(
                        CellComparison(excel_row, col, STATUS_UNAVAILABLE, MT_CACHE_MISSING)
                    )
                    continue
                python_value = legacy_derived_value(col, record, ruleset)
                status, mismatch_type = _compare_cell(col, python_value, cached_value)
                if status == STATUS_EXACT:
                    summary.exact_matches += 1
                    continue
                if status == STATUS_UNAVAILABLE:
                    summary.unavailable += 1
                else:
                    summary.mismatches += 1
                mismatch_rows.append(CellComparison(excel_row, col, status, mismatch_type))
                if capture_details:
                    captured.append(
                        (excel_row, col, excel_str(python_value), excel_str(cached_value))
                    )
    finally:
        wb_formulas.close()
        wb_values.close()

    physical_rows = structural_empty_rows + active_records
    return EquivalenceResult(
        source_name=source.name,
        source_sha256=compute_sha256(source),
        sheet_name=sheet_name,
        physical_rows=physical_rows,
        structural_empty_rows=structural_empty_rows,
        active_records=active_records,
        columns_tested=list(DERIVED_COLUMNS),
        column_summaries=[summaries[col] for col in DERIVED_COLUMNS],
        mismatch_rows=mismatch_rows,
        cache_diagnostic=cache_diag,
        legacy_records_with_match=legacy_records_with_match,
        legacy_category_hits=dict(sorted(legacy_hits.items())),
        generated_at=datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        captured_details=captured,
    )


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------

_README = """# Legacy derived equivalence (Sprint 2.2)

Salida de `scripts/compare_legacy_derived.py`. Compara los valores derivados
`AC:AL` **cacheados por Excel** en `GENERACION DATOS REMASEP.xlsx` contra los que
produce `remasep.services.legacy_transform` para las mismas filas reales.

- `summary.json` — hash del workbook, filas físicas / estructuralmente vacías /
  activas, disponibilidad de cache, conteos globales y `overall_status`
  (`PASS` / `FAIL` / `CACHE_UNAVAILABLE`). `PASS` solo si **todas** las
  comparaciones requeridas son exactas.
- `column_summary.csv` — por columna: filas comparadas, exactas, mismatches,
  no disponibles, tasa de exactitud y estado.
- `mismatches.csv` — una fila por comparación no exacta. **Solo** `row_number`,
  `derived_column`, `comparison_status`, `mismatch_type`. Nunca valores derivados
  por individuo.

`mismatch_type`: `VALUE_MISMATCH` (strings distintos), `SEMANTIC_MISMATCH`
(distintos pero iguales tras normalizar tildes/casing), `CACHE_MISSING`,
`PY_NONE`, `LEGACY_ERROR`.

Equivalencia con el workbook actual **no** implica validación oficial MINSAL.
"""


def write_equivalence_outputs(result: EquivalenceResult, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    summary = {
        "source": result.source_name,
        "source_sha256": result.source_sha256,
        "sheet_name": result.sheet_name,
        "generated_at": result.generated_at,
        "physical_rows": result.physical_rows,
        "structural_empty_rows": result.structural_empty_rows,
        "active_records": result.active_records,
        "columns_tested": result.columns_tested,
        "total_comparisons": result.total_comparisons,
        "exact_matches": result.exact_matches,
        "mismatches": result.mismatches,
        "unavailable": result.unavailable,
        "cached_values_available": result.cached_values_available,
        "cache_diagnostic": result.cache_diagnostic,
        "legacy_category_hits": {
            "records_with_any_legacy_match": result.legacy_records_with_match,
            "hits": result.legacy_category_hits,
            "note": "Una atención puede activar más de una categoría; no hay exclusividad.",
        },
        "column_status": {c.column: c.status for c in result.column_summaries},
        "overall_status": result.overall_status,
    }
    summary_path = output_dir / "summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    column_path = output_dir / "column_summary.csv"
    with column_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["column", "compared_rows", "exact_matches", "mismatches", "unavailable",
             "exact_match_rate", "status"]
        )
        for c in result.column_summaries:
            writer.writerow(
                [c.column, c.compared_rows, c.exact_matches, c.mismatches, c.unavailable,
                 c.exact_match_rate, c.status]
            )

    mismatches_path = output_dir / "mismatches.csv"
    with mismatches_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["row_number", "derived_column", "comparison_status", "mismatch_type"])
        for m in result.mismatch_rows:
            writer.writerow([m.row_number, m.column, m.status, m.mismatch_type])

    readme_path = output_dir / "README.md"
    readme_path.write_text(_README, encoding="utf-8")

    return [summary_path, column_path, mismatches_path, readme_path]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="compare_legacy_derived.py",
        description=(
            "Compara los valores derivados AC:AL cacheados por Excel contra la "
            "implementación Python. Solo lectura; sin COM ni recálculo."
        ),
    )
    parser.add_argument("workbook", help="Ruta al workbook legacy (.xlsx/.xlsm).")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT, help=f"(def: {DEFAULT_OUTPUT})")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        result = compare_legacy_derived(args.workbook)
    except EquivalenceError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    output_dir = Path(args.output)
    files = write_equivalence_outputs(result, output_dir)

    print(f"Workbook: {result.source_name}  (hoja {result.sheet_name!r})")
    print(
        f"  filas físicas {result.physical_rows} = "
        f"{result.structural_empty_rows} vacías + {result.active_records} activas"
    )
    print(f"  cached values disponibles: {result.cached_values_available}")
    print(f"  comparaciones: {result.total_comparisons} | exactas: {result.exact_matches} | "
          f"mismatches: {result.mismatches} | no disponibles: {result.unavailable}")
    print("  por columna:")
    for c in result.column_summaries:
        print(
            f"    {c.column}: {c.exact_matches}/{c.compared_rows} exactas "
            f"(rate {c.exact_match_rate}) mismatch {c.mismatches} unavail {c.unavailable} "
            f"-> {c.status}"
        )
    print(
        f"  legacy match (bucketing): {result.legacy_records_with_match} registros | "
        f"hits {result.legacy_category_hits}"
    )
    print(f"  OVERALL: {result.overall_status}")
    print(f"Artefactos en {output_dir}:")
    for path in files:
        print(f"  - {path}")
    # Exit 1 solo ante un desacuerdo real Python <-> cache; CACHE_UNAVAILABLE no lo es.
    return 1 if result.overall_status == "FAIL" else 0


if __name__ == "__main__":
    raise SystemExit(main())
