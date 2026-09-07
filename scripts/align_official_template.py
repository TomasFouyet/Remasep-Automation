"""Sprint 3.5 — Alineación semántica generador MEDINET → plantilla oficial.

Empareja cada `SemanticMetric` MEDINET **elegible** (`readiness = AUTO_READY`,
Sprint 3.3) con la(s) celda(s) de la plantilla oficial que representan lo mismo,
usando evidencia por dimensión — nunca la coordenada como única prueba.

Sólo lectura: no escribe ni guarda ningún workbook, no usa COM/macros, verifica
el SHA256 de ambos archivos antes y después.

Uso:

    python scripts/align_official_template.py \\
        "data/local/GENERACION DATOS REMASEP.xlsx" \\
        "data/local/REMASEP_V1.4 Julio 2026.xlsm"
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
import validate_semantic_mapping as vs

from remasep.services import semantic_mapping as sm
from remasep.services import template_alignment as ta

DEFAULT_OUTPUT = "artifacts/official_template_semantic_alignment"
DEFAULT_GENERATOR = "data/local/GENERACION DATOS REMASEP.xlsx"
DEFAULT_TEMPLATE = "data/local/REMASEP_V1.4 Julio 2026.xlsm"
TARGET_SHEETS = ("REMASEP 01", "REMASEP B1", "B2 ANEXO", "REMASEP_OD")

APPROVAL_STATUS = "UNCONFIRMED"
CANDIDATE_GOLDEN_STATUS = "CANDIDATE_PENDING_SOURCE_COMPLETENESS"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


# ---------------------------------------------------------------------------
# Análisis
# ---------------------------------------------------------------------------


def build_alignment_analysis(generator_path: str, template_path: str) -> dict:
    gen, tpl = Path(generator_path), Path(template_path)
    for p in (gen, tpl):
        if not p.is_file():
            raise ta.TemplateAlignmentError(f"no es un archivo: {p}")
    gen_sha_pre, tpl_sha_pre = _sha256(gen), _sha256(tpl)

    inventory, _mapcfg, metrics, _mpol, _ovr, result = vs.build_validation(str(gen))
    readiness_by_id = {r.metric_id: r for r in result.readiness}

    policy = ta.load_alignment_policy()
    regions = ta.load_source_regions()

    wb_vals = openpyxl.load_workbook(tpl, data_only=True, keep_vba=True)
    wb_form = openpyxl.load_workbook(tpl, data_only=False, keep_vba=True)
    try:
        targets_by_form: dict[str, list[ta.TargetMetricContext]] = defaultdict(list)
        for sheet in TARGET_SHEETS:
            if sheet not in wb_vals.sheetnames:
                continue
            inv = ta.build_target_inventory(
                wb_vals[sheet], wb_form[sheet], policy=policy, regions=regions
            )
            for t in inv:
                targets_by_form[t.form].append(t)
    finally:
        wb_vals.close()
        wb_form.close()

    report = ta.align(metrics, readiness_by_id, dict(targets_by_form), policy, regions)

    gen_sha_post, tpl_sha_post = _sha256(gen), _sha256(tpl)
    if (gen_sha_pre, tpl_sha_pre) != (gen_sha_post, tpl_sha_post):
        raise ta.TemplateAlignmentError("¡el SHA256 de un workbook cambió durante el análisis!")

    return {
        "generator": gen, "template": tpl,
        "generator_sha256": gen_sha_pre, "template_sha256": tpl_sha_pre,
        "source_hash_verified_unchanged": True,
        "inventory": inventory, "metrics": metrics, "readiness_by_id": readiness_by_id,
        "policy": policy, "regions": regions, "report": report,
    }


# ---------------------------------------------------------------------------
# Artefactos
# ---------------------------------------------------------------------------


def _csv(path: Path, header: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def _ctx(*parts) -> str:
    flat: list[str] = []
    for part in parts:
        flat.extend(part if isinstance(part, (list, tuple)) else [part])
    return " :: ".join(str(p) for p in flat if str(p))


def _manual_review_sample(report: ta.AlignmentReport, metrics_by_id: dict) -> list[dict]:
    rows = sorted(report.rows, key=lambda r: r.source_metric_id)
    buckets = [
        lambda r: r.match_status == ta.EXACT_SEMANTIC_MATCH,
        lambda r: r.match_status == ta.STRONG_MATCH,
        lambda r: r.match_status == ta.AMBIGUOUS,
        lambda r: r.match_status == ta.CONFLICT,
        lambda r: r.source_form == "REMASEP_OD",
        lambda r: r.source_form == "REMASEP_01",
        lambda r: r.source_form == "B2_ANEXO",
        lambda r: r.source_kind == "BASE_AGGREGATION",
        lambda r: r.source_kind == "DERIVED_AGGREGATION",
        lambda r: r.source_kind == "DOWNSTREAM_TOTAL",
        lambda r: r.sex_match == ta.EV_MATCH,
        lambda r: r.age_match == ta.EV_MATCH,
        lambda r: r.procedure_match == ta.EV_MATCH,
        lambda r: r.row_match == ta.EV_PARTIAL,
        lambda r: r.target_alignment_role == ta.ROLE_DERIVED_TARGET,
        lambda r: r.target_alignment_role == ta.ROLE_INPUT_TARGET,
        lambda r: r.needs_human_review,
        lambda r: not r.needs_human_review,
        lambda r: r.target_source_expectation != ta.EXP_MEDINET,
    ]
    picked: dict[str, ta.AlignmentRow] = {}
    for pred in buckets:
        for r in [x for x in rows if pred(x)][:6]:
            picked[r.source_metric_id] = r
    out = []
    for r in sorted(picked.values(), key=lambda r: r.source_metric_id):
        m = metrics_by_id.get(r.source_metric_id)
        out.append({
            "source_metric_id": r.source_metric_id,
            "target_sheet": r.target_sheet,
            "target_cell": r.target_cell,
            "target_kind": r.target_kind,
            "target_alignment_role": r.target_alignment_role,
            "target_source_expectation": r.target_source_expectation,
            "source_context": _ctx(m.section_path_raw, m.row_path_raw, m.column_path_raw)
            if m else "",
            "target_context": r.target_semantic_signature,
            "match_status": r.match_status,
            "evidence_summary": (
                f"score={r.match_score} row={r.row_match} col={r.column_match} "
                f"sex={r.sex_match} age={r.age_match} proc={r.procedure_match}"
            ),
        })
    # excluidos por región NO-MEDINET del target (EGRESOS / RESOURCE_CALCULATION / …)
    non_medinet = sorted(
        (u for u in report.unmatched_targets if u.reason == ta.TGT_NON_MEDINET_REGION),
        key=lambda u: (u.target_sheet, u.target_cell),
    )
    for u in non_medinet[:6]:
        out.append({
            "source_metric_id": "",
            "target_sheet": u.target_sheet,
            "target_cell": u.target_cell,
            "target_kind": u.target_kind,
            "target_alignment_role": u.target_alignment_role,
            "target_source_expectation": u.expected_source,
            "source_context": "",
            "target_context": u.target_context,
            "match_status": f"NON_MEDINET_TARGET/{u.expected_source}",
            "evidence_summary": "celda de plantilla con fuente esperada NO-MEDINET; "
                                "excluida del alignment MEDINET",
        })
    return out


_README = """# Official template semantic alignment (Sprint 3.5)

Salida de `scripts/align_official_template.py`.

Empareja cada `SemanticMetric` **MEDINET** elegible (`readiness = AUTO_READY`,
Sprint 3.3) del workbook **generador** legacy con la(s) celda(s) de la
**plantilla oficial** REMASEP que representan lo mismo.

> La plantilla oficial tiene **otro layout** que el generador. El matching se
> basa en **evidencia por dimensión** (forma · sección · fila · columna · sexo ·
> edad · alcance · código de procedimiento), **no** en la coordenada. Un match
> **no** significa "MINSAL validated": reglas versionadas en
> `config/official_template_alignment_2026/` con
> `status: preliminary_pending_functional_validation`.

## `match_status`

| estado | criterio |
| --- | --- |
| `EXACT_SEMANTIC_MATCH` | todas las dimensiones aplicables coinciden (score ≥ `exact`) |
| `STRONG_MATCH` | evidencia suficiente pese a alguna diferencia estructural (score ≥ `strong`) |
| `AMBIGUOUS` | más de un target dentro de `tie_margin` del mejor score — **no** se elige uno |
| `NO_MATCH` | ningún target con evidencia suficiente (ver `unmatched_sources.csv`) |
| `CONFLICT` | dimensión incompatible (sexo / edad disjunta / código distinto) |

## `target_kind` vs. `target_alignment_role` vs. `target_source_expectation`

Tres conceptos **distintos**:

| campo | pregunta | valores |
| --- | --- | --- |
| `target_kind` | ¿qué **es** físicamente la celda? | `DIRECT_INPUT_TARGET` · `FORMULA_TARGET` · `STRUCTURAL` · `VALIDATION` · `UNKNOWN` |
| `target_alignment_role` | ¿**para qué** sirve en la alineación? | `INPUT_TARGET` · `DERIVED_TARGET` · `STRUCTURAL` · `VALIDATION` · `UNKNOWN` |
| `target_source_expectation` | ¿de qué **fuente** debería venir? | `MEDINET` · `EGRESOS` · `RESOURCE_CALCULATION` · `SURGICAL_TABLE` · `CONTROL_METADATA` · `UNKNOWN` |

- Una celda `STRUCTURAL` **nunca** es un target que requiere source.
- Una `DERIVED_TARGET` (fórmula) depende de inputs pero **no** es punto de
  ingreso.
- `target_source_expectation` **default = `UNKNOWN`**: sólo se marca `MEDINET`
  cuando una región con evidencia estructural positiva lo reclama, o cuando un
  source MEDINET se alinea fuerte a la celda. Sólo las regiones explícitamente
  **NO-MEDINET** (`EGRESOS` / recursos / tabla quirúrgica) **excluyen** el
  emparejamiento (una expectativa `UNKNOWN` no bloquea). Ver
  `config/official_template_alignment_2026/source_regions.yaml`.

## Cobertura MEDINET — no mezclar cantidades

`summary.json` responde por separado:

- inventario físico total de targets; cuántos `INPUT_TARGET` / `DERIVED_TARGET` /
  `STRUCTURAL`;
- de los `INPUT_TARGET`: cuántos se esperan de MEDINET (`medinet_input_targets_expected`),
  cuántos son NO-MEDINET, cuántos `UNKNOWN`;
- cuántos MEDINET input targets quedaron alineados
  (`matched_medinet_input_targets`) y cuántos **genuinamente** sin source
  (`genuine_unmatched_medinet_input_targets`).

`unmatched_target_rows_technical` es el inventario **técnico** de todas las filas
sin match (incluye `STRUCTURAL` / `DERIVED` / `UNKNOWN`); **no** son "targets
MEDINET faltantes".

## Archivos

- `target_metric_inventory.csv` — celdas de la plantilla con `target_kind`,
  `target_alignment_role`, `target_source_expectation` y contexto semántico.
- `source_metric_inventory.csv` — métricas MEDINET con readiness y elegibilidad.
- `semantic_alignment.csv` — una fila por par emparejado, con el desglose de
  evidencia por dimensión, `target_alignment_role` y `match_score`.
- `alignment_evidence.csv` — evidencia dimensión a dimensión (source vs target).
- `ambiguous_matches.csv` — sources con varios targets plausibles.
- `unmatched_sources.csv` — sources sin match (`NO_TARGET` /
  `TARGET_SOURCE_NOT_MEDINET` / `CONTEXT_INSUFFICIENT` / `CONFLICT`).
- `unmatched_targets.csv` — TODAS las celdas sin match, con `reason` tipado
  (`NO_MEDINET_SOURCE_FOUND` = genuino · `NON_MEDINET_REGION` ·
  `DERIVED_TARGET_NOT_INPUT` · `STRUCTURAL_NOT_ALIGNMENT_TARGET` ·
  `VALIDATION_NOT_ALIGNMENT_TARGET` · `UNKNOWN_SOURCE`) y la columna
  `is_genuine_unmatched_medinet_input`.
- `source_exclusions.csv` — sources no elegibles (`BLOCKED_CONFLICT` ·
  `REVIEW_REQUIRED` · `ROLLUP_NOT_INPUT` — readiness `NOT_APPLICABLE` = roll-up
  TOTAL/SUBTOTAL).
- `target_source_expectations.csv` — regiones declaradas (NO-MEDINET y MEDINET).
- `alignment_coverage.csv` — por forma: lado source (match_status) y lado target
  (inventario por rol + cobertura MEDINET input).
- `alignment_manual_review_sample.csv` — muestra determinista y estratificada,
  con `target_kind` / `target_alignment_role` / `target_source_expectation`.
- `summary.json` — cantidades separadas (ver "Cobertura MEDINET").

## Golden reference

El REMASEP final julio se usa como **referencia estructural**, pero
`approval_status = UNCONFIRMED` y
`candidate_golden_status = CANDIDATE_PENDING_SOURCE_COMPLETENESS`. No es un
golden aprobado.

## Privacidad

Sólo `metric_id`, coordenadas, rótulos del formulario, fórmulas y estados. Sin
datos individuales.
"""


COVERAGE_HEADER = [
    "scope",
    # --- lado SOURCE (métricas MEDINET) ---
    "source_count", "eligible_source_count", "exact_matches", "strong_matches",
    "ambiguous", "unmatched_sources", "conflicts",
    # --- lado TARGET (celdas de la plantilla) ---
    "target_total", "input_targets", "formula_targets", "structural_targets",
    "medinet_input_expected", "matched_medinet_input",
    "unmatched_medinet_input_genuine", "non_medinet_input", "unknown_source_input",
]


def _target_form_stats(report: ta.AlignmentReport, form: str | None) -> list:
    """Estadística de targets por rol para una forma (o todas si form=None)."""
    inv = [t for t in report.target_inventory if form is None or t.form == form]
    matched_keys = {
        (r.target_sheet, r.target_cell) for r in report.rows
        if r.match_status in (ta.EXACT_SEMANTIC_MATCH, ta.STRONG_MATCH)
    }
    inputs = [t for t in inv if t.target_alignment_role == ta.ROLE_INPUT_TARGET]
    formulas = [t for t in inv if t.target_alignment_role == ta.ROLE_DERIVED_TARGET]
    structural = [t for t in inv if t.target_alignment_role == ta.ROLE_STRUCTURAL]
    medinet_in = [t for t in inputs if t.target_source_expectation == ta.EXP_MEDINET]
    non_medinet_in = [
        t for t in inputs if t.target_source_expectation in
        (ta.EXP_EGRESOS, ta.EXP_RESOURCE_CALCULATION, ta.EXP_SURGICAL_TABLE,
         ta.EXP_CONTROL_METADATA)
    ]
    unknown_in = [t for t in inputs if t.target_source_expectation == ta.EXP_UNKNOWN]
    matched_medinet_in = [t for t in medinet_in if t.key in matched_keys]
    return [
        len(inv), len(inputs), len(formulas), len(structural),
        len(medinet_in), len(matched_medinet_in),
        len(medinet_in) - len(matched_medinet_in),
        len(non_medinet_in), len(unknown_in),
    ]


def _coverage(report: ta.AlignmentReport, metrics: list) -> list[list]:
    src_by_fk: Counter = Counter((m.form, m.kind) for m in metrics)
    elig = set(report.eligible_source_ids)
    elig_by_fk: Counter = Counter((m.form, m.kind) for m in metrics if m.metric_id in elig)
    fk_by_id = {m.metric_id: (m.form, m.kind) for m in metrics}

    by_fk_status: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for r in report.rows:
        fk = fk_by_id.get(r.source_metric_id)
        if fk:
            by_fk_status[fk][r.match_status] += 1
    unmatched_by_fk: Counter = Counter()
    for u in report.unmatched_sources:
        fk = fk_by_id.get(u.metric_id)
        if fk:
            unmatched_by_fk[fk] += 1

    src_forms = {form for form, _kind in src_by_fk}
    tgt_forms = {t.form for t in report.target_inventory}
    rows: list[list] = []
    for form in sorted(src_forms | tgt_forms):
        for kind in ("BASE_AGGREGATION", "DERIVED_AGGREGATION", "DOWNSTREAM_TOTAL"):
            fk = (form, kind)
            counts = by_fk_status.get(fk, Counter())
            if not src_by_fk.get(fk) and not counts:
                continue
            rows.append([
                f"{form} / {kind}", src_by_fk.get(fk, 0), elig_by_fk.get(fk, 0),
                counts.get(ta.EXACT_SEMANTIC_MATCH, 0), counts.get(ta.STRONG_MATCH, 0),
                counts.get(ta.AMBIGUOUS, 0), unmatched_by_fk.get(fk, 0),
                counts.get(ta.CONFLICT, 0), *([""] * 9),
            ])
        if form in tgt_forms:
            rows.append([f"{form} / *TARGETS*", *([""] * 7),
                         *_target_form_stats(report, form)])
    rows.append(["TOTAL / *TARGETS*", *([""] * 7), *_target_form_stats(report, None)])
    return rows


def write_outputs(analysis: dict, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    report: ta.AlignmentReport = analysis["report"]
    metrics = analysis["metrics"]
    readiness_by_id = analysis["readiness_by_id"]
    metrics_by_id = {m.metric_id: m for m in metrics}
    files: list[Path] = []

    p = output_dir / "target_metric_inventory.csv"
    _csv(p, [
        "target_sheet", "target_cell", "form", "target_kind", "target_alignment_role",
        "section_path", "row_path", "column_path", "sex_value", "age_min", "age_max",
        "aggregation_scope", "procedure_code", "has_formula",
        "target_source_expectation", "target_source_evidence", "semantic_signature",
    ], (
        [t.target_sheet, t.target_cell, t.form, t.target_kind, t.target_alignment_role,
         _ctx(t.section_path), _ctx(t.row_path), _ctx(t.column_path), t.sex_value,
         t.age_min if t.age_min is not None else "",
         t.age_max if t.age_max is not None else "",
         t.aggregation_scope, t.procedure_code, "yes" if t.has_formula else "no",
         t.target_source_expectation, t.target_source_evidence, t.semantic_signature]
        for t in report.target_inventory
    ))
    files.append(p)

    p = output_dir / "source_metric_inventory.csv"
    elig = set(report.eligible_source_ids)
    _csv(p, [
        "metric_id", "form", "kind", "readiness_status", "mapping_status",
        "sex_value", "sex_status", "age_min", "age_max", "aggregation_scope",
        "procedure_code", "eligible", "semantic_signature",
    ], (
        [m.metric_id, m.form, m.kind,
         readiness_by_id[m.metric_id].readiness_status if m.metric_id in readiness_by_id else "",
         m.mapping_status, m.sex_value, m.sex_status,
         m.age_min_years if m.age_min_years is not None else "",
         m.age_max_years if m.age_max_years is not None else "",
         m.aggregation_scope, m.procedure_code_raw,
         "yes" if m.metric_id in elig else "no", m.semantic_signature]
        for m in metrics
    ))
    files.append(p)

    p = output_dir / "semantic_alignment.csv"
    _csv(p, [
        "source_metric_id", "source_form", "source_kind", "source_readiness",
        "target_sheet", "target_cell", "target_kind", "target_alignment_role",
        "match_status", "match_score",
        "source_semantic_signature", "target_semantic_signature",
        "sex_match", "age_match", "procedure_match", "row_match", "column_match",
        "section_match", "target_source_expectation", "needs_human_review",
    ], (
        [r.source_metric_id, r.source_form, r.source_kind, r.source_readiness,
         r.target_sheet, r.target_cell, r.target_kind, r.target_alignment_role,
         r.match_status, r.match_score,
         r.source_semantic_signature, r.target_semantic_signature,
         r.sex_match, r.age_match, r.procedure_match, r.row_match, r.column_match,
         r.section_match, r.target_source_expectation,
         "yes" if r.needs_human_review else "no"]
        for r in sorted(report.rows, key=lambda r: r.source_metric_id)
    ))
    files.append(p)

    p = output_dir / "alignment_evidence.csv"
    _csv(p, ["source_metric_id", "target_sheet", "target_cell", "dimension",
             "source_value", "target_value", "evidence_status"],
         (list(e) for e in report.evidence))
    files.append(p)

    p = output_dir / "ambiguous_matches.csv"
    _csv(p, ["source_metric_id", "candidate_count", "candidate_cells", "reason"],
         ([a.source_metric_id, a.candidate_count, "|".join(a.candidate_cells), a.reason]
          for a in report.ambiguous))
    files.append(p)

    p = output_dir / "unmatched_sources.csv"
    _csv(p, ["source_metric_id", "form", "row_path", "column_path", "reason"],
         ([u.metric_id, u.form, u.row_path, u.column_path, u.reason]
          for u in report.unmatched_sources))
    files.append(p)

    p = output_dir / "unmatched_targets.csv"
    _csv(p, ["target_sheet", "target_cell", "target_kind", "target_alignment_role",
             "target_context", "expected_source", "reason",
             "is_genuine_unmatched_medinet_input"],
         ([u.target_sheet, u.target_cell, u.target_kind, u.target_alignment_role,
           u.target_context, u.expected_source, u.reason,
           "yes" if u.is_genuine_unmatched_medinet_input else "no"]
          for u in report.unmatched_targets))
    files.append(p)

    p = output_dir / "source_exclusions.csv"
    _csv(p, ["metric_id", "form", "readiness_status", "reason", "detail"],
         ([e.metric_id, e.form, e.readiness_status, e.reason, e.detail]
          for e in report.excluded_sources))
    files.append(p)

    p = output_dir / "target_source_expectations.csv"
    _csv(p, ["region_id", "sheet", "expected_source", "section_contains",
             "row_min", "row_max", "evidence"],
         ([r.region_id, r.sheet, r.expected_source, " | ".join(r.section_contains),
           r.row_min if r.row_min is not None else "",
           r.row_max if r.row_max is not None else "", r.evidence]
          for r in analysis["regions"]))
    files.append(p)

    p = output_dir / "alignment_coverage.csv"
    _csv(p, COVERAGE_HEADER, _coverage(report, metrics))
    files.append(p)

    sample = _manual_review_sample(report, metrics_by_id)
    p = output_dir / "alignment_manual_review_sample.csv"
    _csv(p, ["source_metric_id", "target_sheet", "target_cell", "target_kind",
             "target_alignment_role", "target_source_expectation", "source_context",
             "target_context", "match_status", "evidence_summary"],
         ([s["source_metric_id"], s["target_sheet"], s["target_cell"], s["target_kind"],
           s["target_alignment_role"], s["target_source_expectation"],
           s["source_context"], s["target_context"], s["match_status"],
           s["evidence_summary"]] for s in sample))
    files.append(p)

    p = output_dir / "summary.json"
    p.write_text(json.dumps(_summary(analysis, sample), indent=2, ensure_ascii=False),
                 encoding="utf-8")
    files.append(p)

    p = output_dir / "README.md"
    p.write_text(_README, encoding="utf-8")
    files.append(p)
    return files


def _summary(analysis: dict, sample: list[dict]) -> dict:
    report: ta.AlignmentReport = analysis["report"]
    metrics = analysis["metrics"]
    by_status = Counter(r.match_status for r in report.rows)
    by_form_status = defaultdict(Counter)
    for r in report.rows:
        by_form_status[r.source_form][r.match_status] += 1
    excl_by_reason = Counter(e.reason for e in report.excluded_sources)
    unmatched_by_reason = Counter(u.reason for u in report.unmatched_sources)
    tgt_kinds = Counter(t.target_kind for t in report.target_inventory)
    tgt_roles = Counter(t.target_alignment_role for t in report.target_inventory)
    tgt_exp = Counter(t.target_source_expectation for t in report.target_inventory)
    unmatched_tgt_by_reason = Counter(u.reason for u in report.unmatched_targets)
    proc_exact = sum(1 for r in report.rows if r.procedure_match == ta.EV_MATCH)

    inputs = [t for t in report.target_inventory
              if t.target_alignment_role == ta.ROLE_INPUT_TARGET]
    formulas = [t for t in report.target_inventory
                if t.target_alignment_role == ta.ROLE_DERIVED_TARGET]
    structural = [t for t in report.target_inventory
                  if t.target_alignment_role == ta.ROLE_STRUCTURAL]
    _non_medinet = (ta.EXP_EGRESOS, ta.EXP_RESOURCE_CALCULATION,
                    ta.EXP_SURGICAL_TABLE, ta.EXP_CONTROL_METADATA)
    medinet_in = [t for t in inputs if t.target_source_expectation == ta.EXP_MEDINET]
    non_medinet_in = [t for t in inputs if t.target_source_expectation in _non_medinet]
    unknown_in = [t for t in inputs if t.target_source_expectation == ta.EXP_UNKNOWN]
    matched_keys = {
        (r.target_sheet, r.target_cell) for r in report.rows
        if r.match_status in (ta.EXACT_SEMANTIC_MATCH, ta.STRONG_MATCH)
    }
    matched_medinet_in = [t for t in medinet_in if t.key in matched_keys]
    lim_layout = (
        "El layout de la plantilla oficial difiere del generador legacy; "
        "REMASEP 01 se parsea peor que REMASEP_OD."
    )
    lim_regions = (
        "Las regiones de source_regions.yaml (NO-MEDINET y MEDINET) son "
        "preliminares: derivadas de un mes de provenance, sin confirmación funcional."
    )
    lim_unknown = (
        "target_source_expectation = UNKNOWN cuando ninguna región lo reclama: "
        "no se infiere MEDINET por defecto."
    )
    return {
        "generated_at": datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "alignment_policy_version": report.policy_version,
        "alignment_status": "PRELIMINARY_NOT_VALIDATED",
        "generator": analysis["generator"].name,
        "generator_sha256": analysis["generator_sha256"],
        "template": analysis["template"].name,
        "template_sha256": analysis["template_sha256"],
        "template_approval_status": APPROVAL_STATUS,
        "candidate_golden_status": CANDIDATE_GOLDEN_STATUS,
        "source_hash_verified_unchanged": analysis["source_hash_verified_unchanged"],
        "total_source_metrics": len(metrics),
        "eligible_source_metrics": len(report.eligible_source_ids),
        "excluded_source_metrics": len(report.excluded_sources),
        "excluded_by_reason": dict(sorted(excl_by_reason.items())),
        "match_status_counts": dict(sorted(by_status.items())),
        "match_status_by_form": {f: dict(sorted(c.items())) for f, c in sorted(by_form_status.items())},
        "unmatched_sources": len(report.unmatched_sources),
        "unmatched_by_reason": dict(sorted(unmatched_by_reason.items())),
        "ambiguous_sources": len(report.ambiguous),
        "procedure_code_exact_matches": proc_exact,
        # --- inventario físico de targets ---
        "target_inventory_total": len(report.target_inventory),
        "target_by_kind": dict(sorted(tgt_kinds.items())),
        "target_by_alignment_role": dict(sorted(tgt_roles.items())),
        "target_by_source_expectation": dict(sorted(tgt_exp.items())),
        # --- cobertura del alignment MEDINET (sin mezclar conceptos) ---
        "input_targets": len(inputs),
        "formula_targets": len(formulas),
        "structural_targets": len(structural),
        "medinet_input_targets_expected": len(medinet_in),
        "matched_medinet_input_targets": len(matched_medinet_in),
        "genuine_unmatched_medinet_input_targets": len(medinet_in) - len(matched_medinet_in),
        "non_medinet_input_targets": len(non_medinet_in),
        "unknown_source_input_targets": len(unknown_in),
        # --- inventario técnico de filas sin match (todas las clases) ---
        "unmatched_target_rows_technical": len(report.unmatched_targets),
        "unmatched_targets_by_reason": dict(sorted(unmatched_tgt_by_reason.items())),
        "manual_review_sample_size": len(sample),
        "manual_review_sample_by_status": dict(
            sorted(Counter(s["match_status"] for s in sample).items())
        ),
        "limitations": [
            lim_layout,
            "REMASEP B1 no tiene subgrafo MEDINET (todo EGRESOS).",
            lim_regions,
            lim_unknown,
            "Un match no es validación funcional MINSAL.",
        ],
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="align_official_template.py",
        description=(
            "Alineación semántica generador MEDINET -> plantilla oficial REMASEP. "
            "Sólo lectura, sin COM ni macros."
        ),
    )
    parser.add_argument("generator", nargs="?", default=DEFAULT_GENERATOR)
    parser.add_argument("template", nargs="?", default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        analysis = build_alignment_analysis(args.generator, args.template)
    except (ta.TemplateAlignmentError, sm.SemanticMappingError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    files = write_outputs(analysis, Path(args.output))
    summary = _summary(analysis, _manual_review_sample(
        analysis["report"], {m.metric_id: m for m in analysis["metrics"]}))
    report = analysis["report"]

    print(f"generator: {summary['generator']}   template: {summary['template']} "
          f"(approval={summary['template_approval_status']})")
    print(f"  hash de ambos workbooks verificado sin cambios: "
          f"{summary['source_hash_verified_unchanged']}")
    print(f"  source metrics: {summary['total_source_metrics']}  "
          f"elegibles (AUTO_READY): {summary['eligible_source_metrics']}  "
          f"excluidos: {summary['excluded_source_metrics']} {summary['excluded_by_reason']}")
    print(f"  match_status: {summary['match_status_counts']}")
    print(f"  por forma: {summary['match_status_by_form']}")
    print(f"  unmatched sources: {summary['unmatched_sources']} {summary['unmatched_by_reason']}")
    print(f"  ambiguous: {summary['ambiguous_sources']}")
    print(f"  procedure-code exact matches: {summary['procedure_code_exact_matches']}")
    print(f"  TARGET inventario físico: {summary['target_inventory_total']}  "
          f"por rol {summary['target_by_alignment_role']}")
    print(f"  TARGET por expectativa de fuente: {summary['target_by_source_expectation']}")
    print(f"  input targets: {summary['input_targets']}  "
          f"(MEDINET esperados {summary['medinet_input_targets_expected']} · "
          f"NO-MEDINET {summary['non_medinet_input_targets']} · "
          f"UNKNOWN {summary['unknown_source_input_targets']})")
    print(f"  MEDINET input targets alineados: {summary['matched_medinet_input_targets']} / "
          f"{summary['medinet_input_targets_expected']}  "
          f"(genuinos sin source: {summary['genuine_unmatched_medinet_input_targets']})")
    print(f"  filas unmatched_targets (técnico, todas las clases): "
          f"{summary['unmatched_target_rows_technical']} {summary['unmatched_targets_by_reason']}")
    b1_excl = sum(1 for u in report.unmatched_sources if u.reason == ta.REASON_TARGET_NOT_MEDINET)
    print(f"  sources excluidos por target NO-MEDINET: {b1_excl}")
    print(f"  muestra de revisión: {summary['manual_review_sample_size']} "
          f"{summary['manual_review_sample_by_status']}")
    print(f"  candidate_golden_status: {summary['candidate_golden_status']}")
    print(f"Artefactos en {args.output}:")
    for f in files:
        print(f"  - {f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
