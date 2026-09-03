"""Equivalencia de agregaciones legacy directas (Sprint 2.3).

Compara, celda agregada a celda agregada, las fórmulas de ``REMASEP 01``,
``B2 ANEXO`` y ``REMASEP_OD`` que referencian **directamente**
``'Atenciones - Detalles de citas'`` (COUNTIF/COUNTIFS y sumas de ellas):

    Atenciones + AC:AL (legacy_transform)  ->  COUNTIF/COUNTIFS Python
                                           vs
                                           valor cacheado por Excel

- Solo lectura. NO escribe el workbook, NO usa COM ni recalcula.
- La unidad de comparación es una **celda agregada**: los artefactos solo
  contienen fórmulas, coordenadas, columnas y conteos/valores agregados. Ningún
  dato individual de paciente.
- Las caches del workbook pueden estar obsoletas (Sprint 2.2 lo confirmó para AF):
  se separa el estado de evaluación de fórmula del estado de comparación de cache.

Uso:

    python scripts/compare_legacy_aggregations.py \\
        "data/local/GENERACION DATOS REMASEP.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
from inventory_workbook import compute_sha256
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from remasep.adapters.medinet import REQUIRED_FIELDS, is_blank_cell, semantic_column_map
from remasep.core.errors import RemasepError
from remasep.services.legacy_aggregation import (
    UnsupportedFormulaError,
    formula_pattern,
    functions_used,
    padding_can_affect,
    parse_aggregation_formula,
)
from remasep.services.legacy_rules import load_legacy_rules
from remasep.services.legacy_transform import legacy_derived_record

DEFAULT_OUTPUT = "artifacts/legacy_aggregation_equivalence"
TARGET_SHEETS = ("REMASEP 01", "B2 ANEXO", "REMASEP_OD")
DETAIL_SHEET_HINT = "Atenciones - Detalles de citas"

_AGE_RANGE_RE = re.compile(r"!\$?AF\$?:\$?AF", re.IGNORECASE)
_DETAIL_COL_RE = re.compile(r"!\$?([A-Za-z]{1,3})\$?:\$?[A-Za-z]{1,3}")

_STALE_AF_NOTE = (
    "Formula depends on AF; workbook reference contains known stale AF cache."
)

# Valores que Excel deja en las columnas para una fila estructuralmente vacía.
_EMPTY_DERIVED_VALUES: dict[str, object] = {
    "AC": "", "AD": "", "AE": "",
    "AF": 0,
    "AG": "#N/A", "AH": "#N/A", "AI": "#N/A",
    "AJ": "0", "AK": "0", "AL": "0",
}

CLS_DIRECT = "DIRECT_SOURCE_AGGREGATION"
CLS_UNSUPPORTED = "UNSUPPORTED"
CLS_OTHER = "OTHER"

CACHE_MATCH = "MATCH"
CACHE_DIFFERENCE = "CACHE_DIFFERENCE"
CACHE_UNAVAILABLE = "CACHE_UNAVAILABLE"


class AggregationCompareError(RemasepError):
    """Error de entrada del comparador con mensaje claro."""


# ---------------------------------------------------------------------------
# Modelos
# ---------------------------------------------------------------------------


@dataclass
class CellResult:
    sheet: str
    cell: str
    formula: str
    functions: tuple[str, ...]
    referenced_columns: tuple[str, ...]
    depends_on_age: bool
    pattern: str
    classification: str
    cached_available: bool
    cached_value: object = None
    python_value: int | None = None
    unsupported_reason: str = ""
    cache_status: str = ""
    padding_can_affect: bool | None = None
    padding_reason: str = ""
    notes: str = ""


@dataclass
class AggregationResult:
    source_name: str
    source_sha256: str
    detail_sheet: str
    physical_rows: int
    structural_empty_rows: int
    active_records: int
    target_sheets: list[str]
    cells: list[CellResult]
    observed_functions: dict[str, int]
    generated_at: str

    def for_sheet(self, sheet: str) -> list[CellResult]:
        return [c for c in self.cells if c.sheet == sheet]

    @property
    def direct_cells(self) -> list[CellResult]:
        return [c for c in self.cells if c.classification == CLS_DIRECT]

    @property
    def supported_cells(self) -> list[CellResult]:
        return [c for c in self.direct_cells if c.python_value is not None]

    @property
    def unsupported_cells(self) -> list[CellResult]:
        return [c for c in self.cells if c.classification == CLS_UNSUPPORTED]

    @property
    def overall_formula_support_status(self) -> str:
        direct = [c for c in self.cells if c.classification in (CLS_DIRECT, CLS_UNSUPPORTED)]
        if not direct:
            return "PASS"
        if not self.unsupported_cells:
            return "PASS"
        return "PARTIAL" if self.supported_cells else "FAIL"

    @property
    def cache_consistency_status(self) -> str:
        evaluated = self.supported_cells
        if any(c.cache_status == CACHE_DIFFERENCE for c in evaluated):
            return "DIFFERENCES"
        if any(c.cache_status == CACHE_UNAVAILABLE for c in evaluated):
            return "UNAVAILABLE"
        return "PASS"


# ---------------------------------------------------------------------------
# Lectura del workbook
# ---------------------------------------------------------------------------


def _resolve_detail_sheet(workbook: openpyxl.Workbook) -> str:
    best_name, best_score = workbook.sheetnames[0], -1
    for name in workbook.sheetnames:
        ws = workbook[name]
        header = [ws.cell(row=1, column=c).value for c in range(1, (ws.max_column or 0) + 1)]
        score = len(set(semantic_column_map(header).values()) & set(REQUIRED_FIELDS))
        if score > best_score:
            best_name, best_score = name, score
    return best_name


def _row_values(cells: list, columns) -> dict[int, object]:
    return {
        idx: (getattr(cells[idx - 1], "value", None) if idx - 1 < len(cells) else None)
        for idx in columns
    }


def _make_resolver(cached: dict[str, object]):
    def resolve(ref: str) -> object:
        return cached.get(ref.replace("$", ""))

    return resolve


def _build_dataset(ws_v) -> tuple[list[dict], int, int, dict[str, str]]:
    header = [ws_v.cell(row=1, column=c).value for c in range(1, (ws_v.max_column or 0) + 1)]
    raw_map = semantic_column_map(header)
    sem_to_index: dict[str, int] = {}
    for col_index, cell_value in enumerate(header, start=1):
        semantic = raw_map.get(str(cell_value))
        if semantic and semantic not in sem_to_index:
            sem_to_index[semantic] = col_index

    missing = [f for f in REQUIRED_FIELDS if f not in sem_to_index]
    if missing:
        raise AggregationCompareError(
            "La hoja de detalle no expone columnas para: " + ", ".join(missing)
        )
    sem_to_letter = {sem: get_column_letter(idx) for sem, idx in sem_to_index.items()}
    detected = list(sem_to_index)

    ruleset = load_legacy_rules()
    active_rows: list[dict] = []
    structural_empty = 0
    physical = 0
    max_row = ws_v.max_row or 1
    columns_needed = set(sem_to_index.values())

    for v_row in ws_v.iter_rows(min_row=2, max_row=max_row):
        v_by_col = _row_values(list(v_row), columns_needed)
        rec = {sem: v_by_col.get(idx) for sem, idx in sem_to_index.items()}
        physical += 1
        if all(is_blank_cell(rec.get(sem)) for sem in detected):
            structural_empty += 1
            continue
        derived = legacy_derived_record(rec, ruleset)
        combined: dict[str, object] = {sem_to_letter[sem]: rec[sem] for sem in sem_to_index}
        combined.update(derived)
        active_rows.append(combined)

    return active_rows, physical, structural_empty, sem_to_letter


def _scan_target_sheet(ws_f, ws_v, detail_sheet: str) -> tuple[dict[str, str], dict[str, object]]:
    max_row = ws_v.max_row or 1
    formulas: dict[str, str] = {}
    cached: dict[str, object] = {}
    for row_number, (f_row, v_row) in enumerate(
        zip(
            ws_f.iter_rows(min_row=1, max_row=max_row),
            ws_v.iter_rows(min_row=1, max_row=max_row),
            strict=False,
        ),
        start=1,
    ):
        v_cells = list(v_row)
        f_cells = list(f_row)
        for col_index, cell in enumerate(v_cells, start=1):
            value = getattr(cell, "value", None)
            if value is not None:
                cached[f"{get_column_letter(col_index)}{row_number}"] = value
        for col_index, cell in enumerate(f_cells, start=1):
            if getattr(cell, "data_type", None) != "f":
                continue
            value = getattr(cell, "value", None)
            if isinstance(value, str) and detail_sheet in value:
                formulas[f"{get_column_letter(col_index)}{row_number}"] = value
    return formulas, cached


def _detail_columns_in(formula: str) -> tuple[str, ...]:
    return tuple(sorted({m.group(1).upper() for m in _DETAIL_COL_RE.finditer(formula)}))


# ---------------------------------------------------------------------------
# Comparación
# ---------------------------------------------------------------------------


def compare_legacy_aggregations(path: str | Path) -> AggregationResult:
    source = Path(path)
    if not source.is_file():
        raise AggregationCompareError(f"El archivo no existe o no es un archivo: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise AggregationCompareError(f"Formato no soportado: {source.suffix!r}")
    try:
        wb_f = openpyxl.load_workbook(source, data_only=False, read_only=True)
        wb_v = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise AggregationCompareError(f"No se pudo abrir el workbook: {exc}") from exc

    try:
        detail_sheet = _resolve_detail_sheet(wb_v)
        active_rows, physical, structural_empty, _sem_to_letter = _build_dataset(
            wb_v[detail_sheet]
        )
        empty_row_values = _empty_row_values(_sem_to_letter)

        target_sheets = [s for s in TARGET_SHEETS if s in wb_f.sheetnames]
        cells: list[CellResult] = []
        observed: Counter[tuple[str, ...]] = Counter()

        for sheet in target_sheets:
            formulas, cached = _scan_target_sheet(wb_f[sheet], wb_v[sheet], detail_sheet)
            resolve_ref = _make_resolver(cached)

            for cell_coord, formula in sorted(formulas.items(), key=_coord_key):
                funcs = functions_used(formula)
                observed[funcs] += 1
                columns = _detail_columns_in(formula)
                depends_age = _AGE_RANGE_RE.search(formula) is not None
                cached_value = cached.get(cell_coord)
                result = CellResult(
                    sheet=sheet,
                    cell=cell_coord,
                    formula=formula,
                    functions=funcs,
                    referenced_columns=columns,
                    depends_on_age=depends_age,
                    pattern=formula_pattern(formula),
                    classification=CLS_OTHER,
                    cached_available=cached_value is not None,
                    cached_value=cached_value,
                )

                if set(funcs) - {"COUNTIF", "COUNTIFS"}:
                    result.classification = CLS_OTHER
                    cells.append(result)
                    continue

                try:
                    parsed = parse_aggregation_formula(formula, detail_sheet, resolve_ref)
                except UnsupportedFormulaError as exc:
                    result.classification = CLS_UNSUPPORTED
                    result.unsupported_reason = str(exc)
                    cells.append(result)
                    continue

                result.classification = CLS_DIRECT
                result.referenced_columns = parsed.referenced_columns
                result.depends_on_age = parsed.depends_on_age
                result.python_value = parsed.evaluate(active_rows)

                affect, reason = padding_can_affect(parsed, empty_row_values)
                result.padding_can_affect = affect
                result.padding_reason = reason

                result.cache_status, note = _compare_cache(
                    result.python_value, cached_value, parsed.depends_on_age
                )
                notes = [note] if note else []
                if affect:
                    notes.append(
                        f"Padding-sensitive: hasta {structural_empty} filas estructuralmente "
                        "vacías podrían contar; el evaluador usa solo filas activas."
                    )
                result.notes = " ".join(notes)
                cells.append(result)
    finally:
        wb_f.close()
        wb_v.close()

    return AggregationResult(
        source_name=source.name,
        source_sha256=compute_sha256(source),
        detail_sheet=detail_sheet,
        physical_rows=physical,
        structural_empty_rows=structural_empty,
        active_records=len(active_rows),
        target_sheets=target_sheets,
        cells=cells,
        observed_functions={"+".join(k) or "(none)": v for k, v in observed.items()},
        generated_at=datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
    )


def _coord_key(item: tuple[str, object]) -> tuple[int, int]:
    match = re.match(r"([A-Z]+)([0-9]+)", item[0])
    letters, digits = match.group(1), int(match.group(2))
    col = 0
    for ch in letters:
        col = col * 26 + (ord(ch) - 64)
    return digits, col


def _empty_row_values(sem_to_letter: dict[str, str]) -> dict[str, object]:
    values: dict[str, object] = {letter: "" for letter in sem_to_letter.values()}
    values.update(_EMPTY_DERIVED_VALUES)
    return values


def _compare_cache(python_value: int, cached_value: object, depends_on_age: bool) -> tuple[str, str]:
    if cached_value is None:
        return CACHE_UNAVAILABLE, ""
    try:
        same = int(python_value) == int(float(cached_value))
    except (TypeError, ValueError):
        same = str(python_value) == str(cached_value)
    if same:
        return CACHE_MATCH, ""
    return CACHE_DIFFERENCE, (_STALE_AF_NOTE if depends_on_age else "")


# ---------------------------------------------------------------------------
# Escritura de artefactos
# ---------------------------------------------------------------------------

_README = """# Legacy aggregation equivalence (Sprint 2.3)

Salida de `scripts/compare_legacy_aggregations.py`. Compara las fórmulas de
`REMASEP 01`, `B2 ANEXO` y `REMASEP_OD` que referencian **directamente** la hoja
`Atenciones - Detalles de citas` (COUNTIF/COUNTIFS y sumas de ellas) contra:

- **evaluación Python** sobre `Atenciones + AC:AL` (AC:AL vía `legacy_transform`);
- **valor cacheado por Excel** (`data_only=True`) de la misma celda.

## Archivos

- `direct_aggregation_inventory.csv` — toda celda de las 3 hojas cuya fórmula
  referencia la hoja de detalle, con su `classification`
  (`DIRECT_SOURCE_AGGREGATION` / `UNSUPPORTED` / `OTHER`).
- `unsupported_formulas.csv` — fórmulas que usan algún construct fuera del subset.
- `padding_sensitivity.csv` — por fórmula soportada: si alguna fila
  estructuralmente vacía podría contar (`padding_can_affect_result`).
- `aggregation_equivalence.csv` — valor Python vs cache, `cache_comparison_status`
  (`MATCH` / `CACHE_DIFFERENCE` / `CACHE_UNAVAILABLE`).
- `sheet_summary.csv` — conteos por hoja.
- `legacy_metrics_long.csv` — `metric_id` técnico (`LEGACY::<hoja>::<celda>`) ->
  valor agregado. Frontera hacia el modelo semántico (sprint posterior).
- `summary.json` — totales, `formula_support_status`, `cache_consistency_status`.

## Evaluator vs cache

El evaluador Python usa `AF` calculado por `legacy_transform` (edad real). El
workbook de referencia tiene **cache de AF obsoleta** para parte de los registros
(Sprint 2.2). Por eso una `CACHE_DIFFERENCE` en una fórmula `depends_on_age` **no**
implica que Python esté equivocado; se registra con nota y sin convertir a `MATCH`.

## Aviso

Equivalencia con `GENERACION DATOS REMASEP.xlsx` **no** es validación oficial
MINSAL.
"""


def _b(value: object) -> str:
    return "True" if value else "False"


def write_outputs(result: AggregationResult, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    inventory = output_dir / "direct_aggregation_inventory.csv"
    _csv(
        inventory,
        ["target_sheet", "target_cell", "formula", "functions", "referenced_detail_columns",
         "depends_on_age", "formula_pattern", "cached_value_available", "classification"],
        (
            [c.sheet, c.cell, c.formula, "|".join(c.functions), "|".join(c.referenced_columns),
             _b(c.depends_on_age), c.pattern, _b(c.cached_available), c.classification]
            for c in result.cells
        ),
    )

    unsupported = output_dir / "unsupported_formulas.csv"
    _csv(
        unsupported,
        ["target_sheet", "target_cell", "formula", "functions", "reason"],
        (
            [c.sheet, c.cell, c.formula, "|".join(c.functions), c.unsupported_reason]
            for c in result.unsupported_cells
        ),
    )

    padding = output_dir / "padding_sensitivity.csv"
    _csv(
        padding,
        ["target_sheet", "target_cell", "formula", "padding_can_affect_result", "reason"],
        (
            [c.sheet, c.cell, c.formula, _b(c.padding_can_affect), c.padding_reason]
            for c in result.supported_cells
        ),
    )

    equivalence = output_dir / "aggregation_equivalence.csv"
    _csv(
        equivalence,
        ["target_sheet", "target_cell", "formula_pattern", "referenced_detail_columns",
         "python_evaluated_value", "cached_excel_value", "cache_comparison_status",
         "supported", "notes"],
        (
            [
                c.sheet, c.cell, c.pattern, "|".join(c.referenced_columns),
                "" if c.python_value is None else c.python_value,
                "" if c.cached_value is None else c.cached_value,
                c.cache_status or ("" if c.classification != CLS_DIRECT else CACHE_UNAVAILABLE),
                _b(c.python_value is not None),
                c.notes,
            ]
            for c in result.direct_cells
        ),
    )

    sheet_summary = output_dir / "sheet_summary.csv"
    _csv(
        sheet_summary,
        ["sheet", "direct_formula_cells", "supported_cells", "unsupported_cells",
         "cache_matches", "cache_differences", "cache_unavailable"],
        (_sheet_row(result, sheet) for sheet in result.target_sheets),
    )

    metrics = output_dir / "legacy_metrics_long.csv"
    _csv(
        metrics,
        ["metric_id", "generator_sheet", "generator_cell", "value"],
        (
            [f"LEGACY::{c.sheet.replace(' ', '_')}::{c.cell}", c.sheet, c.cell, c.python_value]
            for c in result.supported_cells
        ),
    )

    summary_path = output_dir / "summary.json"
    summary_path.write_text(
        json.dumps(_summary_dict(result), indent=2, ensure_ascii=False), encoding="utf-8"
    )

    readme = output_dir / "README.md"
    readme.write_text(_README, encoding="utf-8")

    return [inventory, unsupported, padding, equivalence, sheet_summary, metrics, summary_path,
            readme]


def _sheet_row(result: AggregationResult, sheet: str) -> list:
    cells = result.for_sheet(sheet)
    direct = [c for c in cells if c.classification == CLS_DIRECT]
    supported = [c for c in direct if c.python_value is not None]
    return [
        sheet,
        len(direct) + len([c for c in cells if c.classification == CLS_UNSUPPORTED]),
        len(supported),
        len([c for c in cells if c.classification == CLS_UNSUPPORTED]),
        len([c for c in supported if c.cache_status == CACHE_MATCH]),
        len([c for c in supported if c.cache_status == CACHE_DIFFERENCE]),
        len([c for c in supported if c.cache_status == CACHE_UNAVAILABLE]),
    ]


def _summary_dict(result: AggregationResult) -> dict:
    supported = result.supported_cells
    direct_like = [c for c in result.cells if c.classification in (CLS_DIRECT, CLS_UNSUPPORTED)]
    padding_sensitive = [c for c in supported if c.padding_can_affect]
    af_cells = [c for c in supported if c.depends_on_age]
    af_differences = [c for c in af_cells if c.cache_status == CACHE_DIFFERENCE]
    return {
        "source": result.source_name,
        "source_sha256": result.source_sha256,
        "detail_sheet": result.detail_sheet,
        "generated_at": result.generated_at,
        "physical_rows": result.physical_rows,
        "structural_empty_rows": result.structural_empty_rows,
        "active_records": result.active_records,
        "target_sheets": result.target_sheets,
        "observed_formula_families": result.observed_functions,
        "direct_formula_cells": len(direct_like),
        "supported_formula_cells": len(supported),
        "unsupported_formula_cells": len(result.unsupported_cells),
        "evaluated_cells": len(supported),
        "cache_matches": len([c for c in supported if c.cache_status == CACHE_MATCH]),
        "cache_differences": len([c for c in supported if c.cache_status == CACHE_DIFFERENCE]),
        "cache_unavailable": len([c for c in supported if c.cache_status == CACHE_UNAVAILABLE]),
        "cells_depending_on_AF": len(af_cells),
        "cache_differences_depending_on_AF": len(af_differences),
        "padding_sensitive_cells": len(padding_sensitive),
        "padding_sensitive": bool(padding_sensitive),
        "formula_support_status": result.overall_formula_support_status,
        "cache_consistency_status": result.cache_consistency_status,
    }


def _csv(path: Path, header: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="compare_legacy_aggregations.py",
        description=(
            "Compara las fórmulas COUNTIF/COUNTIFS que referencian directamente la "
            "hoja de detalle contra su valor cacheado. Solo lectura, sin COM."
        ),
    )
    parser.add_argument("workbook", help="Ruta al workbook legacy (.xlsx/.xlsm).")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT, help=f"(def: {DEFAULT_OUTPUT})")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        result = compare_legacy_aggregations(args.workbook)
    except AggregationCompareError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    files = write_outputs(result, Path(args.output))

    print(f"Workbook: {result.source_name}  (hoja de detalle {result.detail_sheet!r})")
    print(
        f"  filas físicas {result.physical_rows} = "
        f"{result.structural_empty_rows} vacías + {result.active_records} activas"
    )
    print(f"  familias de fórmula observadas: {result.observed_functions}")
    for sheet in result.target_sheets:
        row = _sheet_row(result, sheet)
        print(
            f"  {sheet}: directas {row[1]} | soportadas {row[2]} | no soportadas {row[3]} "
            f"| cache MATCH {row[4]} DIFF {row[5]} UNAVAIL {row[6]}"
        )
    summary = _summary_dict(result)
    print(f"  celdas que dependen de AF: {summary['cells_depending_on_AF']} "
          f"(cache differences dependientes de AF: {summary['cache_differences_depending_on_AF']})")
    print(f"  padding-sensitive: {summary['padding_sensitive_cells']}")
    print(f"  formula_support_status: {summary['formula_support_status']}")
    print(f"  cache_consistency_status: {summary['cache_consistency_status']}")
    print(f"Artefactos en {args.output}:")
    for path in files:
        print(f"  - {path}")
    return 0 if result.overall_formula_support_status != "FAIL" else 1


if __name__ == "__main__":
    raise SystemExit(main())
