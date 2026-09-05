"""Análisis estático de un workbook Excel para ingeniería inversa de REMASEP.

Herramienta de solo lectura. NO modifica el archivo de entrada, no ejecuta
Excel/COM y no almacena valores de celdas (solo estructura y fórmulas).

Uso:

    python scripts/inventory_workbook.py \\
        "data/local/reference/GENERACION DATOS REMASEP.xlsx" \\
        --output "artifacts/workbook_inventory"

Depende únicamente de la librería estándar y de ``openpyxl`` (se lee con
``data_only=False`` para conservar las fórmulas). Funciona en Linux/Ubuntu.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.formula import ArrayFormula

try:  # openpyxl >= 3.1
    from openpyxl.worksheet.formula import DataTableFormula
except ImportError:  # pragma: no cover - defensivo para versiones antiguas
    DataTableFormula = ()  # type: ignore[assignment]


class InventoryError(Exception):
    """Error de entrada/lectura con mensaje apto para el usuario final."""


# ---------------------------------------------------------------------------
# Modelos
# ---------------------------------------------------------------------------


@dataclass
class SheetStats:
    sheet: str
    state: str
    max_row: int
    max_column: int
    non_empty_cells: int
    formula_cells: int
    merged_ranges: list[str] = field(default_factory=list)
    frozen_panes: str | None = None


@dataclass
class FormulaRecord:
    sheet: str
    cell: str
    formula: str
    formula_type: str
    function_names: str
    row: int
    column: int


@dataclass
class PatternRecord:
    pattern: str
    count: int
    function_names: str
    formula_type: str
    sheets: str
    example_sheet: str
    example_cell: str
    example_formula: str


@dataclass
class NamedRangeRecord:
    name: str
    scope: str
    refers_to: str
    hidden: bool


# ---------------------------------------------------------------------------
# Clasificación de fórmulas (heurística, sin parser)
# ---------------------------------------------------------------------------

_AGGREGATION = {
    "SUM", "SUMIF", "SUMIFS", "COUNT", "COUNTA", "COUNTBLANK", "COUNTIF",
    "COUNTIFS", "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "SUMPRODUCT", "SUBTOTAL",
    "AGGREGATE", "MAX", "MAXIFS", "MIN", "MINIFS", "MEDIAN",
}
_LOOKUP = {
    "VLOOKUP", "HLOOKUP", "XLOOKUP", "LOOKUP", "INDEX", "MATCH", "XMATCH",
    "OFFSET", "INDIRECT", "CHOOSE", "GETPIVOTDATA",
}
_LOGICAL = {"IF", "IFS", "IFERROR", "IFNA", "AND", "OR", "NOT", "SWITCH"}
_TEXT = {
    "CONCAT", "CONCATENATE", "TEXTJOIN", "LEFT", "RIGHT", "MID", "TRIM",
    "UPPER", "LOWER", "SUBSTITUTE", "REPLACE", "TEXT", "VALUE", "LEN", "FIND",
    "SEARCH",
}
_DATE = {
    "DATE", "DATEDIF", "YEAR", "MONTH", "DAY", "TODAY", "NOW", "EOMONTH",
    "EDATE", "WEEKDAY", "DATEVALUE",
}

_FUNC_CALL_RE = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z_][A-Za-z0-9_.]*)\s*\(")
_XL_PREFIX_RE = re.compile(r"^_xl(fn|ws|udf|fn\.)\.", re.IGNORECASE)

# Tokens de referencia (con prefijo de hoja opcional, comillas opcionales).
_SHEET = r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!"
_A1 = r"\$?[A-Z]{1,3}\$?[0-9]+"
_COL = r"\$?[A-Z]{1,3}"
_ROW = r"\$?[0-9]+"

_STR_RE = re.compile(r'"(?:[^"]|"")*"')
_COLRANGE_RE = re.compile(rf"(?:{_SHEET})?{_COL}:(?:{_SHEET})?{_COL}")
_ROWRANGE_RE = re.compile(rf"(?:{_SHEET})?{_ROW}:(?:{_SHEET})?{_ROW}")
_RANGE_RE = re.compile(rf"(?:{_SHEET})?{_A1}:(?:{_SHEET})?{_A1}")
_CELL_RE = re.compile(rf"(?:{_SHEET})?{_A1}")
_NUM_RE = re.compile(r"(?<![A-Za-z0-9_])\d+(?:\.\d+)?")
_WS_RE = re.compile(r"\s+")


def extract_function_names(formula: str) -> list[str]:
    """Nombres de función usados en la fórmula, normalizados y ordenados.

    Los literales de texto de Excel (``"..."``) se eliminan antes de buscar
    llamadas para no confundir palabras seguidas de ``(`` dentro de un string
    (p.ej. ``"CONTROL CLINICO (ADULTOS)"``) con funciones reales.
    """
    stripped = _STR_RE.sub(" ", formula)
    names: set[str] = set()
    for raw in _FUNC_CALL_RE.findall(stripped):
        name = _XL_PREFIX_RE.sub("", raw).upper().strip(".")
        if name:
            names.add(name)
    return sorted(names)


def classify_formula(function_names: list[str], formula: str) -> str:
    """Bucket estático grueso. NO es una regla REMASEP, solo categoriza sintaxis."""
    if not function_names:
        return "reference" if not re.search(r"[+\-*/^&]", formula) else "arithmetic"

    fns = set(function_names)
    buckets: list[str] = []
    if fns & _AGGREGATION:
        buckets.append("aggregation")
    if fns & _LOOKUP:
        buckets.append("lookup")
    if fns & _LOGICAL:
        buckets.append("logical")
    if fns & _TEXT:
        buckets.append("text")
    if fns & _DATE:
        buckets.append("date")
    return "+".join(buckets) if buckets else "other"


def normalize_formula(formula: str) -> str:
    """Sustituye literales y referencias por marcadores para agrupar familias.

    Ejemplo::

        =COUNTIFS(A2:A100,"X",B2:B100,"F")
        =COUNTIFS(A2:A100,"Y",B2:B100,"M")

    ambas -> ``=COUNTIFS(<RANGE>,<STR>,<RANGE>,<STR>)``.
    """
    text = formula.strip()
    text = _STR_RE.sub("<STR>", text)
    text = _COLRANGE_RE.sub("<COLRANGE>", text)
    text = _ROWRANGE_RE.sub("<ROWRANGE>", text)
    text = _RANGE_RE.sub("<RANGE>", text)
    text = _CELL_RE.sub("<CELL>", text)
    text = _NUM_RE.sub("<NUM>", text)
    text = _WS_RE.sub("", text)
    return text


# ---------------------------------------------------------------------------
# Lectura del workbook
# ---------------------------------------------------------------------------


def compute_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_workbook_safely(path: Path, *, keep_vba: bool = False) -> openpyxl.Workbook:
    if not path.exists():
        raise InventoryError(f"El archivo no existe: {path}")
    if not path.is_file():
        raise InventoryError(f"La ruta no es un archivo: {path}")
    try:
        return openpyxl.load_workbook(
            path, data_only=False, read_only=False, keep_vba=keep_vba
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
        raise InventoryError(
            f"No se pudo leer el workbook '{path}': {exc}. "
            "¿Es un archivo .xlsx/.xlsm válido y no está corrupto?"
        ) from exc
    except OSError as exc:
        raise InventoryError(f"No se pudo abrir '{path}': {exc}") from exc
    except Exception as exc:  # queremos un mensaje claro, no un traceback
        raise InventoryError(
            f"Error inesperado leyendo '{path}': {exc.__class__.__name__}: {exc}"
        ) from exc


def cell_formula_text(value: object) -> str | None:
    """Texto de la fórmula de una celda openpyxl, o ``None`` si no es fórmula."""
    if isinstance(value, ArrayFormula):
        return value.text or None
    if DataTableFormula and isinstance(value, DataTableFormula):
        return "=<DATA_TABLE>"
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def analyze_worksheet(ws) -> tuple[SheetStats, list[FormulaRecord]]:
    non_empty = 0
    formulas: list[FormulaRecord] = []

    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            non_empty += 1

            formula = cell_formula_text(value) if cell.data_type == "f" else None
            if formula is None and cell.data_type == "f":
                # data_type dice fórmula pero el valor no lo parece; lo registramos igual.
                formula = str(value)
            if formula is None:
                continue

            function_names = extract_function_names(formula)
            formulas.append(
                FormulaRecord(
                    sheet=ws.title,
                    cell=cell.coordinate,
                    formula=formula,
                    formula_type=classify_formula(function_names, formula),
                    function_names="|".join(function_names),
                    row=cell.row,
                    column=cell.column,
                )
            )

    stats = SheetStats(
        sheet=ws.title,
        state=ws.sheet_state,
        max_row=ws.max_row or 0,
        max_column=ws.max_column or 0,
        non_empty_cells=non_empty,
        formula_cells=len(formulas),
        merged_ranges=[str(rng) for rng in sorted(ws.merged_cells.ranges, key=str)],
        frozen_panes=ws.freeze_panes,
    )
    return stats, formulas


def build_patterns(formulas: list[FormulaRecord]) -> list[PatternRecord]:
    groups: dict[str, list[FormulaRecord]] = defaultdict(list)
    for record in formulas:
        groups[normalize_formula(record.formula)].append(record)

    patterns: list[PatternRecord] = []
    for pattern, members in groups.items():
        function_names = sorted({fn for m in members for fn in m.function_names.split("|") if fn})
        sheets = sorted({m.sheet for m in members})
        example = members[0]
        patterns.append(
            PatternRecord(
                pattern=pattern,
                count=len(members),
                function_names="|".join(function_names),
                formula_type=example.formula_type,
                sheets="|".join(sheets),
                example_sheet=example.sheet,
                example_cell=example.cell,
                example_formula=example.formula,
            )
        )

    patterns.sort(key=lambda p: (-p.count, p.pattern))
    return patterns


def collect_named_ranges(wb: openpyxl.Workbook) -> list[NamedRangeRecord]:
    records: list[NamedRangeRecord] = []

    for name, defined in wb.defined_names.items():
        records.append(
            NamedRangeRecord(
                name=name,
                scope="workbook",
                refers_to=defined.value or "",
                hidden=bool(defined.hidden),
            )
        )

    for ws in wb.worksheets:
        sheet_names = getattr(ws, "defined_names", None)
        if not sheet_names:
            continue
        for name, defined in sheet_names.items():
            records.append(
                NamedRangeRecord(
                    name=name,
                    scope=f"sheet:{ws.title}",
                    refers_to=defined.value or "",
                    hidden=bool(defined.hidden),
                )
            )

    records.sort(key=lambda r: (r.scope, r.name))
    return records


# ---------------------------------------------------------------------------
# Escritura de resultados
# ---------------------------------------------------------------------------

_README = """# Workbook Inventory

Salida de `scripts/inventory_workbook.py`: análisis estático y de solo lectura
del archivo de referencia REMASEP. No contiene valores de celdas ni datos
nominales de pacientes, únicamente estructura y fórmulas.

## Archivos

### `workbook_summary.json`
Resumen global:

- `file`: nombre, ruta relativa analizada, tamaño y hash SHA256 del archivo.
- `generated_at`: timestamp UTC de la ejecución.
- `workbook`: nº de hojas, nombres, totales de celdas con datos y de fórmulas,
  nº de named ranges.
- `sheets[]`: por hoja -> `state` (visible/hidden/veryHidden), `max_row`,
  `max_column`, `non_empty_cells`, `formula_cells`, `merged_ranges` (lista y
  conteo) y `frozen_panes`.
- `formula_cells_by_sheet`: conteo de fórmulas por hoja.

### `sheets.csv`
Una fila por hoja: `sheet`, `state`, `max_row`, `max_column`,
`non_empty_cells`, `formula_cells`. Vista tabular rápida de `workbook_summary.json`.

### `formulas.csv`
Una fila por celda que contiene una fórmula:

- `sheet`, `cell` (p.ej. `E4`), `row`, `column` (índice numérico).
- `formula`: texto exacto de la fórmula (con `=`).
- `function_names`: funciones usadas, separadas por `|` (prefijos `_xlfn.`
  eliminados). Vacío si la fórmula no llama funciones.
- `formula_type`: bucket sintáctico grueso, combinable con `+`:
  `aggregation`, `lookup`, `logical`, `text`, `date`, `arithmetic`,
  `reference`, `other`. **No** es una clasificación de reglas REMASEP; solo
  describe la sintaxis.

No se guardan los valores de las celdas fuente.

### `formula_patterns.csv`
Fórmulas agrupadas por "familia". Cada fórmula se normaliza sustituyendo:

- literales de texto -> `<STR>`
- números -> `<NUM>`
- referencias -> `<CELL>`, `<RANGE>`, `<COLRANGE>` (`A:A`), `<ROWRANGE>` (`1:1`)
- se elimina el espacio en blanco

Así `=COUNTIFS(A2:A100,"X",B2:B100,"F")` y `=COUNTIFS(A2:A100,"Y",B2:B100,"M")`
caen en el mismo patrón. Columnas: `pattern`, `count`, `function_names`,
`formula_type`, `sheets`, `example_sheet`, `example_cell`, `example_formula`.
Ordenado por `count` descendente. Es una heurística de tokenización, no un
parser: fórmulas muy distintas rara vez colisionan, pero conviene revisar
`example_formula`.

### `named_ranges.csv`
Named ranges definidos (`name`, `scope` = `workbook` o `sheet:<hoja>`,
`refers_to`, `hidden`). Puede quedar vacío (solo cabecera) si no hay ninguno.

## Notas

- Generado con `openpyxl` (`data_only=False`), sin Excel ni COM.
- El archivo de entrada no se modifica.
- Reejecutar sobrescribe estos archivos.
"""


def _write_csv(path: Path, header: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def write_outputs(
    output_dir: Path,
    *,
    input_path: Path,
    sha256: str,
    sheet_stats: list[SheetStats],
    formulas: list[FormulaRecord],
    patterns: list[PatternRecord],
    named_ranges: list[NamedRangeRecord],
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    summary = {
        "file": {
            "name": input_path.name,
            "path": str(input_path),
            "size_bytes": input_path.stat().st_size,
            "sha256": sha256,
        },
        "generated_at": datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "workbook": {
            "sheet_count": len(sheet_stats),
            "sheet_names": [s.sheet for s in sheet_stats],
            "total_non_empty_cells": sum(s.non_empty_cells for s in sheet_stats),
            "total_formula_cells": sum(s.formula_cells for s in sheet_stats),
            "named_range_count": len(named_ranges),
            "formula_pattern_count": len(patterns),
        },
        "sheets": [
            {
                "name": s.sheet,
                "state": s.state,
                "max_row": s.max_row,
                "max_column": s.max_column,
                "non_empty_cells": s.non_empty_cells,
                "formula_cells": s.formula_cells,
                "merged_range_count": len(s.merged_ranges),
                "merged_ranges": s.merged_ranges,
                "frozen_panes": s.frozen_panes,
            }
            for s in sheet_stats
        ],
        "formula_cells_by_sheet": {s.sheet: s.formula_cells for s in sheet_stats},
    }

    summary_path = output_dir / "workbook_summary.json"
    summary_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")

    sheets_path = output_dir / "sheets.csv"
    _write_csv(
        sheets_path,
        ["sheet", "state", "max_row", "max_column", "non_empty_cells", "formula_cells"],
        (
            [s.sheet, s.state, s.max_row, s.max_column, s.non_empty_cells, s.formula_cells]
            for s in sheet_stats
        ),
    )

    formulas_path = output_dir / "formulas.csv"
    _write_csv(
        formulas_path,
        ["sheet", "cell", "formula", "formula_type", "function_names", "row", "column"],
        (
            [f.sheet, f.cell, f.formula, f.formula_type, f.function_names, f.row, f.column]
            for f in formulas
        ),
    )

    patterns_path = output_dir / "formula_patterns.csv"
    _write_csv(
        patterns_path,
        [
            "pattern", "count", "function_names", "formula_type", "sheets",
            "example_sheet", "example_cell", "example_formula",
        ],
        (
            [
                p.pattern, p.count, p.function_names, p.formula_type, p.sheets,
                p.example_sheet, p.example_cell, p.example_formula,
            ]
            for p in patterns
        ),
    )

    named_ranges_path = output_dir / "named_ranges.csv"
    _write_csv(
        named_ranges_path,
        ["name", "scope", "refers_to", "hidden"],
        ([r.name, r.scope, r.refers_to, r.hidden] for r in named_ranges),
    )

    readme_path = output_dir / "README.md"
    readme_path.write_text(_README, encoding="utf-8")

    return [
        summary_path,
        sheets_path,
        formulas_path,
        patterns_path,
        named_ranges_path,
        readme_path,
    ]


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def inventory_workbook(input_path: Path, output_dir: Path) -> dict:
    """Ejecuta el análisis y escribe los artefactos. Devuelve el resumen."""
    sha256 = compute_sha256(input_path) if input_path.is_file() else ""
    wb = load_workbook_safely(input_path)
    try:
        sheet_stats: list[SheetStats] = []
        all_formulas: list[FormulaRecord] = []
        for ws in wb.worksheets:
            stats, formulas = analyze_worksheet(ws)
            sheet_stats.append(stats)
            all_formulas.extend(formulas)

        patterns = build_patterns(all_formulas)
        named_ranges = collect_named_ranges(wb)
    finally:
        wb.close()

    written = write_outputs(
        output_dir,
        input_path=input_path,
        sha256=sha256,
        sheet_stats=sheet_stats,
        formulas=all_formulas,
        patterns=patterns,
        named_ranges=named_ranges,
    )

    return {
        "output_dir": output_dir,
        "files": written,
        "sheet_count": len(sheet_stats),
        "formula_count": len(all_formulas),
        "pattern_count": len(patterns),
        "named_range_count": len(named_ranges),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="inventory_workbook.py",
        description=(
            "Análisis estático de solo lectura de un workbook Excel "
            "(hojas, fórmulas, patrones y named ranges). No modifica el archivo."
        ),
    )
    parser.add_argument("workbook", help="Ruta al archivo .xlsx/.xlsm a analizar.")
    parser.add_argument(
        "--output",
        "-o",
        default="artifacts/workbook_inventory",
        help="Directorio de salida (se crea si no existe). "
        "Por defecto: artifacts/workbook_inventory",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    input_path = Path(args.workbook)
    output_dir = Path(args.output)

    try:
        result = inventory_workbook(input_path, output_dir)
    except InventoryError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    print(f"Workbook analizado: {input_path}")
    print(
        f"  {result['sheet_count']} hojas, {result['formula_count']} fórmulas, "
        f"{result['pattern_count']} patrones, {result['named_range_count']} named ranges"
    )
    print(f"Artefactos escritos en: {output_dir}")
    for path in result["files"]:
        print(f"  - {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
