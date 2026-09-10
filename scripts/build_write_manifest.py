"""Sprint 3.6 — Writable Target Mapping / write manifest.

Convierte los alignments semánticos (Sprint 3.5) en un manifiesto de instrucciones
de escritura versionable y auditable. **No escribe Excel**, no usa COM, verifica
el SHA256 de generador y plantilla antes y después.

Uso:

    python scripts/build_write_manifest.py \\
        "data/local/GENERACION DATOS REMASEP.xlsx" \\
        "data/local/REMASEP_V1.4 Julio 2026.xlsm"
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from collections import Counter, defaultdict
from datetime import UTC, datetime
from pathlib import Path

import align_official_template as af
import openpyxl

from remasep.core.errors import RemasepError
from remasep.services import writable_target_mapping as wtm

DEFAULT_OUTPUT = "artifacts/writable_target_mapping"
DEFAULT_GENERATOR = "data/local/GENERACION DATOS REMASEP.xlsx"
DEFAULT_TEMPLATE = "data/local/REMASEP_V1.4 Julio 2026.xlsm"

ESTADO_FILTER_STATUS = "PENDING_FUNCTIONAL_CONFIRMATION"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _csv(path: Path, header: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def _lock_str(v: bool | None) -> str:
    return "unknown" if v is None else ("locked" if v else "unlocked")


# ---------------------------------------------------------------------------
# Análisis
# ---------------------------------------------------------------------------


def build_write_analysis(generator_path: str, template_path: str) -> dict:
    gen, tpl = Path(generator_path), Path(template_path)
    for p in (gen, tpl):
        if not p.is_file():
            raise RemasepError(f"no es un archivo: {p}")
    sha_pre = (_sha256(gen), _sha256(tpl))

    alignment = af.build_alignment_analysis(str(gen), str(tpl))
    report = alignment["report"]

    policy = wtm.load_write_policy()

    wb_form = openpyxl.load_workbook(tpl, data_only=False, keep_vba=True)
    wb_vals = openpyxl.load_workbook(tpl, data_only=True, keep_vba=True)
    try:
        tfp = wtm.structural_template_fingerprint(
            wb_form, file_sha256=alignment["template_sha256"]
        )
        # valor actual + revalidación defensiva de la protección por celda-target
        target_values: dict[tuple[str, str], object] = {}
        revalidated_locked: dict[tuple[str, str], bool | None] = {}
        for row in report.rows:
            if row.target_sheet not in wb_vals.sheetnames:
                continue
            key = (row.target_sheet, row.target_cell)
            target_values[key] = wb_vals[row.target_sheet][row.target_cell].value
            prot = getattr(wb_form[row.target_sheet][row.target_cell], "protection", None)
            revalidated_locked[key] = getattr(prot, "locked", None)
    finally:
        wb_form.close()
        wb_vals.close()

    # defensa en profundidad (§17): la protección re-leída del workbook gana.
    write_report = wtm.build_write_mapping(
        report, tfp, policy,
        target_values=target_values, locked_by_target=revalidated_locked,
    )
    write_report.estado_filter_status = ESTADO_FILTER_STATUS

    sha_post = (_sha256(gen), _sha256(tpl))
    if sha_pre != sha_post:
        raise RemasepError("¡el SHA256 de un workbook fuente cambió durante el análisis!")

    return {
        "generator": gen, "template": tpl,
        "generator_sha256": sha_pre[0], "template_sha256": sha_pre[1],
        "source_hash_verified_unchanged": True,
        "alignment": alignment, "policy": policy,
        "template_fingerprint": tfp, "write_report": write_report,
    }


# ---------------------------------------------------------------------------
# Artefactos
# ---------------------------------------------------------------------------

_README = """# Writable target mapping (Sprint 3.6)

Salida de `scripts/build_write_manifest.py`.

Capa por ENCIMA del semantic alignment (Sprint 3.5). Decide, por alignment
resuelto, si puede convertirse en una **instrucción de escritura** automática.
**No escribe Excel.**

## `write_status` (≠ `readiness` ≠ `match_status`)

| estado | significado |
| --- | --- |
| `WRITE_READY` | mapping seguro para automatizar: `EXACT` (o `STRONG` que cumple la policy), target `DIRECT_INPUT_TARGET` / `INPUT_TARGET` / **unlocked** / `MEDINET`, único, sin colisión |
| `WRITE_REVIEW_REQUIRED` | `STRONG` sin evidencia de una dimensión requerida, o expectativa de fuente `UNKNOWN`, o tipo de valor incompatible |
| `WRITE_BLOCKED` | `AMBIGUOUS`, conflicto de dimensiones, source `BLOCKED_CONFLICT`, target **locked**, target NO-MEDINET, o **colisión de escritura** |
| `NOT_WRITABLE` | roll-up calculado por la plantilla, target fórmula, source `REVIEW_REQUIRED`, source no-MEDINET, o sin alignment |

## Identidad

`instruction_id` = hash estable de `source_semantic_signature` +
`target_semantic_signature` + `structural_template_fingerprint_id`. **La
coordenada es ubicación física, no identidad semántica.**

## Template fingerprint

`structural_template_fingerprint_id` ata el manifiesto a la **estructura
semántica** (hojas, orden, layout de fórmulas, celdas desbloqueadas, merges) —
**no** al SHA256 del archivo. Un re-guardado que cambie VBA/hash pero preserve la
estructura **no** invalida el manifiesto; un cambio estructural sí.

## Contrato de período (Sprint 3.7)

Los `MetricValue` MEDINET deben calcularse **exclusivamente** desde
`processing_scope_records` (registros válidos ∩ mes/año) — **no** desde todos los
válidos. El filtro por ESTADO sigue `PENDING_FUNCTIONAL_CONFIRMATION` y **no** se
aplica: 3.6 decide DÓNDE escribir, no QUÉ población entra.

`zero_write_policy = UNRESOLVED`: escribir `0` vs. dejar la celda vacía es una
decisión del *writer* (Sprint 3.7), no de la elegibilidad de escritura.

## Archivos

- `write_readiness.csv` — una fila por source relevante, con `write_status` / `write_reason`.
- `write_manifest_ready.csv` — SÓLO instrucciones `WRITE_READY` (sin PII, sin valores).
- `write_review_queue.csv` + `write_review_clusters.csv` — `STRONG` / `UNKNOWN` a revisar, agrupados por causa.
- `write_blocked.csv` — `AMBIGUOUS`, conflictos, locked, colisiones (con candidatos).
- `write_collisions.csv` — colisiones de escritura detectadas.
- `write_evidence.csv` — evidencia por dimensión (reexportada del alignment).
- `coverage_by_form.csv` / `coverage_by_match_status.csv`.
- `template_fingerprint.json` — `file_sha256` vs. `structural_template_fingerprint`.
- `write_mapping_manual_review_sample.csv` — muestra determinista y estratificada.
- `summary.json`.

Sin PII, sin valores reales del mes. Es un **mapping**, no la salida de una
ejecución.
"""


def _sample(readiness: list[wtm.WriteReadiness]) -> list[wtm.WriteReadiness]:
    rows = sorted(readiness, key=lambda w: w.source_metric_id)
    buckets = [
        lambda w: w.write_status == wtm.WRITE_READY and w.match_status == "EXACT_SEMANTIC_MATCH",
        lambda w: w.write_status == wtm.WRITE_READY and w.match_status == "STRONG_MATCH",
        lambda w: w.write_status == wtm.WRITE_REVIEW_REQUIRED,
        lambda w: w.write_status == wtm.WRITE_BLOCKED and w.match_status == "AMBIGUOUS",
        lambda w: w.write_status == wtm.WRITE_BLOCKED,
        lambda w: w.write_status == wtm.NOT_WRITABLE,
        lambda w: w.source_form == "REMASEP_OD",
        lambda w: w.source_form == "REMASEP_01",
        lambda w: w.source_form == "B2_ANEXO",
        lambda w: w.source_kind == "BASE_AGGREGATION",
        lambda w: w.source_kind == "DERIVED_AGGREGATION",
        lambda w: w.source_kind == "DOWNSTREAM_TOTAL",
        lambda w: "proc=MATCH" in w.evidence_summary,
        lambda w: "sex=MATCH" in w.evidence_summary and "age=MATCH" in w.evidence_summary,
    ]
    picked: dict[str, wtm.WriteReadiness] = {}
    for pred in buckets:
        for w in [x for x in rows if pred(x)][:5]:
            picked[w.source_metric_id] = w
    return sorted(picked.values(), key=lambda w: w.source_metric_id)


def write_outputs(analysis: dict, output_dir: Path) -> list[Path]:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    wr: wtm.WriteMappingReport = analysis["write_report"]
    align_report = analysis["alignment"]["report"]
    files: list[Path] = []

    p = out / "template_fingerprint.json"
    p.write_text(json.dumps(analysis["template_fingerprint"], indent=2, ensure_ascii=False),
                 encoding="utf-8")
    files.append(p)

    p = out / "write_readiness.csv"
    _csv(p, [
        "source_metric_id", "source_semantic_signature", "source_readiness",
        "match_status", "target_sheet", "target_cell", "target_semantic_signature",
        "target_kind", "target_alignment_role", "target_source_expectation",
        "target_locked", "write_status", "write_reason", "expected_value_type",
        "missing_evidence", "needs_human_review",
    ], (
        [w.source_metric_id, w.source_semantic_signature, w.source_readiness,
         w.match_status, w.target_sheet, w.target_cell, w.target_semantic_signature,
         w.target_kind, w.target_alignment_role, w.target_source_expectation,
         _lock_str(w.target_locked), w.write_status, w.write_reason,
         w.expected_value_type, "|".join(w.missing_evidence),
         "yes" if w.needs_human_review else "no"]
        for w in wr.readiness
    ))
    files.append(p)

    p = out / "write_manifest_ready.csv"
    _csv(p, [
        "instruction_id", "source_metric_id", "source_semantic_signature",
        "target_sheet", "target_cell", "target_semantic_signature",
        "expected_value_type", "template_fingerprint_id", "policy_version",
        "match_status",
    ], (
        [i.instruction_id, i.source_metric_id, i.source_semantic_signature,
         i.target_sheet, i.target_cell, i.target_semantic_signature,
         i.expected_value_type, i.template_fingerprint_id, i.policy_version,
         i.match_status]
        for i in wr.instructions
    ))
    files.append(p)

    review = [w for w in wr.readiness if w.write_status == wtm.WRITE_REVIEW_REQUIRED]
    p = out / "write_review_queue.csv"
    _csv(p, [
        "source_metric_id", "form", "match_status", "candidate_target", "reason",
        "missing_evidence", "evidence_summary", "recommended_action",
    ], (
        [w.source_metric_id, w.source_form, w.match_status,
         f"{w.target_sheet}!{w.target_cell}", w.write_reason,
         "|".join(w.missing_evidence), w.evidence_summary, "HUMAN_REVIEW"]
        for w in review
    ))
    files.append(p)

    p = out / "write_review_clusters.csv"
    _csv(p, [
        "cluster_id", "form", "match_status", "write_reason", "missing_evidence",
        "metric_count", "example_metric_ids",
    ], (
        [c.cluster_id, c.form, c.match_status, c.write_reason, c.missing_evidence,
         c.metric_count, "|".join(c.example_metric_ids)]
        for c in wtm.build_review_clusters(wr.readiness)
    ))
    files.append(p)

    blocked = [w for w in wr.readiness if w.write_status == wtm.WRITE_BLOCKED]
    p = out / "write_blocked.csv"
    _csv(p, [
        "source_metric_id", "form", "match_status", "target_sheet", "target_cell",
        "write_reason", "evidence_summary",
    ], (
        [w.source_metric_id, w.source_form, w.match_status, w.target_sheet,
         w.target_cell, w.write_reason, w.evidence_summary]
        for w in blocked
    ))
    files.append(p)

    p = out / "write_collisions.csv"
    _csv(p, ["collision_type", "key", "involved", "detail"],
         ([c.collision_type, c.key, "|".join(c.involved), c.detail]
          for c in wr.collisions))
    files.append(p)

    p = out / "write_evidence.csv"
    _csv(p, ["source_metric_id", "target_sheet", "target_cell", "dimension",
             "source_value", "target_value", "evidence_status"],
         (list(e) for e in align_report.evidence))
    files.append(p)

    p = out / "coverage_by_form.csv"
    by_form: dict[str, Counter] = defaultdict(Counter)
    for w in wr.readiness:
        by_form[w.source_form or "<none>"][w.write_status] += 1
    _csv(p, ["form", wtm.WRITE_READY, wtm.WRITE_REVIEW_REQUIRED, wtm.WRITE_BLOCKED,
             wtm.NOT_WRITABLE, "total"],
         ([f, c.get(wtm.WRITE_READY, 0), c.get(wtm.WRITE_REVIEW_REQUIRED, 0),
           c.get(wtm.WRITE_BLOCKED, 0), c.get(wtm.NOT_WRITABLE, 0), sum(c.values())]
          for f, c in sorted(by_form.items())))
    files.append(p)

    p = out / "coverage_by_match_status.csv"
    by_ms: dict[str, Counter] = defaultdict(Counter)
    for w in wr.readiness:
        by_ms[w.match_status][w.write_status] += 1
    _csv(p, ["match_status", wtm.WRITE_READY, wtm.WRITE_REVIEW_REQUIRED,
             wtm.WRITE_BLOCKED, wtm.NOT_WRITABLE, "total"],
         ([ms, c.get(wtm.WRITE_READY, 0), c.get(wtm.WRITE_REVIEW_REQUIRED, 0),
           c.get(wtm.WRITE_BLOCKED, 0), c.get(wtm.NOT_WRITABLE, 0), sum(c.values())]
          for ms, c in sorted(by_ms.items())))
    files.append(p)

    sample = _sample(wr.readiness)
    p = out / "write_mapping_manual_review_sample.csv"
    _csv(p, [
        "source_metric_id", "source_context", "match_status", "target_sheet",
        "target_cell", "target_context", "target_kind", "target_alignment_role",
        "target_locked", "target_source_expectation", "expected_value_type",
        "write_status", "write_reason", "evidence_summary",
    ], (
        [w.source_metric_id, w.source_semantic_signature, w.match_status,
         w.target_sheet, w.target_cell, w.target_semantic_signature, w.target_kind,
         w.target_alignment_role, _lock_str(w.target_locked),
         w.target_source_expectation, w.expected_value_type, w.write_status,
         w.write_reason, w.evidence_summary]
        for w in sample
    ))
    files.append(p)

    p = out / "summary.json"
    p.write_text(json.dumps(_summary(analysis, sample), indent=2, ensure_ascii=False),
                 encoding="utf-8")
    files.append(p)

    p = out / "README.md"
    p.write_text(_README, encoding="utf-8")
    files.append(p)
    return files


def _summary(analysis: dict, sample: list) -> dict:
    wr: wtm.WriteMappingReport = analysis["write_report"]
    align = analysis["alignment"]
    report = align["report"]
    metrics = align["metrics"]
    by_status = Counter(w.write_status for w in wr.readiness)
    ready = [w for w in wr.readiness if w.write_status == wtm.WRITE_READY]
    ready_by_form = Counter(w.source_form for w in ready)
    ready_by_match = Counter(w.match_status for w in ready)
    review_reasons = Counter(
        (w.write_reason, "|".join(w.missing_evidence))
        for w in wr.readiness if w.write_status == wtm.WRITE_REVIEW_REQUIRED
    )
    blocked_reasons = Counter(
        w.write_reason for w in wr.readiness if w.write_status == wtm.WRITE_BLOCKED
    )
    not_writable_reasons = Counter(
        w.write_reason for w in wr.readiness if w.write_status == wtm.NOT_WRITABLE
    )
    align_status = Counter(r.match_status for r in report.rows)
    # celdas de FÓRMULA físicas en la plantilla (inventario de Sprint 3.5)
    template_formula_targets = sum(
        1 for t in report.target_inventory if t.target_kind == "FORMULA_TARGET"
    )
    # SOURCE rows cuyo candidato resuelto cae en una celda de fórmula (todos son
    # además AMBIGUOUS -> WRITE_BLOCKED). NO es lo mismo que lo anterior.
    source_rows_blocked_by_formula_target = sum(
        1 for w in wr.readiness if w.target_kind == "FORMULA_TARGET"
    )
    locked_blocked = sum(
        1 for w in wr.readiness if w.write_reason == wtm.R_TARGET_LOCKED
    )
    return {
        "generated_at": datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "write_policy_version": wr.policy_version,
        "write_status": "PRELIMINARY_NOT_VALIDATED",
        "generator": analysis["generator"].name,
        "generator_sha256": analysis["generator_sha256"],
        "template": analysis["template"].name,
        "template_file_sha256": analysis["template_sha256"],
        "structural_template_fingerprint_id":
            analysis["template_fingerprint"]["structural_template_fingerprint_id"],
        "source_hash_verified_unchanged": analysis["source_hash_verified_unchanged"],
        "total_source_metrics": len(metrics),
        "source_auto_ready": len(report.eligible_source_ids),
        "resolved_alignments": len(report.rows),
        "alignment_status_counts": dict(sorted(align_status.items())),
        "write_status_counts": dict(sorted(by_status.items())),
        "write_ready_by_form": dict(sorted(ready_by_form.items())),
        "write_ready_by_match_status": dict(sorted(ready_by_match.items())),
        "write_review_reasons": {f"{r}::{m or '-'}": n
                                 for (r, m), n in sorted(review_reasons.items())},
        "write_blocked_reasons": dict(sorted(blocked_reasons.items())),
        "not_writable_reasons": dict(sorted(not_writable_reasons.items())),
        "collision_count": len(wr.collisions),
        "collision_types": dict(sorted(Counter(c.collision_type for c in wr.collisions).items())),
        "template_formula_targets": template_formula_targets,
        "source_rows_blocked_by_formula_target": source_rows_blocked_by_formula_target,
        "locked_targets_blocked": locked_blocked,
        "manifest_instruction_count": len(wr.instructions),
        "review_cluster_count": len(wtm.build_review_clusters(wr.readiness)),
        "zero_write_policy": wr.zero_write_policy,
        "estado_filter_status": wr.estado_filter_status,
        "period_contract": (
            "MetricValue MEDINET = processing_scope_records (válidos ∩ mes/año); "
            "NO all_valid_records; NO filtro ESTADO pendiente de confirmación."
        ),
        "manual_review_sample_size": len(sample),
        "manual_review_sample_by_status":
            dict(sorted(Counter(w.write_status for w in sample).items())),
        "limitations": [
            "3.6 decide DÓNDE escribir, no QUÉ valor ni QUÉ población clínica.",
            ("STRONG con ROW_PATH PARTIAL / SEX MISSING -> WRITE_REVIEW_REQUIRED "
             "(no se bajan umbrales para cobertura)."),
            "Un match no es validación funcional MINSAL.",
            "zero_write_policy UNRESOLVED: precondición del Sprint 3.7.",
        ],
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="build_write_manifest.py",
        description=(
            "Write readiness / manifiesto de escritura desde el semantic alignment. "
            "Sólo lectura, sin COM ni macros."
        ),
    )
    parser.add_argument("generator", nargs="?", default=DEFAULT_GENERATOR)
    parser.add_argument("template", nargs="?", default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        analysis = build_write_analysis(args.generator, args.template)
    except RemasepError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    files = write_outputs(analysis, Path(args.output))
    s = _summary(analysis, _sample(analysis["write_report"].readiness))

    print(f"generator: {s['generator']}   template: {s['template']}")
    print(f"  hash de fuentes verificado sin cambios: {s['source_hash_verified_unchanged']}")
    print(f"  structural_template_fingerprint_id: {s['structural_template_fingerprint_id']}")
    print(f"  file_sha256 (informativo): {s['template_file_sha256'][:16]}…")
    print(f"  source metrics: {s['total_source_metrics']}  AUTO_READY: {s['source_auto_ready']}  "
          f"alineados: {s['resolved_alignments']} {s['alignment_status_counts']}")
    print(f"  write_status: {s['write_status_counts']}")
    print(f"  WRITE_READY por forma: {s['write_ready_by_form']}")
    print(f"  WRITE_READY por match: {s['write_ready_by_match_status']}")
    print(f"  review reasons: {s['write_review_reasons']}")
    print(f"  blocked reasons: {s['write_blocked_reasons']}")
    print(f"  not_writable reasons: {s['not_writable_reasons']}")
    print(f"  colisiones: {s['collision_count']} {s['collision_types']}")
    print(f"  celdas de fórmula físicas en la plantilla: {s['template_formula_targets']}")
    print(f"  source rows bloqueados por candidato de fórmula: "
          f"{s['source_rows_blocked_by_formula_target']}  "
          f"locked targets bloqueados: {s['locked_targets_blocked']}")
    print(f"  manifest instructions: {s['manifest_instruction_count']}  "
          f"review clusters: {s['review_cluster_count']}")
    print(f"  zero_write_policy: {s['zero_write_policy']}  "
          f"estado_filter_status: {s['estado_filter_status']}")
    print(f"  muestra: {s['manual_review_sample_size']} {s['manual_review_sample_by_status']}")
    print(f"Artefactos en {args.output}:")
    for f in files:
        print(f"  - {f}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
