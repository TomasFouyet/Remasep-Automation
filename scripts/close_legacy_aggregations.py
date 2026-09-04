"""Cierre por dependencias de las agregaciones legacy (Sprint 2.4).

Sprint 2.3 dejó 102 fórmulas ``UNSUPPORTED`` de la forma ``COUNTIFS(...) - F30``:
una **agregación base** sobre ``Atenciones - Detalles de citas`` menos otra celda
agregada de la misma hoja. Este script:

- inventaría esas 102 y mide cuántas familias estructurales son;
- construye el **grafo de dependencias** entre celdas agregadas de
  ``REMASEP 01`` / ``B2 ANEXO`` / ``REMASEP_OD``;
- evalúa en **orden topológico**: primero las métricas base, después las
  derivadas (``base ± celdas agregadas``), sin ``eval()`` — AST explícito en
  ``remasep.services.legacy_aggregation``;
- amplía el universo Medinet-dependiente de forma **transitiva** (una celda que
  depende de otra que depende de ``Atenciones``);
- clasifica cada celda: ``BASE_AGGREGATION`` / ``DERIVED_AGGREGATION`` /
  ``DOWNSTREAM_TOTAL`` / ``VALIDATION`` / ``UNSUPPORTED_OTHER``.

Solo lectura. No ejecuta Excel, COM ni macros. No escribe el workbook. Los
artefactos contienen únicamente fórmulas, coordenadas y valores **agregados**:
ninguna fila individual de paciente.

Uso:

    python scripts/close_legacy_aggregations.py \\
        "data/local/GENERACION DATOS REMASEP.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

import openpyxl
from compare_legacy_aggregations import (
    CACHE_DIFFERENCE,
    CACHE_MATCH,
    CACHE_UNAVAILABLE,
    AggregationCompareError,
    _build_dataset,
    _compare_cache,
    _make_resolver,
    _resolve_detail_sheet,
    compute_sha256,
)
from formula_refs import col_to_index, parse_references
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from remasep.services.legacy_aggregation import (
    _VALIDATION_FUNCS,
    UnsupportedFormulaError,
    arithmetic_constructs,
    formula_pattern,
    parse_derived_formula,
)

DEFAULT_OUTPUT = "artifacts/legacy_aggregation_closure"
TARGET_SHEETS = ("REMASEP 01", "B2 ANEXO", "REMASEP_OD")

KIND_BASE = "BASE_AGGREGATION"
KIND_DERIVED = "DERIVED_AGGREGATION"
KIND_TOTAL = "DOWNSTREAM_TOTAL"
KIND_VALIDATION = "VALIDATION"
KIND_OTHER = "UNSUPPORTED_OTHER"

STATUS_OK = "OK"
STATUS_MISSING = "MISSING_DEPENDENCY"
STATUS_CYCLE = "CYCLE"

_AGE_RANGE_RE = re.compile(r"!\$?AF\$?:\$?AF", re.IGNORECASE)
_DETAIL_COL_RE = re.compile(r"!\$?([A-Za-z]{1,3})\$?:\$?[A-Za-z]{1,3}")
_COORD_RE = re.compile(r"^([A-Za-z]{1,3})([0-9]+)$")


class ClosureError(AggregationCompareError):
    """Error de entrada del cierre de agregaciones, con mensaje claro."""


# ---------------------------------------------------------------------------
# Modelo
# ---------------------------------------------------------------------------


@dataclass
class Node:
    sheet: str
    cell: str
    formula: str
    kind: str
    direct: bool
    detail_columns: tuple[str, ...]
    value_ref_coords: tuple[str, ...]  # refs usadas como VALOR (misma hoja)
    depends_on_age_local: bool
    constructs: tuple[str, ...] = ()
    # dependencias en el universo Medinet (cross/same sheet, celdas y rangos):
    graph_deps: set[tuple[str, str]] = field(default_factory=set)
    # resultado:
    medinet: bool = False
    depth: int = 0
    depends_on_age: bool = False
    python_value: object = None
    cached_value: object = None
    cache_status: str = ""
    evaluation_supported: bool = False
    status: str = STATUS_OK
    notes: str = ""

    @property
    def key(self) -> tuple[str, str]:
        return (self.sheet, self.cell)

    @property
    def metric_id(self) -> str:
        return f"LEGACY::{self.sheet.replace(' ', '_')}::{self.cell}"


@dataclass
class ClosureResult:
    source_name: str
    source_sha256: str
    detail_sheet: str
    physical_rows: int
    structural_empty_rows: int
    active_records: int
    nodes: dict[tuple[str, str], Node]
    edges: list[tuple[str, str, str, str, str]]
    unsupported_family_rows: list[list]
    unsupported_family_count: int
    unsupported_variant_count: int
    max_depth: int
    max_depth_evaluable: int
    cycle_nodes: list[tuple[str, str]]
    missing_dependencies: list[tuple[str, str, str]]
    generated_at: str

    def medinet_nodes(self) -> list[Node]:
        return [n for n in self.nodes.values() if n.medinet]

    def by_kind(self, kind: str) -> list[Node]:
        return [n for n in self.medinet_nodes() if n.kind == kind]

    def evaluable(self) -> list[Node]:
        return [n for n in self.medinet_nodes() if n.kind in (KIND_BASE, KIND_DERIVED)]

    @property
    def formula_support_status(self) -> str:
        base = self.by_kind(KIND_BASE)
        derived = self.by_kind(KIND_DERIVED)
        if any(not n.evaluation_supported for n in base):
            return "FAIL"
        if derived and all(n.evaluation_supported for n in derived):
            return "PASS"
        if not derived:
            return "PASS"
        return "PARTIAL"

    @property
    def cache_consistency_status(self) -> str:
        evaluated = [n for n in self.evaluable() if n.evaluation_supported]
        if any(n.cache_status == CACHE_DIFFERENCE for n in evaluated):
            return "DIFFERENCES"
        if any(n.cache_status == CACHE_UNAVAILABLE for n in evaluated):
            return "UNAVAILABLE"
        return "PASS"


# ---------------------------------------------------------------------------
# Lectura del workbook
# ---------------------------------------------------------------------------


def _scan_sheet(ws_f, ws_v) -> tuple[dict[str, str], dict[str, object]]:
    """(fórmulas, valores cacheados) de una hoja, por coordenada (``"F46"``)."""
    max_row = ws_v.max_row or 1
    formulas: dict[str, str] = {}
    cached: dict[str, object] = {}
    rows_f = ws_f.iter_rows(min_row=1, max_row=max_row)
    rows_v = ws_v.iter_rows(min_row=1, max_row=max_row)
    for row_number, (f_row, v_row) in enumerate(zip(rows_f, rows_v, strict=False), start=1):
        for col_index, cell in enumerate(list(v_row), start=1):
            value = getattr(cell, "value", None)
            if value is not None:
                cached[f"{get_column_letter(col_index)}{row_number}"] = value
        for col_index, cell in enumerate(list(f_row), start=1):
            if getattr(cell, "data_type", None) != "f":
                continue
            value = getattr(cell, "value", None)
            if isinstance(value, str):
                formulas[f"{get_column_letter(col_index)}{row_number}"] = value
    return formulas, cached


def _detail_columns_regex(formula: str) -> tuple[str, ...]:
    return tuple(sorted({m.group(1).upper() for m in _DETAIL_COL_RE.finditer(formula)}))


def _universe_refs(
    formula: str, host_sheet: str, sheet_set: set[str]
) -> tuple[set[tuple[str, str]], list[tuple[str, object]]]:
    """Referencias de la fórmula que caen en las hojas objetivo (no la de detalle).

    Devuelve ``(celdas, rangos)`` para el grafo de dependencias transitivo.
    """
    parsed = parse_references(formula)
    cells: set[tuple[str, str]] = set()
    ranges: list[tuple[str, object]] = []
    for ref in parsed.references:
        ref_sheet = ref.sheet or host_sheet
        if ref_sheet not in sheet_set:
            continue
        if ref.ref_type == "cell":
            cells.add((ref_sheet, f"{ref.col_start}{ref.row_start}"))
        else:
            ranges.append((ref_sheet, ref))
    return cells, ranges


def _ref_contains(ref, col_i: int, row_i: int) -> bool:
    in_cols = not ref.col_start or (
        col_to_index(ref.col_start) <= col_i <= col_to_index(ref.col_end)
    )
    in_rows = not ref.row_start or (ref.row_start <= row_i <= ref.row_end)
    return in_cols and in_rows


def _range_covers(ref, sheet_cells: dict[str, set[str]], ref_sheet: str) -> set[tuple[str, str]]:
    covered: set[tuple[str, str]] = set()
    for coord in sheet_cells.get(ref_sheet, set()):
        match = _COORD_RE.match(coord)
        if match and _ref_contains(ref, col_to_index(match.group(1)), int(match.group(2))):
            covered.add((ref_sheet, coord))
    return covered


# ---------------------------------------------------------------------------
# Clasificación
# ---------------------------------------------------------------------------

_VALIDATION_NAMES = set(_VALIDATION_FUNCS)
_BAD_OPS = {"operator:*", "operator:/", "operator:&", "grouping_parens"}


def _classify(formula: str, detail_sheet: str, resolver, direct: bool):
    """(kind, value_ref_coords, detail_columns, depends_on_age_local, constructs, parsed)."""
    try:
        parsed = parse_derived_formula(formula, detail_sheet, resolver)
    except UnsupportedFormulaError:
        constructs = arithmetic_constructs(formula)
        fns = {c.split(":", 1)[1] for c in constructs if c.startswith("function:")}
        bad = {c for c in constructs if c in _BAD_OPS}
        if fns & _VALIDATION_NAMES:
            kind = KIND_VALIDATION
        elif not bad and fns <= {"SUM"}:
            kind = KIND_TOTAL
        else:
            kind = KIND_OTHER
        return (
            kind,
            (),
            _detail_columns_regex(formula),
            _AGE_RANGE_RE.search(formula) is not None,
            constructs,
            None,
        )

    if parsed.is_pure_aggregation:
        kind = KIND_BASE
    elif parsed.is_derived_aggregation:
        kind = KIND_DERIVED
    else:
        kind = KIND_TOTAL
    return (
        kind,
        parsed.value_ref_coords,
        parsed.referenced_columns,
        parsed.depends_on_age_local,
        (),
        parsed,
    )


# ---------------------------------------------------------------------------
# Grafo
# ---------------------------------------------------------------------------


def _topo_order(
    deps_of: dict[tuple[str, str], set[tuple[str, str]]],
) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    """Kahn sobre ``deps_of``. Devuelve (orden, nodos_en_ciclo)."""
    universe = set(deps_of)
    indeg = {k: 0 for k in universe}
    dependents: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for node, deps in deps_of.items():
        for dep in deps:
            if dep in universe:
                indeg[node] += 1
                dependents[dep].append(node)
    queue = [k for k, d in indeg.items() if d == 0]
    order: list[tuple[str, str]] = []
    while queue:
        current = queue.pop()
        order.append(current)
        for child in dependents[current]:
            indeg[child] -= 1
            if indeg[child] == 0:
                queue.append(child)
    cyclic = [k for k in universe if k not in set(order)]
    return order, cyclic


def _compute_depth(
    order: list[tuple[str, str]],
    deps_of: dict[tuple[str, str], set[tuple[str, str]]],
) -> dict[tuple[str, str], int]:
    """Profundidad = longitud del camino más largo desde una raíz (raíz = 0)."""
    depth: dict[tuple[str, str], int] = {}
    for node in order:
        node_deps = [d for d in deps_of.get(node, ()) if d in deps_of]
        depth[node] = 1 + max((depth.get(d, 0) for d in node_deps), default=-1)
    return depth


# ---------------------------------------------------------------------------
# Orquestación
# ---------------------------------------------------------------------------


def close_legacy_aggregations(path: str | Path) -> ClosureResult:
    source = Path(path)
    if not source.is_file():
        raise ClosureError(f"El archivo no existe o no es un archivo: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ClosureError(f"Formato no soportado: {source.suffix!r}")
    try:
        wb_f = openpyxl.load_workbook(source, data_only=False, read_only=True)
        wb_v = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ClosureError(f"No se pudo abrir el workbook: {exc}") from exc

    try:
        detail_sheet = _resolve_detail_sheet(wb_v)
        active_rows, physical, structural_empty, _sem = _build_dataset(wb_v[detail_sheet])

        present = [s for s in TARGET_SHEETS if s in wb_f.sheetnames]
        sheet_set = set(present)

        formulas: dict[str, dict[str, str]] = {}
        cached: dict[str, dict[str, object]] = {}
        for sheet in present:
            formulas[sheet], cached[sheet] = _scan_sheet(wb_f[sheet], wb_v[sheet])
    finally:
        wb_f.close()
        wb_v.close()

    sheet_cells = {s: set(formulas[s]) for s in present}

    nodes: dict[tuple[str, str], Node] = {}
    parsed_of: dict[tuple[str, str], object] = {}
    range_deps_of: dict[tuple[str, str], list[tuple[str, object]]] = {}

    for sheet in present:
        resolver = _make_resolver(cached[sheet])
        for cell, formula in formulas[sheet].items():
            direct = detail_sheet in formula
            kind, value_refs, detail_cols, age_local, constructs, parsed = _classify(
                formula, detail_sheet, resolver, direct
            )
            node = Node(
                sheet=sheet,
                cell=cell,
                formula=formula,
                kind=kind,
                direct=direct,
                detail_columns=tuple(detail_cols),
                value_ref_coords=tuple(value_refs),
                depends_on_age_local=age_local,
                constructs=tuple(constructs),
            )
            cells_ref, ranges_ref = _universe_refs(formula, sheet, sheet_set)
            node.graph_deps = cells_ref
            range_deps_of[node.key] = ranges_ref
            nodes[node.key] = node
            parsed_of[node.key] = parsed

    # --- expandir rangos a celdas concretas del universo -----------------
    for key, ranges in range_deps_of.items():
        for ref_sheet, ref in ranges:
            nodes[key].graph_deps |= _range_covers(ref, sheet_cells, ref_sheet)
    # descartar auto-referencias y referencias fuera del universo de fórmulas
    for node in nodes.values():
        node.graph_deps = {d for d in node.graph_deps if d in nodes and d != node.key}

    # --- cierre Medinet transitivo -------------------------------------
    for node in nodes.values():
        node.medinet = node.direct
    changed = True
    while changed:
        changed = False
        for node in nodes.values():
            if node.medinet:
                continue
            if any(nodes[d].medinet for d in node.graph_deps):
                node.medinet = True
                changed = True

    medinet_keys = {k for k, n in nodes.items() if n.medinet}
    deps_med = {
        k: {d for d in nodes[k].graph_deps if d in medinet_keys} for k in medinet_keys
    }

    # --- profundidad + ciclos sobre el grafo Medinet completo -----------
    order, cyclic = _topo_order(deps_med)
    depth = _compute_depth(order, deps_med)
    for key, value in depth.items():
        nodes[key].depth = value
    for key in cyclic:
        nodes[key].status = STATUS_CYCLE
        nodes[key].depth = -1

    # --- propagación depends_on_AF (orden topológico) ------------------
    for key in order:
        node = nodes[key]
        node.depends_on_age = node.depends_on_age_local or any(
            nodes[d].depends_on_age for d in deps_med[key]
        )
    for key in cyclic:
        nodes[key].depends_on_age = nodes[key].depends_on_age_local

    # --- evaluación topológica (solo BASE + DERIVED) -------------------
    eval_keys = {k for k in medinet_keys if nodes[k].kind in (KIND_BASE, KIND_DERIVED)}
    eval_deps = {
        k: {(nodes[k].sheet, c) for c in nodes[k].value_ref_coords} & eval_keys
        for k in eval_keys
    }
    eval_order, eval_cyclic = _topo_order(eval_deps)
    missing: list[tuple[str, str, str]] = []

    for key in eval_order:
        node = nodes[key]
        parsed = parsed_of[key]
        if node.kind == KIND_BASE:
            node.python_value = parsed.evaluate(active_rows, {})
            node.evaluation_supported = True
            continue
        lookup: dict[str, float] = {}
        unresolved = None
        for coord in node.value_ref_coords:
            dep = nodes.get((node.sheet, coord))
            if dep is None or not dep.evaluation_supported or dep.python_value is None:
                unresolved = coord
                break
            lookup[coord] = float(dep.python_value)
        if unresolved is not None:
            node.status = STATUS_MISSING
            node.evaluation_supported = False
            missing.append((node.sheet, node.cell, unresolved))
            continue
        node.python_value = parsed.evaluate(active_rows, lookup)
        node.evaluation_supported = True

    for key in eval_cyclic:
        nodes[key].status = STATUS_CYCLE
        nodes[key].evaluation_supported = False

    # --- comparación de cache ----------------------------------------
    for node in nodes.values():
        if not node.evaluation_supported or node.python_value is None:
            continue
        cached_value = cached[node.sheet].get(node.cell)
        node.cached_value = cached_value
        node.cache_status, note = _compare_cache(
            node.python_value, cached_value, node.depends_on_age
        )
        if note:
            node.notes = note

    # --- aristas + inventario Parte A -------------------------------
    edges: list[tuple[str, str, str, str, str]] = []
    for key in sorted(medinet_keys, key=_key_sort):
        node = nodes[key]
        value_set = {(node.sheet, c) for c in node.value_ref_coords}
        for dep in sorted(deps_med[key], key=_key_sort):
            same = dep[0] == node.sheet
            is_value = dep in value_set
            dep_type = (
                ("SAME_SHEET_VALUE" if same else "CROSS_SHEET_VALUE")
                if is_value
                else ("SAME_SHEET_RANGE" if same else "CROSS_SHEET_RANGE")
            )
            edges.append((dep[0], dep[1], node.sheet, node.cell, dep_type))

    family_rows, family_count, variant_count = _unsupported_family_inventory(
        nodes, medinet_keys, parsed_of
    )

    max_depth = max((n.depth for n in nodes.values() if n.medinet), default=0)
    max_depth_evaluable = max(
        (nodes[k].depth for k in eval_keys if nodes[k].depth >= 0), default=0
    )

    return ClosureResult(
        source_name=source.name,
        source_sha256=compute_sha256(source),
        detail_sheet=detail_sheet,
        physical_rows=physical,
        structural_empty_rows=structural_empty,
        active_records=len(active_rows),
        nodes=nodes,
        edges=edges,
        unsupported_family_rows=family_rows,
        unsupported_family_count=family_count,
        unsupported_variant_count=variant_count,
        max_depth=max_depth,
        max_depth_evaluable=max_depth_evaluable,
        cycle_nodes=[*cyclic, *(k for k in eval_cyclic if k not in cyclic)],
        missing_dependencies=missing,
        generated_at=datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ"),
    )


def _key_sort(key: tuple[str, str]) -> tuple[str, int, int]:
    match = _COORD_RE.match(key[1])
    if not match:
        return (key[0], 0, 0)
    return (key[0], int(match.group(2)), col_to_index(match.group(1)))


def _unsupported_family_inventory(
    nodes: dict[tuple[str, str], Node],
    medinet_keys: set[tuple[str, str]],
    parsed_of: dict[tuple[str, str], object],
) -> tuple[list[list], int, int]:
    """Inventario de las derivadas que Sprint 2.3 marcaba ``UNSUPPORTED``
    (``COUNTIF(S)(...) ± celda``). Devuelve ``(filas, familias, variantes)``:
    *familias* = combinación conceptual (operadores / nº de bases / nº de refs);
    *variantes* = además el patrón normalizado de la fórmula."""
    rows: list[list] = []
    families: set[tuple] = set()
    variants: set[tuple] = set()
    derived = sorted(
        (nodes[k] for k in medinet_keys if nodes[k].kind == KIND_DERIVED),
        key=lambda n: _key_sort(n.key),
    )
    for node in derived:
        parsed = parsed_of.get(node.key)
        base_calls = parsed.base_call_count if parsed is not None else 0
        signs: set[str] = set()
        if parsed is not None:
            signs = {"-" if s < 0 else "+" for s, _ in parsed.value_refs} | {
                "-" if s < 0 else "+" for s, _ in parsed.literals
            }
        operators = "|".join(sorted(signs))
        same_refs = "|".join(node.value_ref_coords)
        depth_candidate = _depth_candidate(node, nodes)
        pattern = formula_pattern(node.formula)
        rows.append(
            [
                node.sheet, node.cell, node.formula, base_calls, same_refs,
                operators, pattern, depth_candidate,
            ]
        )
        families.add((operators, base_calls, len(node.value_ref_coords)))
        variants.add((operators, base_calls, len(node.value_ref_coords), pattern))
    return rows, len(families), len(variants)


def _depth_candidate(node: Node, nodes: dict[tuple[str, str], Node]) -> object:
    depths = []
    for coord in node.value_ref_coords:
        dep = nodes.get((node.sheet, coord))
        if dep is None or dep.kind not in (KIND_BASE, KIND_DERIVED):
            return "unresolved"
        depths.append(dep.depth)
    return 1 + max(depths, default=-1)


# ---------------------------------------------------------------------------
# Artefactos
# ---------------------------------------------------------------------------

_README = """# Legacy aggregation closure (Sprint 2.4)

Salida de `scripts/close_legacy_aggregations.py`. Amplía Sprint 2.3: además de
las agregaciones **base** (`COUNTIF`/`COUNTIFS` sobre `Atenciones - Detalles de
citas`), reproduce las agregaciones **derivadas** `base ± celda agregada`
mediante un grafo de dependencias evaluado en orden topológico.

## Archivos

- `unsupported_family_inventory.csv` — las derivadas que Sprint 2.3 marcaba
  `UNSUPPORTED` (`COUNTIF(S)(...) ± celda`), con nº de familias estructurales.
- `aggregation_dependency_edges.csv` — aristas `origen -> destino` entre celdas
  agregadas (`dependency_type`: `SAME_SHEET_VALUE`, `SAME_SHEET_RANGE`, …).
- `aggregation_dependency_summary.csv` — por nodo: grado de entrada/salida,
  profundidad y `status` (`OK` / `CYCLE` / `MISSING_DEPENDENCY`).
- `transitive_medinet_cells.csv` — toda celda de las 3 hojas que depende de
  `Atenciones`, directa o transitivamente.
- `metric_nodes.csv` — DAG de métricas: `metric_id`, `kind`, `dependencies`,
  `value`, `depends_on_age`.
- `equivalence.csv` — valor Python vs. valor cacheado por Excel para las celdas
  evaluadas (`BASE` + `DERIVED`), con `cache_comparison_status`.
- `sheet_summary.csv` — conteos por hoja.
- `summary.json` — totales y estados (`formula_support_status`,
  `cache_consistency_status` — separados).

## Evaluador vs. cache

Igual que Sprint 2.3: se separa el estado de evaluación de fórmula del estado de
comparación de cache (`MATCH` / `CACHE_DIFFERENCE` / `CACHE_UNAVAILABLE`). La
señal `depends_on_age` se **propaga transitivamente**: si `F30` depende de `AF`,
entonces `G30 = COUNTIFS(...) - F30` también. Una `CACHE_DIFFERENCE` en una
fórmula `depends_on_age` se anota (cache de `AF` obsoleta, Sprint 2.2) y **no**
se convierte en `MATCH`.

## Aviso

Equivalencia con `GENERACION DATOS REMASEP.xlsx` **no** es validación oficial
MINSAL.
"""


def _fmt(value: object) -> str:
    return "" if value is None else str(value)


def _b(value: object) -> str:
    return "True" if value else "False"


def _csv(path: Path, header: list[str], rows) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)


def write_outputs(result: ClosureResult, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    med = sorted(result.medinet_nodes(), key=lambda n: _key_sort(n.key))

    family = output_dir / "unsupported_family_inventory.csv"
    _csv(
        family,
        ["sheet", "cell", "formula", "base_aggregation_count", "same_sheet_references",
         "arithmetic_operators", "formula_pattern", "dependency_depth_candidate"],
        result.unsupported_family_rows,
    )

    edges = output_dir / "aggregation_dependency_edges.csv"
    _csv(
        edges,
        ["source_sheet", "source_cell", "target_sheet", "target_cell", "dependency_type"],
        result.edges,
    )

    summary_graph = output_dir / "aggregation_dependency_summary.csv"
    indeg: dict[tuple[str, str], int] = defaultdict(int)
    outdeg: dict[tuple[str, str], int] = defaultdict(int)
    for s_sheet, s_cell, t_sheet, t_cell, _type in result.edges:
        outdeg[(t_sheet, t_cell)] += 1
        indeg[(s_sheet, s_cell)] += 1
    _csv(
        summary_graph,
        ["node", "dependency_count", "dependent_count", "depth", "status"],
        (
            [f"{n.sheet}!{n.cell}", outdeg[n.key], indeg[n.key], n.depth, n.status]
            for n in med
        ),
    )

    transitive = output_dir / "transitive_medinet_cells.csv"
    _csv(
        transitive,
        ["sheet", "cell", "direct_or_transitive", "dependency_depth", "root_detail_columns",
         "formula", "evaluation_supported"],
        (
            [
                n.sheet, n.cell, "direct" if n.direct else "transitive", n.depth,
                "|".join(_root_detail_columns(n, result.nodes)), n.formula,
                _b(n.evaluation_supported),
            ]
            for n in med
        ),
    )

    metrics = output_dir / "metric_nodes.csv"
    _csv(
        metrics,
        ["metric_id", "sheet", "cell", "kind", "dependencies", "detail_columns", "value",
         "depends_on_age"],
        (
            [
                n.metric_id, n.sheet, n.cell, n.kind,
                "|".join(
                    nodes_metric_id(result.nodes, dep) for dep in sorted(
                        {d for d in result.nodes[n.key].graph_deps if d in result.nodes},
                        key=_key_sort,
                    )
                ),
                "|".join(n.detail_columns),
                _fmt(n.python_value),
                _b(n.depends_on_age),
            ]
            for n in med
        ),
    )

    equivalence = output_dir / "equivalence.csv"
    _csv(
        equivalence,
        ["sheet", "cell", "classification", "python_value", "cached_excel_value",
         "cache_comparison_status", "evaluation_supported", "depends_on_AF", "notes"],
        (
            [
                n.sheet, n.cell, n.kind, _fmt(n.python_value), _fmt(n.cached_value),
                n.cache_status or (CACHE_UNAVAILABLE if n.evaluation_supported else ""),
                _b(n.evaluation_supported), _b(n.depends_on_age), n.notes,
            ]
            for n in med
            if n.kind in (KIND_BASE, KIND_DERIVED)
        ),
    )

    sheet_summary = output_dir / "sheet_summary.csv"
    _csv(
        sheet_summary,
        ["sheet", "base_cells", "derived_cells", "downstream_total_cells", "validation_cells",
         "unsupported_other_cells", "base_supported", "derived_supported", "cache_matches",
         "cache_differences", "cache_unavailable"],
        (_sheet_row(result, sheet) for sheet in _present_sheets(result)),
    )

    summary_path = output_dir / "summary.json"
    summary_path.write_text(
        json.dumps(_summary_dict(result), indent=2, ensure_ascii=False), encoding="utf-8"
    )

    readme = output_dir / "README.md"
    readme.write_text(_README, encoding="utf-8")

    return [family, edges, summary_graph, transitive, metrics, equivalence, sheet_summary,
            summary_path, readme]


def nodes_metric_id(nodes: dict[tuple[str, str], Node], key: tuple[str, str]) -> str:
    node = nodes.get(key)
    return node.metric_id if node else f"LEGACY::{key[0].replace(' ', '_')}::{key[1]}"


def _root_detail_columns(node: Node, nodes: dict[tuple[str, str], Node]) -> list[str]:
    seen: set[tuple[str, str]] = set()
    stack = [node.key]
    columns: set[str] = set()
    while stack:
        current = stack.pop()
        if current in seen:
            continue
        seen.add(current)
        cur_node = nodes.get(current)
        if cur_node is None:
            continue
        columns.update(cur_node.detail_columns)
        stack.extend(d for d in cur_node.graph_deps if d in nodes)
    return sorted(columns)


def _present_sheets(result: ClosureResult) -> list[str]:
    ordered = [s for s in TARGET_SHEETS if any(n.sheet == s for n in result.medinet_nodes())]
    return ordered


def _sheet_row(result: ClosureResult, sheet: str) -> list:
    cells = [n for n in result.medinet_nodes() if n.sheet == sheet]
    base = [n for n in cells if n.kind == KIND_BASE]
    derived = [n for n in cells if n.kind == KIND_DERIVED]
    evaluated = [n for n in base + derived if n.evaluation_supported]
    return [
        sheet,
        len(base),
        len(derived),
        len([n for n in cells if n.kind == KIND_TOTAL]),
        len([n for n in cells if n.kind == KIND_VALIDATION]),
        len([n for n in cells if n.kind == KIND_OTHER]),
        len([n for n in base if n.evaluation_supported]),
        len([n for n in derived if n.evaluation_supported]),
        len([n for n in evaluated if n.cache_status == CACHE_MATCH]),
        len([n for n in evaluated if n.cache_status == CACHE_DIFFERENCE]),
        len([n for n in evaluated if n.cache_status == CACHE_UNAVAILABLE]),
    ]


def _summary_dict(result: ClosureResult) -> dict:
    med = result.medinet_nodes()
    base = result.by_kind(KIND_BASE)
    derived = result.by_kind(KIND_DERIVED)
    evaluated = [n for n in base + derived if n.evaluation_supported]
    af_cells = [n for n in med if n.depends_on_age]
    af_diff = [
        n for n in evaluated if n.cache_status == CACHE_DIFFERENCE and n.depends_on_age
    ]
    return {
        "source": result.source_name,
        "source_sha256": result.source_sha256,
        "detail_sheet": result.detail_sheet,
        "generated_at": result.generated_at,
        "physical_rows": result.physical_rows,
        "structural_empty_rows": result.structural_empty_rows,
        "active_records": result.active_records,
        "direct_medinet_cells": len([n for n in med if n.direct]),
        "transitive_medinet_cells": len([n for n in med if not n.direct]),
        "total_medinet_cells": len(med),
        "base_aggregation_cells": len(base),
        "derived_aggregation_cells": len(derived),
        "downstream_total_cells": len(result.by_kind(KIND_TOTAL)),
        "validation_cells": len(result.by_kind(KIND_VALIDATION)),
        "unsupported_other_cells": len(result.by_kind(KIND_OTHER)),
        "base_supported": len([n for n in base if n.evaluation_supported]),
        "derived_supported": len([n for n in derived if n.evaluation_supported]),
        "unsupported_family_count": result.unsupported_family_count,
        "unsupported_structural_variants": result.unsupported_variant_count,
        "max_dag_depth": result.max_depth,
        "max_dag_depth_evaluable": result.max_depth_evaluable,
        "cycles": len(result.cycle_nodes),
        "missing_dependencies": len(result.missing_dependencies),
        "cache_matches": len([n for n in evaluated if n.cache_status == CACHE_MATCH]),
        "cache_differences": len([n for n in evaluated if n.cache_status == CACHE_DIFFERENCE]),
        "cache_unavailable": len([n for n in evaluated if n.cache_status == CACHE_UNAVAILABLE]),
        "cache_differences_depending_on_AF": len(af_diff),
        "depends_on_AF_cells": len(af_cells),
        "formula_support_status": result.formula_support_status,
        "cache_consistency_status": result.cache_consistency_status,
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="close_legacy_aggregations.py",
        description=(
            "Cierre por dependencias de las agregaciones legacy: evalúa las "
            "derivadas 'base ± celda agregada' en orden topológico. Solo lectura."
        ),
    )
    parser.add_argument("workbook", help="Ruta al workbook legacy (.xlsx/.xlsm).")
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT, help=f"(def: {DEFAULT_OUTPUT})")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        result = close_legacy_aggregations(args.workbook)
    except ClosureError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    files = write_outputs(result, Path(args.output))
    summary = _summary_dict(result)

    print(f"Workbook: {result.source_name}  (hoja de detalle {result.detail_sheet!r})")
    print(
        f"  filas físicas {result.physical_rows} = "
        f"{result.structural_empty_rows} vacías + {result.active_records} activas"
    )
    print(
        f"  celdas Medinet-dependent: {summary['total_medinet_cells']} "
        f"({summary['direct_medinet_cells']} directas + "
        f"{summary['transitive_medinet_cells']} transitivas)"
    )
    print(
        f"  base {summary['base_aggregation_cells']} "
        f"(soportadas {summary['base_supported']}) | "
        f"derived {summary['derived_aggregation_cells']} "
        f"(soportadas {summary['derived_supported']}) | "
        f"downstream_total {summary['downstream_total_cells']} | "
        f"validation {summary['validation_cells']} | "
        f"unsupported_other {summary['unsupported_other_cells']}"
    )
    print(
        f"  derivadas: {summary['unsupported_family_count']} familia(s) conceptual(es), "
        f"{summary['unsupported_structural_variants']} variantes estructurales | "
        f"profundidad máx DAG: {summary['max_dag_depth']} (evaluables: "
        f"{summary['max_dag_depth_evaluable']}) | ciclos: {summary['cycles']} | "
        f"dependencias faltantes: {summary['missing_dependencies']}"
    )
    for sheet in _present_sheets(result):
        row = _sheet_row(result, sheet)
        print(
            f"  {sheet}: base {row[1]}/{row[6]} | derived {row[2]}/{row[7]} | "
            f"total {row[3]} | validation {row[4]} | other {row[5]} | "
            f"cache MATCH {row[8]} DIFF {row[9]} UNAVAIL {row[10]}"
        )
    print(
        f"  cache: {summary['cache_matches']} MATCH / {summary['cache_differences']} DIFF "
        f"/ {summary['cache_unavailable']} UNAVAIL "
        f"(diferencias dependientes de AF: {summary['cache_differences_depending_on_AF']})"
    )
    print(f"  depends_on_AF: {summary['depends_on_AF_cells']} celdas")
    print(f"  formula_support_status: {summary['formula_support_status']}")
    print(f"  cache_consistency_status: {summary['cache_consistency_status']}")
    print(f"Artefactos en {args.output}:")
    for path in files:
        print(f"  - {path}")
    return 0 if result.formula_support_status != "FAIL" else 1


if __name__ == "__main__":
    raise SystemExit(main())
