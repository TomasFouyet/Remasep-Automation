"""Sprint C (F02) — congela el contrato semántico de plantilla (dev-only).

Lee TODAS las celdas de fórmula estática de la plantilla oficial en blanco
(``data/local/REMASEP 2026_V1.4.xlsm`` por defecto) y las congela en
``config/excel_writer_2026/``:

    template_static_formulas.csv   sheet,cell,formula — una fila por cada
                                    celda con ``data_type == "f"`` en la
                                    plantilla oficial.
    template_semantic_contract.yaml  versión + fingerprint estructural
                                    esperado + sha256 del CSV.

Estrategia (reemplaza el contrato de 28 celdas "críticas" curadas a mano,
insuficiente porque dejaba pasar el repro original de la auditoría en
REMASEP 01!C15): certificar el texto de **toda** fórmula estática del
template, sin resolver dependencias. Un input manual/valor mensual nunca es
``data_type == "f"`` en la plantilla EN BLANCO, así que queda fuera sin
necesidad de curarlo a mano — y se verifica automáticamente que ningún
target de ``write_manifest.csv`` coincide con una celda de fórmula (si
coincidiera, el propio contrato bloquearía cada generación real).

Esta es una herramienta de desarrollo: **su salida se versiona**, el workbook
NO (vive en ``data/local/``, gitignored). Se regenera sólo explícitamente
cuando cambia la plantilla oficial — nunca en tiempo de ejecución de la app.

Uso::

    python scripts/build_template_semantic_contract.py \\
        --template "data/local/REMASEP 2026_V1.4.xlsm" \\
        --write-manifest config/runtime_2026/write_manifest.csv
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import sys
from pathlib import Path

import openpyxl
import yaml

DEFAULT_TEMPLATE = "data/local/REMASEP 2026_V1.4.xlsm"
DEFAULT_WRITE_MANIFEST = "config/runtime_2026/write_manifest.csv"
DEFAULT_OUT_DIR = "config/excel_writer_2026"
CSV_NAME = "template_static_formulas.csv"
YAML_NAME = "template_semantic_contract.yaml"


def _structural_fingerprint(template: Path) -> str:
    from remasep.services.writable_target_mapping import structural_template_fingerprint

    wb = openpyxl.load_workbook(template, data_only=False, keep_vba=True)
    try:
        return structural_template_fingerprint(wb)["structural_template_fingerprint_id"]
    finally:
        wb.close()


def _extract_static_formulas(template: Path) -> list[tuple[str, str, str]]:
    wb = openpyxl.load_workbook(template, data_only=False, keep_vba=True, read_only=True)
    rows: list[tuple[str, str, str]] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        rows.append((ws.title, cell.coordinate, str(cell.value)))
    finally:
        wb.close()
    return rows


def _read_write_manifest_targets(path: Path) -> set[tuple[str, str]]:
    if not path.is_file():
        return set()
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        return {(row["target_sheet"], row["target_cell"]) for row in reader}


def build(*, template: Path, write_manifest: Path, out_dir: Path) -> int:
    rows = _extract_static_formulas(template)
    targets = _read_write_manifest_targets(write_manifest)

    collisions = sorted({(s, c) for s, c, _ in rows} & targets)
    if collisions:
        print(
            f"ABORTADO: {len(collisions)} target(s) de write_manifest.csv son "
            "celdas de fórmula en la plantilla en blanco (contradice el "
            "contrato: un target debe ser siempre un input, nunca una "
            f"fórmula). Ejemplos: {collisions[:5]}",
            file=sys.stderr,
        )
        return 1

    out_dir.mkdir(parents=True, exist_ok=True)
    csv_path = out_dir / CSV_NAME
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["sheet", "cell", "formula"])
        for sheet, cell, formula in rows:
            writer.writerow([sheet, cell, formula])

    csv_sha256 = hashlib.sha256(csv_path.read_bytes()).hexdigest()
    per_sheet: dict[str, int] = {}
    for sheet, _cell, _formula in rows:
        per_sheet[sheet] = per_sheet.get(sheet, 0) + 1

    yaml_path = out_dir / YAML_NAME
    doc = {
        "version": "template_semantic_contract_2026",
        "strategy": "ALL_STATIC_TEMPLATE_FORMULAS",
        "structural_template_fingerprint_id": _structural_fingerprint(template),
        "formulas_file": CSV_NAME,
        "formulas_file_sha256": csv_sha256,
        "total_formulas": len(rows),
        "formulas_per_sheet": dict(sorted(per_sheet.items())),
        "notes": (
            "Congela TODA celda de formula estatica (data_type=='f') de la "
            "plantilla oficial en blanco -- no solo un subconjunto curado a "
            "mano. Un input manual/valor mensual nunca es una formula en la "
            "plantilla en blanco, asi que queda fuera automaticamente. "
            "Verificado en build time: ningun target de write_manifest.csv "
            "coincide con una celda de formula. Regenerar SOLO con este "
            "script cuando cambie la plantilla oficial -- nunca en runtime."
        ),
    }
    yaml_path.write_text(
        yaml.safe_dump(doc, sort_keys=False, allow_unicode=True, default_flow_style=False),
        encoding="utf-8",
    )

    print(f"OK: {len(rows)} formulas estaticas en {len(per_sheet)} hojas -> {csv_path}")
    for sheet, n in sorted(per_sheet.items()):
        print(f"  {sheet}: {n}")
    print(f"contrato: {yaml_path}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="build_template_semantic_contract.py")
    parser.add_argument("--template", default=DEFAULT_TEMPLATE)
    parser.add_argument("--write-manifest", default=DEFAULT_WRITE_MANIFEST)
    parser.add_argument("--out-dir", default=DEFAULT_OUT_DIR)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    template = Path(args.template)
    if not template.is_file():
        print(f"ABORTADO: no existe la plantilla {template}", file=sys.stderr)
        return 1
    return build(
        template=template,
        write_manifest=Path(args.write_manifest),
        out_dir=Path(args.out_dir),
    )


if __name__ == "__main__":
    raise SystemExit(main())
